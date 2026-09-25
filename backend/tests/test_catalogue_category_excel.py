from io import BytesIO
from types import SimpleNamespace
from zipfile import ZipFile

from openpyxl import Workbook
from PIL import Image

from app import catalogue_share_links


def test_product_image_is_embedded_as_a_bounded_thumbnail(tmp_path, monkeypatch):
    source_path = tmp_path / "large-product.bmp"
    Image.effect_noise((1800, 1800), 100).convert("RGB").save(source_path)
    monkeypatch.setattr(
        catalogue_share_links.storage,
        "resolve",
        lambda _key: source_path,
    )

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["Image"])
    worksheet.append([""])
    product = SimpleNamespace(main_image_url="/uploads/products/large-product.bmp")

    catalogue_share_links._add_public_product_image(worksheet, product, 2)
    output = BytesIO()
    workbook.save(output)

    with ZipFile(BytesIO(output.getvalue())) as archive:
        media_files = [name for name in archive.namelist() if name.startswith("xl/media/")]
        assert len(media_files) == 1
        embedded = archive.read(media_files[0])

    assert len(embedded) < 50_000
    with Image.open(BytesIO(embedded)) as thumbnail:
        assert thumbnail.width <= 144
        assert thumbnail.height <= 144


def test_missing_product_image_gets_a_readable_placeholder():
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["Image"])
    worksheet.append([""])
    product = SimpleNamespace(main_image_url="")

    catalogue_share_links._add_public_product_image(worksheet, product, 2)

    assert worksheet.cell(2, 1).value == "Image unavailable"
