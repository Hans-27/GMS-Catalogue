"""Secure multi-audience catalogue link acceptance smoke test.

Run from ``backend`` with: ``python -m tests.smoke_catalogue_share_links``.
"""

import os
import shutil
import tempfile
import uuid
from datetime import UTC, datetime, timedelta
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from urllib.parse import urlparse

test_directory = Path(tempfile.mkdtemp(prefix="gms-share-links-test-"))
os.environ["DATABASE_URL"] = f"sqlite+pysqlite:///{(test_directory / 'catalogue.db').as_posix()}"
os.environ["SECRET_KEY"] = "catalogue-share-link-smoke-secret"
os.environ["AUTO_CREATE_TABLES"] = "true"
os.environ["SEED_DEMO_DATA"] = "true"
os.environ["PUBLIC_APP_URL"] = "http://catalogue.test"

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader
from sqlalchemy import select

from app.commerce_models import CatalogueShareLink, PriceList, ProductPrice
from app.database import SessionLocal
from app.erp_models import ErpCustomerPriceLevel, ErpProductCustomerPrice
from app.main import app
from app.models import AuditLog, Product, Role, User
from app.security import hash_password


def ok(response, expected=200):
    assert response.status_code == expected, response.text
    return response.json()


def token_from(link: dict) -> str:
    return urlparse(link["public_url"]).path.rsplit("/", 1)[-1]


def main() -> None:
    try:
        with TestClient(app) as client:
            with SessionLocal() as db:
                role = db.scalar(select(Role).where(Role.name == "superadmin"))
                product = db.scalar(select(Product).order_by(Product.sku))
                assert role and product
                user = User(username="share.admin", email="share.admin@example.com", full_name="Share Administrator", password_hash=hash_password("StrongPassword123!"), roles=[role])
                sales_role = db.scalar(select(Role).where(Role.name == "sales_user"))
                customer_role = db.scalar(select(Role).where(Role.name == "customer_user"))
                assert sales_role and customer_role
                second_user = User(username="share.second", email="share.second@example.com", full_name="Second Sales User", password_hash=hash_password("StrongPassword456!"), roles=[sales_role])
                customer_user = User(username="share.customer", email="share.customer@example.com", full_name="Shared Customer", password_hash=hash_password("CustomerPassword123!"), roles=[customer_role])
                db.add_all([user, second_user, customer_user]); db.flush()
                now = datetime.now(UTC) - timedelta(minutes=1)
                lists = {item.code: item for item in db.scalars(select(PriceList))}
                amounts = {"NORMAL": "100.00", "BIG_CUSTOMER": "90.00", "VIP": "80.00", "VVIP": "70.00"}
                for order, (code, amount) in enumerate(amounts.items(), start=1):
                    level = ErpCustomerPriceLevel(
                        erp_price_type_id=order,
                        source_code=f"SP{order}",
                        source_name=code.replace("_", " ").title(),
                        price_list_id=lists[code].id,
                        product_count=1,
                        sort_order=order,
                        is_active=True,
                    )
                    db.add(level)
                    db.flush()
                    db.add(ErpProductCustomerPrice(
                        product_id=product.id,
                        price_level_id=level.id,
                        erp_product_id=product.sku,
                        amount=Decimal(amount),
                        currency="THB",
                    ))
                    db.add(ProductPrice(product_id=product.id, price_list_id=lists[code].id, amount=Decimal(amount), currency="THB", effective_from=now, status="active", reason="Share-link test price.", created_by_id=user.id, approved_by_id=user.id, approved_at=now))
                card_prices = {
                    "SP5": (19477, "DEALER", "110.00"),
                    "SP6": (19478, "WHOLESALE", "90.00"),
                    "SRP": (4481, "RETAIL", "120.00"),
                }
                for sort_order, (source_code, (source_id, list_code, amount)) in enumerate(card_prices.items(), start=5):
                    level = ErpCustomerPriceLevel(
                        erp_price_type_id=source_id,
                        source_code=source_code,
                        source_name=source_code,
                        price_list_id=lists[list_code].id,
                        product_count=1,
                        sort_order=sort_order,
                        is_active=True,
                    )
                    db.add(level)
                    db.flush()
                    db.add(ErpProductCustomerPrice(
                        product_id=product.id,
                        price_level_id=level.id,
                        erp_product_id=product.sku,
                        amount=Decimal(amount),
                        currency="THB",
                    ))
                product_id = str(product.id)
                normal_id = lists["NORMAL"].id
                db.commit()

            ok(client.post("/api/auth/login", json={"identifier": "share.admin", "password": "StrongPassword123!", "remember_me": False}))
            audiences = ok(client.get("/api/v1/catalogue-audience-types"))
            assert [item["code"] for item in audiences] == ["normal", "big_customer_vip", "vip", "vip_province", "vvip", "no_price"]
            assert next(item for item in audiences if item["code"] == "vip")["display_name"] == "VIP BKK"

            catalogue = ok(client.post("/api/v1/catalogues", json={"title": "Secure Audience Catalogue", "slug": "secure-audience-catalogue", "description": "Audience pricing security test.", "brand": "TestBrand", "audience": "Customers", "price_list_id": normal_id, "show_prices": True, "currency": "THB", "language": "en-th", "valid_from": None, "valid_until": None, "is_public": True}), 201)
            catalogue_id = catalogue["id"]
            ok(client.put(f"/api/v1/catalogues/{catalogue_id}/products", json={"products": [{"product_id": product_id, "section_title": "Products", "override_description": "Version one description", "hide_price": False}]}))
            cover_image = BytesIO()
            Image.new("RGB", (64, 64), (240, 20, 20)).save(cover_image, format="PNG")
            uploaded_cover = ok(
                client.post(
                    f"/api/v1/catalogues/{catalogue_id}/cover/assets",
                    data={"asset_type": "background", "alt_text": "Red PDF cover test"},
                    files={"file": ("cover.png", cover_image.getvalue(), "image/png")},
                ),
                201,
            )
            assert uploaded_cover["asset_type"] == "background"
            ok(client.post(f"/api/v1/catalogues/{catalogue_id}/publish"))
            draft_catalogue = ok(client.post("/api/v1/catalogues", json={"title": "Sales Hidden Draft", "slug": "sales-hidden-draft", "description": "Must never be visible to Sales accounts.", "brand": "TestBrand", "audience": "Customers", "price_list_id": normal_id, "show_prices": True, "currency": "THB", "language": "en-th", "valid_from": None, "valid_until": None, "is_public": False}), 201)
            draft_catalogue_id = draft_catalogue["id"]

            links = ok(client.get(f"/api/v1/catalogues/{catalogue_id}/share-links"))
            assert len(links) == 6
            assert len({token_from(link) for link in links}) == 6
            cards = ok(client.get("/api/v1/catalogue-share-links/cards"))
            assert len(cards[catalogue_id]) == 6

            by_code = {item["audience_code"]: item for item in links}
            expected_prices = {"normal": "100.00", "big_customer_vip": "90.00", "vip": "80.00", "vip_province": "80.00", "vvip": "70.00"}
            for code, expected in expected_prices.items():
                token = token_from(by_code[code])
                public = ok(client.get(f"/api/v1/public/catalogues/{token}", params={"price_list": "NO_PRICE", "audience": "no_price"}))
                assert public["audience_code"] == code
                assert public["products"][0]["price"] == expected
                assert public["products"][0]["wholesale_price"] == "90.00"
                assert public["products"][0]["online_price"] == "120.00"
                assert public["products"][0]["retail_price"] == "110.00"

            # One shared staff/customer portal account can issue separate,
            # customer-scoped links. The opaque token selects the stored
            # profile; recipients never change pricing with query parameters.
            siam_link = ok(client.post(
                f"/api/v1/catalogues/{catalogue_id}/share-links",
                json={
                    "audience_type_id": by_code["normal"]["audience_type_id"],
                    "customer_name": "Siam Retail Co.",
                },
            ), 201)
            metro_link = ok(client.post(
                f"/api/v1/catalogues/{catalogue_id}/share-links",
                json={
                    "audience_type_id": by_code["vip"]["audience_type_id"],
                    "customer_name": "Metro Dealer Co.",
                },
            ), 201)
            assert siam_link["customer_code"] == "SIAM-RETAIL-CO"
            assert metro_link["customer_code"] == "METRO-DEALER-CO"
            assert token_from(siam_link) != token_from(metro_link)
            assert client.post(
                f"/api/v1/catalogues/{catalogue_id}/share-links",
                json={
                    "audience_type_id": by_code["vip"]["audience_type_id"],
                    "customer_name": "Siam Retail Co.",
                },
            ).status_code == 409
            siam_public = ok(client.get(
                f"/api/v1/public/catalogues/{token_from(siam_link)}",
                params={"audience": "vip", "price_list": "VVIP"},
            ))
            metro_public = ok(client.get(
                f"/api/v1/public/catalogues/{token_from(metro_link)}",
                params={"audience": "normal", "price_list": "NORMAL"},
            ))
            assert siam_public["customer_name"] == "Siam Retail Co."
            assert siam_public["products"][0]["price"] == "100.00"
            assert metro_public["customer_name"] == "Metro Dealer Co."
            assert metro_public["products"][0]["price"] == "80.00"
            ok(client.post("/api/auth/login", json={"identifier": "share.customer", "password": "CustomerPassword123!", "remember_me": False}))
            customer_portal = ok(client.get(
                "/api/v1/customer-portal",
                headers={"X-Customer-Access": token_from(siam_link)},
            ))
            assert customer_portal["customer"] == {
                "code": "SIAM-RETAIL-CO",
                "name": "Siam Retail Co.",
                "audience_code": "normal",
                "audience_name": "Normal",
            }
            assert [item["id"] for item in customer_portal["catalogues"]] == [catalogue_id]
            assert customer_portal["catalogues"][0]["public_url"] == siam_link["public_url"]
            assert customer_portal["catalogues"][0]["pdf_url"].endswith(f"/{token_from(siam_link)}/pdf")
            test_brand_price = next(
                item for item in customer_portal["brand_prices"]
                if item["brand"] == "TestBrand"
            )
            assert test_brand_price["price_list_code"] == "SP1"
            assert client.get("/api/v1/customer-portal", headers={"X-Customer-Access": "x" * 32}).status_code == 404
            ok(client.post("/api/auth/login", json={"identifier": "share.admin", "password": "StrongPassword123!", "remember_me": False}))
            assert len(ok(client.get(
                "/api/v1/catalogue-share-links/cards",
                params={"catalogue_ids": catalogue_id},
            ))[catalogue_id]) == 6
            no_price_token = token_from(by_code["no_price"])
            no_price = ok(client.get(f"/api/v1/public/catalogues/{no_price_token}", params={"price_list": "VIP"}))
            assert no_price["show_prices"] is False
            assert "catalogue_id" not in no_price and "status" not in no_price
            assert no_price["products"][0]["id"] == product_id
            assert "product_status" not in no_price["products"][0]
            assert "price" not in no_price["products"][0]
            assert "currency" not in no_price["products"][0]
            assert "wholesale_price" not in no_price["products"][0]
            assert "online_price" not in no_price["products"][0]
            assert "retail_price" not in no_price["products"][0]
            assert "storage_key" not in no_price["cover"]["assets"][0]
            public_cover_url = no_price["cover"]["assets"][0]["preview_url"]
            assert f"/api/v1/public/catalogues/{no_price_token}/cover/assets/" in public_cover_url
            with TestClient(app) as public_client:
                public_cover = public_client.get(public_cover_url)
            assert public_cover.status_code == 200, public_cover.text
            assert public_cover.headers["content-type"] == "image/webp"

            mappings = ok(client.get("/api/v1/my-catalogue-price-mappings"))
            assert len(mappings) == 6
            remapped = ok(
                client.put(
                    "/api/v1/my-catalogue-price-mappings",
                    json={
                        "mappings": [
                            {
                                "audience_type_id": item["audience_type_id"],
                                "price_list_id": normal_id
                                if item["audience_code"] == "vip"
                                else item["price_list_id"],
                            }
                            for item in mappings
                        ]
                    },
                )
            )
            assert next(item for item in remapped if item["audience_code"] == "vip")["price_list_code"] == "NORMAL"
            updated_links = ok(client.get(f"/api/v1/catalogues/{catalogue_id}/share-links"))
            assert next(item for item in updated_links if item["audience_code"] == "vip")["price_list_name"] == "Normal"
            assert ok(client.get(f"/api/v1/public/catalogues/{token_from(by_code['vip'])}"))["products"][0]["price"] == "100.00"
            mapped_preview = ok(client.get(
                f"/api/v1/catalogues/{catalogue_id}/preview",
                params={"audience_type_id": by_code["vip"]["audience_type_id"]},
            ))
            assert mapped_preview["price_list"]["name"] == "Normal"
            assert mapped_preview["products"][0]["price"] == "100.00"
            assert mapped_preview["products"][0]["wholesale_price"] == "90.00"
            assert mapped_preview["products"][0]["online_price"] == "120.00"
            assert mapped_preview["products"][0]["retail_price"] == "110.00"
            brand_mappings = ok(client.get(
                "/api/v1/my-catalogue-price-mappings",
                params={"brand": "TestBrand"},
            ))
            assert all(item["brand"] == "TestBrand" for item in brand_mappings)
            vip_price_list_id = lists["VIP"].id
            saved_brand_mappings = ok(client.put(
                "/api/v1/my-catalogue-price-mappings",
                json={
                    "brand": "TestBrand",
                    "mappings": [
                        {
                            "audience_type_id": item["audience_type_id"],
                            "price_list_id": vip_price_list_id
                            if item["audience_code"] == "vip"
                            else item["price_list_id"],
                        }
                        for item in brand_mappings
                    ],
                },
            ))
            assert next(item for item in saved_brand_mappings if item["audience_code"] == "vip")["is_brand_override"] is True
            brand_links = ok(client.get(f"/api/v1/catalogues/{catalogue_id}/share-links"))
            brand_vip_link = next(item for item in brand_links if item["audience_code"] == "vip")
            assert brand_vip_link["price_list_name"] == "VIP BKK"
            assert ok(client.get(f"/api/v1/public/catalogues/{token_from(brand_vip_link)}"))["products"][0]["price"] == "80.00"
            brand_preview = ok(client.get(
                f"/api/v1/catalogues/{catalogue_id}/preview",
                params={"audience_type_id": brand_vip_link["audience_type_id"]},
            ))
            assert brand_preview["products"][0]["price"] == "80.00"

            # A category spreadsheet must use the exact public-link view of
            # stock and customer pricing, and must not include another category.
            category_public = ok(client.get(
                f"/api/v1/public/catalogues/{token_from(brand_vip_link)}",
            ))
            category = category_public["categories"][0]
            category_export = client.get(
                f"/api/v1/public/catalogues/{token_from(brand_vip_link)}/categories/{category['slug']}/excel",
            )
            assert category_export.status_code == 200, category_export.text
            assert category_export.headers["content-type"].startswith(
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
            assert ".xlsx" in category_export.headers["content-disposition"]
            workbook = load_workbook(BytesIO(category_export.content))
            worksheet = workbook.active
            headings = [cell.value for cell in worksheet[1]]
            assert headings == [
                "Image", "Product code", "Product name", "Brand", "Category",
                "Model", "Barcode", "Description", "Stock", "Price", "Currency",
            ]
            assert worksheet.max_row == 2
            assert worksheet.cell(2, 2).value == product.sku
            assert worksheet.cell(2, 5).value == category["name"]
            assert worksheet.cell(2, 9).value == category_public["products"][0]["stock_quantity"]
            assert worksheet.cell(2, 10).value == 80
            assert worksheet.cell(2, 9).data_type == "n"
            assert worksheet.cell(2, 10).data_type == "n"
            missing_category = client.get(
                f"/api/v1/public/catalogues/{token_from(brand_vip_link)}/categories/not-a-real-category/excel",
            )
            assert missing_category.status_code == 404

            restricted_link = siam_link
            ok(client.patch(
                f"/api/v1/catalogues/{catalogue_id}/share-links/{restricted_link['id']}",
                json={"allow_pdf_download": False},
            ))
            restricted_export = client.get(
                f"/api/v1/public/catalogues/{token_from(restricted_link)}/categories/{category['slug']}/excel",
            )
            assert restricted_export.status_code == 403
            ok(client.patch(
                f"/api/v1/catalogues/{catalogue_id}/share-links/{restricted_link['id']}",
                json={"allow_pdf_download": True},
            ))
            mapped_products = ok(client.get(
                "/api/catalogue/products",
                params={
                    "audience_type_id": by_code["vip"]["audience_type_id"],
                    "q": product.sku,
                },
            ))
            assert mapped_products["items"][0]["price"] == "100.00", mapped_products["items"][0]
            assert mapped_products["items"][0]["price_list_name"] == "VIP BKK"
            assert mapped_products["items"][0]["price_level_code"] is None
            mapped_pdf = client.get(
                f"/api/v1/catalogues/{catalogue_id}/export/pdf",
                params={"audience_type_id": by_code["vip"]["audience_type_id"]},
            )
            assert mapped_pdf.status_code == 200, mapped_pdf.text
            mapped_pdf_text = "\n".join(
                page.extract_text() or ""
                for page in PdfReader(BytesIO(mapped_pdf.content)).pages
            )
            assert "80.00" in mapped_pdf_text
            assert "VIP BKK" not in mapped_pdf_text
            assert "Normal" not in mapped_pdf_text

            public_pdf = client.get(f"/api/v1/public/catalogues/{no_price_token}/pdf")
            assert public_pdf.status_code == 200, public_pdf.text
            pdf = PdfReader(BytesIO(public_pdf.content))
            assert len(pdf.pages) >= 2
            pdf_text = "\n".join(page.extract_text() or "" for page in pdf.pages)
            assert "No Price" not in pdf_text
            assert "no_price" not in public_pdf.headers.get("content-disposition", "")
            cover_images = list(pdf.pages[0].images)
            assert cover_images, "The downloaded PDF must begin with the rendered cover page."
            rendered_cover = Image.open(BytesIO(cover_images[0].data)).convert("RGB")
            red, green, blue = rendered_cover.getpixel((rendered_cover.width // 2, rendered_cover.height // 2))
            assert red > green * 2 and red > blue * 2, (red, green, blue)

            vip = by_code["vip"]
            vip_token = token_from(vip)
            protected = ok(client.patch(f"/api/v1/catalogues/{catalogue_id}/share-links/{vip['id']}", json={"password": "CustomerPass123!"}))
            assert protected["has_password"] is True
            assert client.get(f"/api/v1/public/catalogues/{vip_token}").status_code == 401
            assert client.get(f"/api/v1/public/catalogues/{vip_token}", headers={"X-Catalogue-Password": "wrong-password"}).status_code == 401
            ok(client.get(f"/api/v1/public/catalogues/{vip_token}", headers={"X-Catalogue-Password": "CustomerPass123!"}))

            fixed = by_code["vvip"]
            fixed_token = token_from(fixed)
            ok(client.patch(f"/api/v1/catalogues/{catalogue_id}/share-links/{fixed['id']}", json={"version_mode": "fixed_published", "fixed_version_number": 1}))
            ok(client.put(f"/api/v1/catalogues/{catalogue_id}/products", json={"products": [{"product_id": product_id, "section_title": "Products", "override_description": "Version two description", "hide_price": False}]}))
            ok(client.post(f"/api/v1/catalogues/{catalogue_id}/publish"))
            assert ok(client.get(f"/api/v1/public/catalogues/{fixed_token}"))["version"] == 1
            assert ok(client.get(f"/api/v1/public/catalogues/{token_from(by_code['normal'])}"))["version"] == 2

            normal = by_code["normal"]
            normal_token = token_from(normal)
            ok(client.post(f"/api/v1/catalogues/{catalogue_id}/share-links/{normal['id']}/revoke"))
            assert client.get(f"/api/v1/public/catalogues/{normal_token}").status_code == 410
            regenerated = ok(client.post(f"/api/v1/catalogues/{catalogue_id}/share-links/{normal['id']}/regenerate"))
            assert token_from(regenerated) != normal_token
            assert client.get(f"/api/v1/public/catalogues/{normal_token}").status_code == 404
            ok(client.get(f"/api/v1/public/catalogues/{token_from(regenerated)}"))

            expired_at = (datetime.now(UTC) - timedelta(minutes=5)).isoformat()
            ok(client.patch(f"/api/v1/catalogues/{catalogue_id}/share-links/{normal['id']}", json={"expires_at": expired_at}))
            assert client.get(f"/api/v1/public/catalogues/{token_from(regenerated)}").status_code == 410

            ok(client.post("/api/auth/login", json={"identifier": "share.second", "password": "StrongPassword456!", "remember_me": False}))
            sales_catalogues = ok(client.get("/api/v1/catalogues"))
            assert sales_catalogues and all(item["status"] == "published" for item in sales_catalogues)
            assert catalogue_id in {item["id"] for item in sales_catalogues}
            assert draft_catalogue_id not in {item["id"] for item in sales_catalogues}
            assert client.get(f"/api/v1/catalogues/{draft_catalogue_id}").status_code == 404
            assert client.get(f"/api/v1/catalogues/{draft_catalogue_id}/preview").status_code == 404
            assert ok(client.get("/api/v1/catalogue-share-links/cards", params={"catalogue_ids": draft_catalogue_id})) == {}
            assert ok(client.get(f"/api/v1/catalogues/{catalogue_id}/share-links")) == []
            second_vip = ok(client.post(
                f"/api/v1/catalogues/{catalogue_id}/share-links",
                json={"audience_type_id": vip["audience_type_id"]},
            ), 201)
            assert second_vip["price_list_name"] == "VIP BKK"
            assert ok(client.get(f"/api/v1/public/catalogues/{token_from(second_vip)}"))["products"][0]["price"] == "80.00"

            with SessionLocal() as db:
                stored = db.scalar(select(CatalogueShareLink).where(CatalogueShareLink.id == uuid.UUID(normal["id"])))
                assert stored and len(stored.token_hash) == 64 and regenerated["public_url"] not in stored.encrypted_token
                audits = list(db.scalars(select(AuditLog).where(AuditLog.module == "catalogue_share_links")))
                assert audits
                raw_tokens = {token_from(item) for item in links}
                assert all(not any(token in str(audit.details) for token in raw_tokens) for audit in audits)

            unauthenticated = TestClient(app)
            assert unauthenticated.get(f"/api/v1/catalogues/{catalogue_id}/share-links").status_code == 401
            assert unauthenticated.get(f"/api/v1/public/catalogues/{no_price_token}").status_code == 200
            unauthenticated.close()
        print("Secure catalogue share-link smoke test passed.")
    finally:
        shutil.rmtree(test_directory, ignore_errors=True)


if __name__ == "__main__":
    main()
