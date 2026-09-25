"""End-to-end pricing and catalogue-version smoke test.

Run from ``backend`` with:
    python -m tests.smoke_pricing_catalogues
"""

import os
import shutil
import tempfile
import uuid
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
from zipfile import is_zipfile


test_directory = Path(tempfile.mkdtemp(prefix="gms-commerce-test-"))
os.environ["DATABASE_URL"] = (
    f"sqlite+pysqlite:///{(test_directory / 'commerce.db').as_posix()}"
)
os.environ["SECRET_KEY"] = "commerce-smoke-test-secret"
os.environ["AUTO_CREATE_TABLES"] = "true"
os.environ["SEED_DEMO_DATA"] = "true"
os.environ["UPLOAD_DIR"] = str(test_directory / "uploads")

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import event, select, update

from app.commerce_models import CatalogueVersion, PriceChangeRequest, ProductPrice
from app.database import SessionLocal, engine
from app.main import app
from app.models import AuditLog, Permission, Product, Role, User, role_permissions
from app.security import hash_password


def assert_a4_landscape(content: bytes) -> None:
    reader = PdfReader(BytesIO(content))
    assert reader.pages
    dimensions = []
    for page in reader.pages:
        width = float(page.mediabox.width)
        height = float(page.mediabox.height)
        assert width > height, (width, height)
        assert abs(width - 841.89) < 1.0
        assert abs(height - 595.28) < 1.0
        dimensions.append((round(width, 2), round(height, 2)))
    assert len(set(dimensions)) == 1


@event.listens_for(engine, "connect")
def enable_sqlite_foreign_keys(dbapi_connection, _) -> None:
    cursor = dbapi_connection.cursor()
    cursor.execute("PRAGMA foreign_keys=ON")
    cursor.close()


def assert_ok(response, expected_status: int = 200) -> dict:
    assert response.status_code == expected_status, response.text
    return response.json()


def price_payload(
    product_id: str,
    price_list_id: int,
    amount: str,
    effective_from: str,
    reason: str,
) -> dict:
    return {
        "product_id": product_id,
        "price_list_id": price_list_id,
        "proposed_amount": amount,
        "effective_from": effective_from,
        "reason": reason,
    }


def catalogue_payload(
    *,
    title: str,
    slug: str,
    price_list_id: int,
    show_prices: bool = True,
) -> dict:
    return {
        "title": title,
        "slug": slug,
        "description": f"Smoke-test catalogue for {title}.",
        "brand": None,
        "audience": "Internal",
        "price_list_id": price_list_id,
        "show_prices": show_prices,
        "currency": "THB",
        "language": "en",
        "valid_from": None,
        "valid_until": None,
        "is_public": False,
    }


def main() -> None:
    try:
        with TestClient(app) as client:
            assert_ok(client.get("/api/health"))

            with SessionLocal() as db:
                superadmin_role = db.scalar(
                    select(Role).where(Role.name == "superadmin")
                )
                system_user_role = db.scalar(
                    select(Role).where(Role.name == "system_user")
                )
                assert superadmin_role is not None
                assert system_user_role is not None
                preview_permissions = list(
                    db.scalars(
                        select(Permission).where(
                            Permission.code.in_(
                                {
                                    "catalogues.preview",
                                    "catalogues.export_pdf",
                                    "catalogues.print",
                                }
                            )
                        )
                    )
                )
                assert len(preview_permissions) == 3
                restricted_role = Role(
                    name="preview_without_prices",
                    description="Can preview and export without price access.",
                    permissions=preview_permissions,
                )
                db.add(restricted_role)
                db.flush()
                db.execute(
                    update(role_permissions)
                    .where(role_permissions.c.role_id == restricted_role.id)
                    .values(access_scope="all", effect="allow")
                )
                db.add(
                    User(
                        username="commerce.admin",
                        email="commerce.admin@example.com",
                        full_name="Commerce Administrator",
                        password_hash=hash_password("StrongPassword123!"),
                        roles=[superadmin_role],
                    )
                )
                db.add_all(
                    [
                        User(
                            username="preview.denied",
                            email="preview.denied@example.com",
                            full_name="Preview Denied",
                            password_hash=hash_password("DeniedPassword123!"),
                            roles=[system_user_role],
                        ),
                        User(
                            username="preview.restricted",
                            email="preview.restricted@example.com",
                            full_name="Price Restricted Preview",
                            password_hash=hash_password("RestrictedPassword123!"),
                            roles=[restricted_role],
                        ),
                    ]
                )
                db.commit()

            assert_ok(
                client.post(
                    "/api/auth/login",
                    json={
                        "identifier": "commerce.admin",
                        "password": "StrongPassword123!",
                        "remember_me": False,
                    },
                )
            )

            generation_result = {
                "created": 3,
                "updated": 1,
                "brand_count": 4,
                "product_count": 120,
                "products_with_images": 100,
                "products_missing_images": 20,
                "category_setting_count": 12,
                "customer_level_count": 7,
                "legacy_category_rows": 95,
                "legacy_matched_brands": 3,
                "cover_count": 4,
                "covers_with_erp_image": 3,
                "covers_with_fallback_design": 1,
                "warnings": [],
            }
            with patch(
                "app.commerce.generate_erp_brand_catalogues",
                return_value=generation_result,
            ):
                generated = assert_ok(
                    client.post("/api/v1/catalogues/generate-erp-brands")
                )
            assert generated == generation_result

            # Seeded price lists are database records, and additional lists can be
            # created without changing a schema or a Python enum.
            seeded_lists = assert_ok(client.get("/api/v1/price-lists"))
            by_code = {item["code"]: item for item in seeded_lists}
            assert {
                "NORMAL",
                "VIP",
                "BIG_CUSTOMER",
                "DEALER",
                "WHOLESALE",
                "RETAIL",
                "NO_PRICE",
            }.issubset(by_code)
            assert by_code["NO_PRICE"]["is_no_price"] is True

            corporate_payload = {
                "code": "corporate_2027",
                "name": "Corporate 2027",
                "description": "A configurable smoke-test price list.",
                "currency": "thb",
                "is_no_price": False,
                "is_active": True,
            }
            corporate = assert_ok(
                client.post("/api/v1/price-lists", json=corporate_payload),
                201,
            )
            assert corporate["code"] == "CORPORATE_2027"
            assert corporate["currency"] == "THB"
            duplicate_list = client.post(
                "/api/v1/price-lists", json=corporate_payload
            )
            assert duplicate_list.status_code == 409, duplicate_list.text

            with SessionLocal() as db:
                products = list(db.scalars(select(Product).order_by(Product.sku)))
                assert len(products) >= 2
                first_product_id = str(products[0].id)
                second_product_id = str(products[1].id)

            # Direct price changes append records and close the preceding interval.
            first_price = assert_ok(
                client.post(
                    "/api/v1/product-prices",
                    json=price_payload(
                        first_product_id,
                        corporate["id"],
                        "100.00",
                        "2026-01-01T00:00:00",
                        "Initial corporate price.",
                    ),
                ),
                201,
            )
            second_price = assert_ok(
                client.post(
                    "/api/v1/product-prices",
                    json=price_payload(
                        first_product_id,
                        corporate["id"],
                        "120.00",
                        "2026-02-01T00:00:00",
                        "Quarterly corporate adjustment.",
                    ),
                ),
                201,
            )
            assert first_price["id"] != second_price["id"]
            price_history = assert_ok(
                client.get(
                    "/api/v1/product-prices",
                    params={
                        "product_id": first_product_id,
                        "price_list_id": corporate["id"],
                    },
                )
            )
            history_by_id = {item["id"]: item for item in price_history}
            assert set(history_by_id) == {first_price["id"], second_price["id"]}
            assert history_by_id[first_price["id"]]["amount"] == "100.00"
            assert history_by_id[first_price["id"]]["reason"] == (
                "Initial corporate price."
            )
            assert history_by_id[first_price["id"]]["expires_at"].startswith(
                "2026-02-01T00:00:00"
            )
            assert history_by_id[second_price["id"]]["amount"] == "120.00"
            assert history_by_id[second_price["id"]]["expires_at"] is None

            no_price_rejection = client.post(
                "/api/v1/product-prices",
                json=price_payload(
                    first_product_id,
                    by_code["NO_PRICE"]["id"],
                    "1.00",
                    "2026-02-15T00:00:00",
                    "This must be rejected.",
                ),
            )
            assert no_price_rejection.status_code == 422, no_price_rejection.text

            # Proposal approval produces another price record and preserves both
            # earlier prices as history.
            proposal = assert_ok(
                client.post(
                    "/api/v1/price-change-requests",
                    json=price_payload(
                        first_product_id,
                        corporate["id"],
                        "130.00",
                        "2026-03-01T00:00:00",
                        "Approved annual contract increase.",
                    ),
                ),
                201,
            )
            assert proposal["status"] == "pending"
            assert proposal["current_amount"] == "120.00"
            approved = assert_ok(
                client.post(
                    f"/api/v1/price-change-requests/{proposal['id']}/review",
                    json={
                        "approve": True,
                        "reason": "Reviewed against the signed contract.",
                    },
                )
            )
            assert approved["status"] == "approved"
            assert approved["resulting_price_id"]
            assert approved["current_amount"] == "130.00"
            already_reviewed = client.post(
                f"/api/v1/price-change-requests/{proposal['id']}/review",
                json={"approve": False, "reason": "A second review is invalid."},
            )
            assert already_reviewed.status_code == 409, already_reviewed.text

            with SessionLocal() as db:
                stored_request = db.get(
                    PriceChangeRequest, uuid.UUID(proposal["id"])
                )
                assert stored_request is not None
                assert stored_request.status == "approved"
                stored_prices = list(
                    db.scalars(
                        select(ProductPrice)
                        .where(
                            ProductPrice.product_id == uuid.UUID(first_product_id),
                            ProductPrice.price_list_id == corporate["id"],
                        )
                        .order_by(ProductPrice.effective_from)
                    )
                )
                assert [str(item.amount) for item in stored_prices] == [
                    "100.00",
                    "120.00",
                    "130.00",
                ]
                assert [item.reason for item in stored_prices] == [
                    "Initial corporate price.",
                    "Quarterly corporate adjustment.",
                    "Approved annual contract increase.",
                ]

            # Catalogue membership preserves order and catalogue-specific content.
            priced_catalogue = assert_ok(
                client.post(
                    "/api/v1/catalogues",
                    json=catalogue_payload(
                        title="Corporate Product Catalogue",
                        slug="corporate-product-catalogue",
                        price_list_id=corporate["id"],
                    ),
                ),
                201,
            )
            priced_catalogue_id = priced_catalogue["id"]
            membership = {
                "products": [
                    {
                        "product_id": first_product_id,
                        "section_title": "Featured",
                        "override_description": "Version one description.",
                        "hide_price": False,
                    },
                    {
                        "product_id": second_product_id,
                        "section_title": "Reference",
                        "override_description": "Price intentionally hidden.",
                        "hide_price": True,
                    },
                ]
            }
            with_products = assert_ok(
                client.put(
                    f"/api/v1/catalogues/{priced_catalogue_id}/products",
                    json=membership,
                )
            )
            assert with_products["product_count"] == 2
            assert [item["sort_order"] for item in with_products["products"]] == [
                1,
                2,
            ]
            assert [item["product_id"] for item in with_products["products"]] == [
                first_product_id,
                second_product_id,
            ]

            reordered_membership = {
                "products": [
                    {
                        "product_id": second_product_id,
                        "section_title": "Reference Reordered",
                        "override_description": "Hidden product moved first.",
                        "hide_price": True,
                    },
                    {
                        "product_id": first_product_id,
                        "section_title": "Featured Reordered",
                        "override_description": "Priced product moved second.",
                        "hide_price": False,
                    },
                ]
            }
            reordered = assert_ok(
                client.put(
                    f"/api/v1/catalogues/{priced_catalogue_id}/products",
                    json=reordered_membership,
                )
            )
            assert [item["product_id"] for item in reordered["products"]] == [
                second_product_id,
                first_product_id,
            ]
            assert [item["sort_order"] for item in reordered["products"]] == [1, 2]
            assert reordered["products"][0]["section_title"] == (
                "Reference Reordered"
            )
            assert reordered["products"][1]["override_description"] == (
                "Priced product moved second."
            )
            persisted_reorder = assert_ok(
                client.get(f"/api/v1/catalogues/{priced_catalogue_id}")
            )
            assert [
                (item["product_id"], item["sort_order"], item["section_title"])
                for item in persisted_reorder["products"]
            ] == [
                (second_product_id, 1, "Reference Reordered"),
                (first_product_id, 2, "Featured Reordered"),
            ]

            preview_v1 = assert_ok(
                client.get(f"/api/v1/catalogues/{priced_catalogue_id}/preview")
            )
            assert preview_v1["is_draft"] is True
            assert preview_v1.get("version") is None
            assert preview_v1["status"] == "draft"
            assert preview_v1["product_count"] == 2
            assert [item["display_order"] for item in preview_v1["products"]] == [1, 2]
            assert [item["code"] for item in preview_v1["products"]] == [
                reordered["products"][0]["sku"],
                reordered["products"][1]["sku"],
            ]
            assert preview_v1["show_prices"] is True
            assert "price" not in preview_v1["products"][0]
            assert preview_v1["products"][1]["price"] == "130.00"
            assert preview_v1["products"][1]["currency"] == "THB"
            assert preview_v1["price_list"]["name"] == "Corporate 2027"
            assert any(
                item.get("main_image_url") is None for item in preview_v1["products"]
            )

            invalid_catalogue = client.get(
                f"/api/v1/catalogues/{uuid.uuid4()}/preview"
            )
            assert invalid_catalogue.status_code == 404, invalid_catalogue.text
            invalid_version = client.get(
                f"/api/v1/catalogues/{priced_catalogue_id}/preview",
                params={"version": 999},
            )
            assert invalid_version.status_code == 404, invalid_version.text
            assert invalid_version.json()["detail"] == (
                "The selected catalogue version is unavailable."
            )

            with TestClient(app) as denied_client:
                assert_ok(
                    denied_client.post(
                        "/api/auth/login",
                        json={
                            "identifier": "preview.denied",
                            "password": "DeniedPassword123!",
                            "remember_me": False,
                        },
                    )
                )
                denied_preview = denied_client.get(
                    f"/api/v1/catalogues/{priced_catalogue_id}/preview"
                )
                assert denied_preview.status_code == 403, denied_preview.text
                denied_pdf = denied_client.get(
                    f"/api/v1/catalogues/{priced_catalogue_id}/export/pdf"
                )
                assert denied_pdf.status_code == 403, denied_pdf.text

            with TestClient(app) as restricted_client:
                assert_ok(
                    restricted_client.post(
                        "/api/auth/login",
                        json={
                            "identifier": "preview.restricted",
                            "password": "RestrictedPassword123!",
                            "remember_me": False,
                        },
                    )
                )
                restricted_preview = assert_ok(
                    restricted_client.get(
                        f"/api/v1/catalogues/{priced_catalogue_id}/preview"
                    )
                )
                assert restricted_preview["show_prices"] is False
                assert "currency" not in restricted_preview
                assert "price_list" not in restricted_preview
                assert not any(
                    {"price", "currency"}.intersection(item)
                    for item in restricted_preview["products"]
                )
                restricted_pdf = restricted_client.get(
                    f"/api/v1/catalogues/{priced_catalogue_id}/export/pdf"
                )
                assert restricted_pdf.status_code == 200, restricted_pdf.text
                assert restricted_pdf.content.startswith(b"%PDF")
                assert_a4_landscape(restricted_pdf.content)
                printed = restricted_client.post(
                    f"/api/v1/catalogues/{priced_catalogue_id}/print"
                )
                assert printed.status_code == 204, printed.text

            published_v1 = assert_ok(
                client.post(f"/api/v1/catalogues/{priced_catalogue_id}/publish")
            )
            assert published_v1["version_number"] == 1
            immutable_v1 = published_v1["snapshot"]
            assert immutable_v1["products"][1]["price"] == "130.00"
            latest_published_preview = assert_ok(
                client.get(f"/api/v1/catalogues/{priced_catalogue_id}/preview")
            )
            assert latest_published_preview["version"] == 1
            assert latest_published_preview["is_draft"] is False

            # Change live pricing and catalogue metadata after publication.
            # Version one must continue returning its original self-contained
            # snapshot.
            assert_ok(
                client.post(
                    "/api/v1/product-prices",
                    json=price_payload(
                        first_product_id,
                        corporate["id"],
                        "140.00",
                        "2026-04-01T00:00:00",
                        "Post-publication price change.",
                    ),
                ),
                201,
            )
            published_catalogue = assert_ok(
                client.get(f"/api/v1/catalogues/{priced_catalogue_id}")
            )
            version_two_update = catalogue_payload(
                title="Corporate Product Catalogue Version Two",
                slug="corporate-product-catalogue",
                price_list_id=corporate["id"],
            )
            version_two_update.update(
                {
                    "status": "draft",
                    "expected_revision": published_catalogue["revision"],
                }
            )
            draft_again = assert_ok(
                client.put(
                    f"/api/v1/catalogues/{priced_catalogue_id}",
                    json=version_two_update,
                )
            )
            assert draft_again["status"] == "draft"
            draft_version_two_preview = assert_ok(
                client.get(f"/api/v1/catalogues/{priced_catalogue_id}/preview")
            )
            assert draft_version_two_preview["is_draft"] is True
            assert draft_version_two_preview.get("version") is None
            assert draft_version_two_preview["title"] == (
                "Corporate Product Catalogue Version Two"
            )
            historical_preview = assert_ok(
                client.get(
                    f"/api/v1/catalogues/{priced_catalogue_id}/preview",
                    params={"version": 1},
                )
            )
            assert historical_preview["version"] == 1
            assert historical_preview["title"] == immutable_v1["title"]
            assert historical_preview["products"][1]["price"] == "130.00"
            persisted_v1 = assert_ok(
                client.get(
                    f"/api/v1/catalogues/{priced_catalogue_id}/versions/1"
                )
            )
            assert persisted_v1["snapshot"] == immutable_v1

            published_v2 = assert_ok(
                client.post(f"/api/v1/catalogues/{priced_catalogue_id}/publish")
            )
            assert published_v2["version_number"] == 2
            assert published_v2["snapshot"]["title"] == (
                "Corporate Product Catalogue Version Two"
            )
            assert published_v2["snapshot"]["products"][1]["price"] == "140.00"
            versions = assert_ok(
                client.get(f"/api/v1/catalogues/{priced_catalogue_id}/versions")
            )
            assert [item["version_number"] for item in versions] == [2, 1]
            assert versions[1]["snapshot"] == immutable_v1

            priced_excel = client.get(
                f"/api/v1/catalogues/{priced_catalogue_id}/export/excel",
                params={"version_number": 2},
            )
            assert priced_excel.status_code == 200, priced_excel.text
            assert priced_excel.content.startswith(b"PK")
            assert is_zipfile(BytesIO(priced_excel.content))
            priced_workbook = load_workbook(
                BytesIO(priced_excel.content), read_only=True, data_only=True
            )
            priced_sheet = priced_workbook.active
            priced_headers = [cell.value for cell in priced_sheet[1]]
            assert priced_headers[-2:] == ["Price", "Currency"]
            assert any(
                str(row[6].value) == "140.00" and row[7].value == "THB"
                for row in priced_sheet.iter_rows(min_row=2)
                if row[6].value is not None
            )
            priced_workbook.close()
            priced_pdf = client.get(
                f"/api/v1/catalogues/{priced_catalogue_id}/export/pdf",
                params={"version_number": 2},
            )
            assert priced_pdf.status_code == 200, priced_pdf.text
            assert priced_pdf.headers["content-type"].startswith("application/pdf")
            assert priced_pdf.content.startswith(b"%PDF")
            assert_a4_landscape(priced_pdf.content)
            assert "CORPORATE-PRODUCT-CATALOGUE-VERSION-TWO-v2.pdf" in (
                priced_pdf.headers["content-disposition"]
            ), priced_pdf.headers["content-disposition"]

            with SessionLocal() as db:
                stored_versions = list(
                    db.scalars(
                        select(CatalogueVersion)
                        .where(
                            CatalogueVersion.catalogue_id
                            == uuid.UUID(priced_catalogue_id)
                        )
                        .order_by(CatalogueVersion.version_number)
                    )
                )
                assert len(stored_versions) == 2
                assert stored_versions[0].snapshot == immutable_v1
                audit_actions = set(
                    db.scalars(
                        select(AuditLog.action).where(
                            AuditLog.action.in_(
                                {
                                    "catalogue_preview_opened",
                                    "catalogue_version_selected",
                                    "catalogue_exported_pdf",
                                    "catalogue_printed",
                                }
                            )
                        )
                    )
                )
                assert audit_actions == {
                    "catalogue_preview_opened",
                    "catalogue_version_selected",
                    "catalogue_exported_pdf",
                    "catalogue_printed",
                }

            # A No Price list forces price display off in every backend response
            # and snapshot, even when a client asks to show prices.
            no_price_catalogue = assert_ok(
                client.post(
                    "/api/v1/catalogues",
                    json=catalogue_payload(
                        title="No Price Product Catalogue",
                        slug="no-price-product-catalogue",
                        price_list_id=by_code["NO_PRICE"]["id"],
                        show_prices=True,
                    ),
                ),
                201,
            )
            assert no_price_catalogue["show_prices"] is False
            no_price_id = no_price_catalogue["id"]
            assert_ok(
                client.put(
                    f"/api/v1/catalogues/{no_price_id}/products",
                    json={
                        "products": [
                            {
                                "product_id": first_product_id,
                                "section_title": "Products",
                                "override_description": "No monetary data.",
                                "hide_price": False,
                            }
                        ]
                    },
                )
            )
            no_price_preview = assert_ok(
                client.get(f"/api/v1/catalogues/{no_price_id}/preview")
            )
            assert no_price_preview["show_prices"] is False
            assert "currency" not in no_price_preview
            assert not {
                "price",
                "currency",
                "price_missing",
            }.intersection(no_price_preview["products"][0])
            no_price_version = assert_ok(
                client.post(f"/api/v1/catalogues/{no_price_id}/publish")
            )
            assert no_price_version["snapshot"]["show_prices"] is False
            assert "currency" not in no_price_version["snapshot"]
            assert "price_list" not in no_price_version["snapshot"]
            assert not {
                "price",
                "currency",
                "price_missing",
            }.intersection(no_price_version["snapshot"]["products"][0])
            persisted_no_price = assert_ok(
                client.get(f"/api/v1/catalogues/{no_price_id}/versions/1")
            )
            assert persisted_no_price["snapshot"] == no_price_version["snapshot"]

            no_price_excel = client.get(
                f"/api/v1/catalogues/{no_price_id}/export/excel",
                params={"version_number": 1},
            )
            assert no_price_excel.status_code == 200, no_price_excel.text
            assert no_price_excel.content.startswith(b"PK")
            assert is_zipfile(BytesIO(no_price_excel.content))
            no_price_workbook = load_workbook(
                BytesIO(no_price_excel.content), read_only=True, data_only=True
            )
            no_price_headers = [cell.value for cell in no_price_workbook.active[1]]
            assert "Price" not in no_price_headers
            assert "Currency" not in no_price_headers
            assert no_price_headers == [
                "Order",
                "SKU",
                "Product",
                "Brand",
                "Category",
                "Description",
            ]
            no_price_workbook.close()
            no_price_pdf = client.get(
                f"/api/v1/catalogues/{no_price_id}/export/pdf",
                params={"version_number": 1},
            )
            assert no_price_pdf.status_code == 200, no_price_pdf.text
            assert no_price_pdf.content.startswith(b"%PDF")
            assert b"130.00" not in no_price_pdf.content
            assert b"140.00" not in no_price_pdf.content
            assert_a4_landscape(no_price_pdf.content)

            # Catalogue updates use an expected revision so stale editors cannot
            # silently overwrite a newer change.
            revision_catalogue = assert_ok(
                client.post(
                    "/api/v1/catalogues",
                    json=catalogue_payload(
                        title="Revision Check Catalogue",
                        slug="revision-check-catalogue",
                        price_list_id=corporate["id"],
                    ),
                ),
                201,
            )
            revision_update = catalogue_payload(
                title="Revision Check Catalogue Updated",
                slug="revision-check-catalogue",
                price_list_id=corporate["id"],
            )
            revision_update.update(
                {"status": "draft", "expected_revision": revision_catalogue["revision"]}
            )
            updated_revision = assert_ok(
                client.put(
                    f"/api/v1/catalogues/{revision_catalogue['id']}",
                    json=revision_update,
                )
            )
            assert updated_revision["revision"] == revision_catalogue["revision"] + 1
            stale_update = client.put(
                f"/api/v1/catalogues/{revision_catalogue['id']}",
                json=revision_update,
            )
            assert stale_update.status_code == 409, stale_update.text

            duplicate = assert_ok(
                client.post(
                    f"/api/v1/catalogues/{priced_catalogue_id}/duplicate"
                ),
                201,
            )
            assert duplicate["id"] != priced_catalogue_id
            assert duplicate["status"] == "draft"
            assert duplicate["version"] == 0
            assert duplicate["product_count"] == 2
            archived = assert_ok(
                client.post(f"/api/v1/catalogues/{duplicate['id']}/archive")
            )
            assert archived["status"] == "archived"
            source_after_archive = assert_ok(
                client.get(f"/api/v1/catalogues/{priced_catalogue_id}")
            )
            assert source_after_archive["status"] == "published"
            deleted = client.delete(f"/api/v1/catalogues/{duplicate['id']}")
            assert deleted.status_code == 204, deleted.text
            deleted_lookup = client.get(
                f"/api/v1/catalogues/{duplicate['id']}"
            )
            assert deleted_lookup.status_code == 404, deleted_lookup.text

            with SessionLocal() as db:
                commerce_audits = list(
                    db.scalars(
                        select(AuditLog).where(
                            AuditLog.module.in_(["pricing", "catalogues"])
                        )
                    )
                )
                audit_actions = {item.action for item in commerce_audits}
                assert {
                    "price_list_created",
                    "product_price_created",
                    "price_change_proposed",
                    "price_change_approved",
                    "catalogue_created",
                    "catalogue_products_updated",
                    "catalogue_updated",
                    "catalogue_published",
                    "catalogue_exported_excel",
                    "catalogue_exported_pdf",
                    "catalogue_duplicated",
                    "catalogue_archived",
                    "catalogue_deleted",
                }.issubset(audit_actions)
                proposal_audit = next(
                    item
                    for item in commerce_audits
                    if item.action == "price_change_proposed"
                )
                assert proposal_audit.details["reason"] == (
                    "Approved annual contract increase."
                )
                assert all(item.status == "success" for item in commerce_audits)

        print("Pricing and catalogue-version smoke test: OK")
    finally:
        shutil.rmtree(test_directory, ignore_errors=True)


if __name__ == "__main__":
    main()
