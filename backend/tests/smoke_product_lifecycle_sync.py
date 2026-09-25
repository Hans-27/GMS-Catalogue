"""Product lifecycle and source-sync acceptance smoke test.

Run from ``backend`` with:
    python -m tests.smoke_product_lifecycle_sync
"""

import io
import hashlib
import os
import shutil
import tempfile
import uuid
from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path

from PIL import Image


test_directory = Path(tempfile.mkdtemp(prefix="gms-lifecycle-sync-test-"))
os.environ["DATABASE_URL"] = f"sqlite+pysqlite:///{(test_directory / 'catalogue.db').as_posix()}"
os.environ["SECRET_KEY"] = "lifecycle-sync-smoke-secret"
os.environ["AUTO_CREATE_TABLES"] = "true"
os.environ["SEED_DEMO_DATA"] = "true"
os.environ["UPLOAD_DIR"] = str(test_directory / "uploads")
os.environ["PRODUCT_SYNC_INTERVAL_SECONDS"] = "180"
os.environ["PRODUCT_SYNC_MAX_RETRIES"] = "0"

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import func, select

from app.commerce_models import Catalogue, CatalogueProduct, PriceList, ProductPrice
from app.database import SessionLocal
from app.erp_models import ErpConnectionSetting, ErpSyncRun
from app.main import app
from app.models import AuditLog, Product, ProductImage, ProductStatusHistory, ProductWarehouseStock, Role, User
from app.security import hash_password
from app import product_sync_service


def login(client: TestClient, identifier: str, password: str) -> None:
    response = client.post("/api/auth/login", json={"identifier": identifier, "password": password, "remember_me": False})
    assert response.status_code == 200, response.text


def catalogue_payload(price_list_id: int) -> dict:
    return {
        "title": "Lifecycle Public Catalogue", "slug": "lifecycle-public-catalogue",
        "description": "Lifecycle visibility test.", "brand": None, "audience": "Public",
        "price_list_id": price_list_id, "show_prices": True, "currency": "THB",
        "language": "en", "valid_from": None, "valid_until": None, "is_public": True,
    }


def main() -> None:
    original_batches = product_sync_service._source_batches
    original_image_metadata = product_sync_service._source_image_metadata
    original_image_content = product_sync_service._source_image_content
    original_brand_names = product_sync_service._erp_brand_names
    original_price_level_rows = product_sync_service._erp_price_level_rows
    try:
        with TestClient(app) as admin:
            with SessionLocal() as db:
                super_role = db.scalar(select(Role).where(Role.name == "superadmin"))
                sales_role = db.scalar(select(Role).where(Role.name == "sales_user"))
                assert super_role and sales_role
                db.add_all([
                    User(username="lifecycle.admin", email="lifecycle.admin@example.com", full_name="Lifecycle Administrator", password_hash=hash_password("StrongPassword123!"), roles=[super_role]),
                    User(username="lifecycle.sales", email="lifecycle.sales@example.com", full_name="Lifecycle Sales", password_hash=hash_password("StrongPassword123!"), roles=[sales_role]),
                    ErpConnectionSetting(id=1, server="source.invalid", port=1433, database_name="ERP", username="read_only", password_ciphertext="test-only", is_enabled=True),
                ])
                db.commit()

            login(admin, "lifecycle.admin", "StrongPassword123!")
            all_products = admin.get("/api/v1/products", params={"status": "all"})
            assert all_products.status_code == 200, all_products.text
            assert all(item["product_status"] == "active" for item in all_products.json()["items"])
            product_json = next(item for item in all_products.json()["items"] if item["workflow_status"] == "published")
            product_id, sku = product_json["id"], product_json["sku"]
            product_uuid = uuid.UUID(product_id)

            price_lists = admin.get("/api/v1/price-lists").json()
            normal = next(item for item in price_lists if item["code"] == "NORMAL")
            created = admin.post("/api/v1/catalogues", json=catalogue_payload(normal["id"]))
            assert created.status_code == 201, created.text
            catalogue_id = created.json()["id"]
            membership = admin.put(f"/api/v1/catalogues/{catalogue_id}/products", json={"products": [{"product_id": product_id, "section_title": "Lifecycle", "override_description": "", "hide_price": False}]})
            assert membership.status_code == 200, membership.text
            with SessionLocal() as db:
                product = db.get(Product, product_uuid)
                product.source_system = "gms_erp"
                product.source_record_id = "1001"
                product.source_record_exists = True
                product.brand = "Lifecycle ERP Brand"
                auto_catalogue = Catalogue(
                    title="Lifecycle ERP Brand Catalogue",
                    slug="erp-brand-lifecycle-erp-brand",
                    description="Lifecycle auto-membership test.",
                    brand="Lifecycle ERP Brand",
                    audience="All customer levels",
                )
                db.add(auto_catalogue)
                db.flush()
                auto_catalogue_id = auto_catalogue.id
                db.add(CatalogueProduct(
                    catalogue_id=auto_catalogue.id,
                    product_id=product.id,
                    section_title="Lifecycle",
                    override_description="",
                    hide_price=False,
                    sort_order=1,
                ))
                db.commit()
            published = admin.post(f"/api/v1/catalogues/{catalogue_id}/publish")
            assert published.status_code == 200, published.text
            assert admin.get(f"/api/v1/public/products/{sku}").status_code == 200
            assert admin.get("/api/v1/public/catalogues/by-slug/lifecycle-public-catalogue").json()["product_count"] == 1

            invalid_reason = admin.patch(f"/api/v1/products/{product_id}/status", json={"status": "inactive", "reason": "other", "note": ""})
            assert invalid_reason.status_code == 422
            changed = admin.patch(f"/api/v1/products/{product_id}/status", headers={"x-request-id": "lifecycle-test-1"}, json={"status": "inactive", "reason": "discontinued", "note": "Supplier stopped this model."})
            assert changed.status_code == 200, changed.text
            assert changed.json()["product_status"] == "inactive"
            assert changed.json()["catalogue_assignments"][0]["id"] == catalogue_id
            with SessionLocal() as db:
                assert db.scalar(select(func.count()).select_from(CatalogueProduct).where(
                    CatalogueProduct.catalogue_id == auto_catalogue_id,
                    CatalogueProduct.product_id == product_uuid,
                )) == 0
            inactive_membership = admin.put(f"/api/v1/catalogues/{catalogue_id}/products", json={"products": [{"product_id": product_id, "section_title": "Lifecycle", "override_description": "", "hide_price": False}]})
            assert inactive_membership.status_code == 422

            inactive_list = admin.get("/api/v1/products", params={"status": "inactive"})
            assert inactive_list.status_code == 200
            assert product_id in {item["id"] for item in inactive_list.json()["items"]}
            assert admin.get(f"/api/v1/public/products/{sku}").status_code == 404
            assert admin.get("/api/v1/public/catalogues/by-slug/lifecycle-public-catalogue").json()["product_count"] == 0
            assert admin.get(f"/api/v1/catalogues/{catalogue_id}/preview").json()["product_count"] == 0
            assert admin.get(f"/api/v1/catalogues/{catalogue_id}/preview", params={"include_inactive": "true"}).json()["product_count"] == 1

            pdf = admin.get(f"/api/v1/catalogues/{catalogue_id}/export/pdf")
            assert pdf.status_code == 200
            pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(pdf.content)).pages)
            assert sku not in pdf_text
            printed = admin.post(f"/api/v1/catalogues/{catalogue_id}/print")
            assert printed.status_code == 204
            excel = admin.get(f"/api/v1/catalogues/{catalogue_id}/export/excel")
            assert excel.status_code == 200
            workbook = load_workbook(io.BytesIO(excel.content), read_only=True)
            excel_values = {str(value) for row in workbook.active.iter_rows(values_only=True) for value in row if value is not None}
            assert sku not in excel_values

            with TestClient(app) as sales:
                login(sales, "lifecycle.sales", "StrongPassword123!")
                hidden = sales.get("/api/v1/products", params={"status": "all"})
                assert hidden.status_code == 403
                # This role cannot view the product module at all, so the
                # permission boundary rejects the request before the inactive
                # record can be resolved.
                assert sales.get(f"/api/v1/products/{product_id}").status_code == 403
                assert sales.patch(f"/api/v1/products/{product_id}/status", json={"status": "active", "note": "No authority"}).status_code == 403
                assert sales.post("/api/v1/admin/data-sync/run").status_code == 403

            reactivated = admin.patch(f"/api/v1/products/{product_id}/status", json={"status": "active", "note": "Supplier resumed production."})
            assert reactivated.status_code == 200
            assert admin.get("/api/v1/public/catalogues/by-slug/lifecycle-public-catalogue").json()["product_count"] == 1
            with SessionLocal() as db:
                assert db.scalar(select(func.count()).select_from(CatalogueProduct).where(
                    CatalogueProduct.catalogue_id == auto_catalogue_id,
                    CatalogueProduct.product_id == product_uuid,
                )) == 1
            admin.patch(f"/api/v1/products/{product_id}/status", json={"status": "inactive", "reason": "temporarily_unavailable", "note": "Backend values must continue syncing."})
            with SessionLocal() as db:
                assert db.scalar(select(func.count()).select_from(CatalogueProduct).where(
                    CatalogueProduct.catalogue_id == auto_catalogue_id,
                    CatalogueProduct.product_id == product_uuid,
                )) == 0

            with SessionLocal() as db:
                product = db.scalar(select(Product).where(Product.sku == sku))
                second = db.scalar(select(Product).where(Product.id != product.id).limit(1))
                product.source_system, product.source_record_id = "gms_erp", "1001"
                second.source_system, second.source_record_id = "gms_erp", "1002"
                old_stock, old_price = product.stock_quantity, product.price
                db.commit()

            source_row = {
                "erp_id": "1001", "sku": sku, "product_name": "ERP Controlled Name",
                "brand": "ERP Brand", "category": "ERP Category", "barcode": None, "unit": "piece",
                "Price01": Decimal("199.00"), "Price02": None, "Price03": None, "Price04": None, "Price05": None,
                "is_discontinued": False, "erp_updated_at": datetime.now(UTC),
            }
            image_buffer = io.BytesIO()
            Image.new("RGB", (32, 24), color=(24, 118, 63)).save(image_buffer, format="PNG")
            source_image = image_buffer.getvalue()
            source_image_hash = hashlib.sha256(source_image).hexdigest()

            def fake_batches(_setting):
                yield [source_row], {"1001": [{"warehouse_code": "MAIN", "on_hand": Decimal("42")}]} 
            product_sync_service._source_batches = fake_batches
            product_sync_service._erp_brand_names = lambda _setting: ["ERP Brand"]
            product_sync_service._erp_price_level_rows = lambda _setting: [
                {
                    "erp_price_type_id": 19,
                    "source_name": "SP1",
                    "product_count": 1,
                }
            ]
            product_sync_service._source_image_metadata = lambda _setting, _ids: {
                "1001": {
                    "erp_id": 1001,
                    "image1_size": len(source_image),
                    "image1_hash": source_image_hash,
                    "image1_description": "ERP lifecycle image",
                }
            }
            product_sync_service._source_image_content = lambda _setting, _ids: {
                "1001": {"erp_id": 1001, "image1": source_image}
            }

            first_run_id = product_sync_service.run_product_sync(trigger="automatic", retry_backoffs=())
            with SessionLocal() as db:
                product = db.scalar(select(Product).where(Product.sku == sku))
                second = db.scalar(select(Product).where(Product.source_record_id == "1002"))
                normal_list = db.scalar(select(PriceList).where(PriceList.code == "NORMAL"))
                history_count = db.scalar(select(func.count()).select_from(ProductPrice).where(ProductPrice.product_id == product.id, ProductPrice.price_list_id == normal_list.id))
                assert product.status == "inactive"
                first_run = db.get(ErpSyncRun, first_run_id)
                assert product.stock_quantity == 42 and Decimal(str(product.price)) == Decimal("199.00"), (product.stock_quantity, product.price, first_run.status, first_run.error_summary)
                assert db.scalar(select(func.count()).select_from(ProductWarehouseStock).where(ProductWarehouseStock.product_id == product.id)) == 1
                assert db.scalar(select(func.count()).select_from(ProductImage).where(ProductImage.product_id == product.id)) >= 1
                assert first_run.details["images_imported"] == 1
                assert second.source_sync_status == "missing" and db.get(Product, second.id) is not None
                assert db.get(ErpSyncRun, first_run_id).status in {"completed", "completed_with_warnings"}

            second_run_id = product_sync_service.run_product_sync(trigger="manual", retry_backoffs=())
            with SessionLocal() as db:
                product = db.scalar(select(Product).where(Product.sku == sku))
                normal_list = db.scalar(select(PriceList).where(PriceList.code == "NORMAL"))
                unchanged_count = db.scalar(select(func.count()).select_from(ProductPrice).where(ProductPrice.product_id == product.id, ProductPrice.price_list_id == normal_list.id))
                assert unchanged_count == history_count
                assert db.get(ErpSyncRun, second_run_id).price_values_updated == 0
                assert db.get(ErpSyncRun, second_run_id).details["images_unchanged"] == 1

            preserved = (42, Decimal("199.00"))
            def failed_batches(_setting):
                raise RuntimeError("temporary source failure; password=must-not-leak")
                yield
            product_sync_service._source_batches = failed_batches
            failed_run_id = product_sync_service.run_product_sync(trigger="automatic", retry_backoffs=())
            with SessionLocal() as db:
                product = db.scalar(select(Product).where(Product.sku == sku))
                run = db.get(ErpSyncRun, failed_run_id)
                assert run.status == "failed" and "must-not-leak" not in run.error_summary
                assert (product.stock_quantity, Decimal(str(product.price))) == preserved
                assert db.scalar(select(func.count()).select_from(ProductStatusHistory).where(ProductStatusHistory.product_id == product.id)) == 3
                assert db.scalar(select(func.count()).select_from(AuditLog).where(AuditLog.action == "product_marked_inactive", AuditLog.identifier == sku)) == 2

            status = admin.get("/api/v1/admin/data-sync/status")
            assert status.status_code == 200
            assert status.json()["interval_seconds"] == 180
            assert not {"server", "username", "password", "connection_string"}.intersection(status.json())
            history = admin.get("/api/v1/admin/data-sync/history")
            assert history.status_code == 200 and len(history.json()) >= 3

        print("Product lifecycle and source synchronization smoke test: OK")
    finally:
        product_sync_service._source_batches = original_batches
        product_sync_service._source_image_metadata = original_image_metadata
        product_sync_service._source_image_content = original_image_content
        product_sync_service._erp_brand_names = original_brand_names
        product_sync_service._erp_price_level_rows = original_price_level_rows
        shutil.rmtree(test_directory, ignore_errors=True)


if __name__ == "__main__":
    main()
