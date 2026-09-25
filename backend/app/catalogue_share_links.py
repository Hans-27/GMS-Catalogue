import base64
import hashlib
import hmac
import ipaddress
import json
import re
import secrets
import time
import uuid
from datetime import UTC, datetime
from io import BytesIO
from urllib.parse import urlsplit

from cryptography.fernet import Fernet, InvalidToken
from fastapi import APIRouter, Depends, Header, HTTPException, Query, Request, status
from fastapi.responses import FileResponse, StreamingResponse
from sqlalchemy import func, or_, select, update
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session, noload, selectinload

from app.access import has_permission, has_record_access, require_any_permission, require_permission
from app.auth import get_current_user
from app.commerce import _apply_erp_card_prices, _cached_catalogue_pdf, _catalogue_query, _presentation_source, _safe_presentation
from app.commerce_models import Catalogue, CatalogueAudienceType, CatalogueShareLink, CatalogueVersion, PriceList, ProductPrice, UserBrandCataloguePriceMapping, UserCataloguePriceMapping
from app.commerce_schemas import CatalogueOnlineCover, CataloguePreviewPriceList
from app.config import settings
from app.database import get_db
from app.design_studio import _design_response_with_live_stock, _snapshot, _snapshot_with_current_live_stock
from app.design_studio_browser_export import render_browser_pdf
from app.design_studio_export import render_pdf as render_studio_pdf
from app.design_studio_models import CatalogueDesign, CatalogueDesignVersion, DesignAsset
from app.models import AuditLog, Product, ProductVideo, User
from app.storage import cover_storage, storage, video_storage
from app.security import hash_password, verify_password
from app.share_link_schemas import AudienceTypeCreate, AudienceTypeResponse, AudienceTypeUpdate, PublicCatalogueResponse, ShareLinkCreate, ShareLinkResponse, ShareLinkUpdate


router = APIRouter(tags=["Catalogue share links"])

PUBLIC_RESPONSE_EXCLUDE = {
    "catalogue_id": True,
    "status": True,
    "is_draft": True,
    "generated_at": True,
    "price_list": {"id": True},
    # Studio carousel bindings reference product UUIDs. Keep those identifiers in
    # the public payload so the share viewer can replace protected editor image
    # URLs with the product's public media URLs.
    "products": {"__all__": {"product_status": True}},
    "categories": {"__all__": {"id": True}},
    "cover": {"assets": {"__all__": {"id": True, "original_filename": True, "file_size": True, "width": True, "height": True}}},
}


def _now() -> datetime:
    return datetime.now(UTC)


def _utc(value: datetime | None) -> datetime | None:
    if value is None:
        return None
    return value.replace(tzinfo=UTC) if value.tzinfo is None else value.astimezone(UTC)


def _active_promotion_ids(presentation: PublicCatalogueResponse) -> set[str]:
    return {
        str(item.get("id") if isinstance(item, dict) else item.id)
        for item in (presentation.promotions or [])
    }


def _promotion_product_ids(presentation: PublicCatalogueResponse) -> dict[str, set[str]]:
    promotion_codes = {
        str(item.get("id") if isinstance(item, dict) else item.id):
        str(item.get("code") if isinstance(item, dict) else item.code)
        for item in (presentation.promotions or [])
    }
    return {
        promotion_id: {
            str(product.id)
            for product in presentation.products
            if product.promotion_code == promotion_code
        }
        for promotion_id, promotion_code in promotion_codes.items()
    }


def _audience_safe_promotion_document(document: dict, eligible_product_ids: set[str]) -> dict:
    filtered = dict(document)
    filtered["elements"] = [
        {
            **element,
            "visible": False,
        }
        if element.get("productId") and str(element["productId"]) not in eligible_product_ids
        else element
        for element in document.get("elements", [])
    ]
    return filtered


def _filter_public_promotion_pages(design_response, presentation: PublicCatalogueResponse):
    """Hide inactive campaigns and product frames that are unsafe for this audience."""
    active_promotion_ids = _active_promotion_ids(presentation)
    product_ids_by_promotion = _promotion_product_ids(presentation)
    visible_pages = []
    for page in design_response.pages:
        document = page.page_data_json
        promotion_id = document.get("promotionId") if isinstance(document, dict) else getattr(document, "promotionId", None)
        if page.page_type == "promotion" and promotion_id:
            promotion_key = str(promotion_id)
            if promotion_key not in active_promotion_ids:
                continue
            page.page_data_json = _audience_safe_promotion_document(
                document,
                product_ids_by_promotion.get(promotion_key, set()),
            )
        visible_pages.append(page)
    design_response.pages = visible_pages
    return design_response


def _filter_public_promotion_snapshot(snapshot: dict, presentation: PublicCatalogueResponse) -> dict:
    """Apply the same optional-page rule to browser and native PDF exports."""
    active_promotion_ids = _active_promotion_ids(presentation)
    product_ids_by_promotion = _promotion_product_ids(presentation)
    filtered = dict(snapshot)
    filtered_pages = []
    for page in snapshot.get("pages", []):
        document_key = "pageData" if "pageData" in page else "page_data_json"
        document = page.get(document_key) or {}
        promotion_id = document.get("promotionId")
        if page.get("pageType") == "promotion" and promotion_id:
            promotion_key = str(promotion_id)
            if promotion_key not in active_promotion_ids:
                continue
            page = {
                **page,
                document_key: _audience_safe_promotion_document(
                    document,
                    product_ids_by_promotion.get(promotion_key, set()),
                ),
            }
        filtered_pages.append(page)
    filtered["pages"] = filtered_pages
    return filtered


def _current_public_studio_snapshot(
    db: Session,
    design: CatalogueDesign,
    presentation: PublicCatalogueResponse,
) -> dict:
    """Build the PDF from the same current design shown by the public viewer.

    A published design remains editable through Studio autosave. Its immutable
    version is useful for history, but can have fewer pages than the current
    public design. Rendering that old snapshot made the browser viewer and its
    downloaded PDF disagree about the catalogue's page count.
    """
    return _filter_public_promotion_snapshot(
        _snapshot_with_public_prices(
            _snapshot_with_current_live_stock(db, _snapshot(design, db)),
            presentation,
        ),
        presentation,
    )


def _fernet() -> Fernet:
    digest = hashlib.sha256((settings.secret_key.get_secret_value() + ":catalogue-share-links").encode()).digest()
    return Fernet(base64.urlsafe_b64encode(digest))


def _token_hash(token: str) -> str:
    return hashlib.sha256(token.encode()).hexdigest()


def _new_token() -> tuple[str, str, str]:
    token = secrets.token_urlsafe(32)
    return token, _token_hash(token), _fernet().encrypt(token.encode()).decode()


def _token_value(link: CatalogueShareLink) -> str:
    try:
        return _fernet().decrypt(link.encrypted_token.encode()).decode()
    except (InvalidToken, ValueError) as exc:
        raise HTTPException(status_code=500, detail="The catalogue link token cannot be recovered. Regenerate this link.") from exc


def _trusted_proxy_public_origin(request: Request | None) -> str | None:
    if request is None or request.client is None:
        return None
    try:
        if not ipaddress.ip_address(request.client.host).is_loopback:
            return None
    except ValueError:
        return None

    forwarded_host = request.headers.get("x-forwarded-host", "").split(",", 1)[0].strip()
    forwarded_proto = request.headers.get("x-forwarded-proto", "").split(",", 1)[0].strip().lower()
    if not forwarded_host or forwarded_proto not in {"http", "https"}:
        return None

    try:
        parsed = urlsplit(f"{forwarded_proto}://{forwarded_host}")
        hostname = parsed.hostname
        parsed.port
    except ValueError:
        return None
    if (
        not hostname
        or parsed.username is not None
        or parsed.password is not None
        or parsed.path not in {"", "/"}
        or parsed.query
        or parsed.fragment
    ):
        return None
    try:
        address = ipaddress.ip_address(hostname)
    except ValueError:
        # A dashboard opened on localhost must still copy the configured LAN
        # origin so the resulting customer link works from another device.
        return None
    else:
        if address.is_loopback or not address.is_private:
            return None
    return f"{forwarded_proto}://{forwarded_host}"


def _public_url(link: CatalogueShareLink, *, request: Request | None = None) -> str:
    origin = _trusted_proxy_public_origin(request) or settings.public_app_url.rstrip("/")
    return f"{origin}/c/{_token_value(link)}"


def _audit(db: Session, action: str, catalogue: Catalogue, actor_id: uuid.UUID | None, audience: CatalogueAudienceType, details: dict | None = None) -> None:
    db.add(AuditLog(user_id=actor_id, action=action, module="catalogue_share_links", status="success", identifier=catalogue.slug, details={"catalogue_id": str(catalogue.id), "audience_type": audience.code, **(details or {})}))


def _audience_response(item: CatalogueAudienceType) -> AudienceTypeResponse:
    return AudienceTypeResponse(
        id=item.id, code=item.code, display_name=item.display_name,
        price_list_id=item.price_list_id, price_list_name=item.price_list.name if item.price_list else None,
        show_prices=item.show_prices, display_order=item.display_order,
        button_style_key=item.button_style_key, is_active=item.is_active,
        created_at=item.created_at, updated_at=item.updated_at,
    )


def _link_response(
    link: CatalogueShareLink,
    *,
    expose_url: bool,
    request: Request | None = None,
) -> ShareLinkResponse:
    price_list = link.price_list or link.audience_type.price_list
    return ShareLinkResponse(
        id=link.id, catalogue_id=link.catalogue_id,
        audience_type_id=link.audience_type_id, audience_code=link.audience_type.code,
        audience_name=link.audience_type.display_name,
        customer_code=link.customer_code or None,
        customer_name=link.customer_name,
        price_list_id=price_list.id if price_list else None,
        price_list_name=price_list.name if price_list else None,
        show_prices=link.show_prices and not bool(price_list and price_list.is_no_price),
        button_style_key=link.audience_type.button_style_key, status=link.status,
        version_mode=link.version_mode,
        fixed_version_number=link.fixed_version.version_number if link.fixed_version else None,
        expires_at=link.expires_at, has_password=bool(link.password_hash),
        allow_pdf_download=link.allow_pdf_download, allow_print=link.allow_print,
        created_at=link.created_at, updated_at=link.updated_at,
        last_accessed_at=link.last_accessed_at, view_count=link.view_count,
        public_url=_public_url(link, request=request) if expose_url else None,
    )


def _get_catalogue(db: Session, catalogue_id: uuid.UUID, actor: User, permission: str) -> Catalogue:
    catalogue = db.scalar(_catalogue_query().where(Catalogue.id == catalogue_id))
    if not catalogue:
        raise HTTPException(status_code=404, detail="Catalogue not found.")
    if not has_record_access(db, actor, permission, catalogue):
        raise HTTPException(status_code=403, detail="You do not have access to this catalogue.")
    return catalogue


def _get_link(db: Session, catalogue_id: uuid.UUID, link_id: uuid.UUID, actor_id: uuid.UUID) -> CatalogueShareLink:
    link = db.scalar(select(CatalogueShareLink).options(
        selectinload(CatalogueShareLink.audience_type).selectinload(CatalogueAudienceType.price_list),
        selectinload(CatalogueShareLink.fixed_version),
    ).where(
        CatalogueShareLink.id == link_id,
        CatalogueShareLink.catalogue_id == catalogue_id,
        CatalogueShareLink.created_by_id == actor_id,
    ))
    if not link:
        raise HTTPException(status_code=404, detail="Catalogue link not found.")
    return link


def _fixed_version(db: Session, catalogue: Catalogue, mode: str, number: int | None) -> CatalogueVersion | None:
    if mode == "latest_published":
        return None
    if not number:
        raise HTTPException(status_code=422, detail="Choose a published version for a fixed link.")
    version = db.scalar(select(CatalogueVersion).where(CatalogueVersion.catalogue_id == catalogue.id, CatalogueVersion.version_number == number))
    if not version:
        raise HTTPException(status_code=422, detail="The selected published version does not exist.")
    return version


def _customer_identity(values: ShareLinkCreate) -> tuple[str, str | None]:
    """Return a stable internal customer code and a display-safe name.

    Customer recipients never type or select this value. It is stored on the
    opaque link so one shared portal account can safely deliver different
    pricing contexts to different customers.
    """
    customer_name = (values.customer_name or "").strip()
    customer_code = (values.customer_code or "").strip().upper()
    if not customer_name and not customer_code:
        return "", None
    if not customer_name:
        customer_name = customer_code.replace("_", " ").replace("-", " ").title()
    if not customer_code:
        customer_code = re.sub(r"[^A-Z0-9]+", "-", customer_name.upper()).strip("-")
        if len(customer_code) < 2:
            customer_code = f"CUSTOMER-{customer_code or 'ACCOUNT'}"
    return customer_code[:80], customer_name[:160]


def _create_link(db: Session, catalogue: Catalogue, audience: CatalogueAudienceType, actor_id: uuid.UUID | None, payload: ShareLinkCreate | None = None) -> CatalogueShareLink:
    if catalogue.status not in {"published", "draft"} or catalogue.version < 1:
        raise HTTPException(status_code=409, detail="Publish at least one catalogue version before creating a public link.")
    values = payload or ShareLinkCreate(audience_type_id=audience.id)
    customer_code, customer_name = _customer_identity(values)
    token, digest, encrypted = _new_token()
    del token
    brand_mapping = db.scalar(
        select(UserBrandCataloguePriceMapping)
        .options(selectinload(UserBrandCataloguePriceMapping.price_list))
        .where(
            UserBrandCataloguePriceMapping.user_id == actor_id,
            UserBrandCataloguePriceMapping.brand_key == catalogue.brand.strip().casefold(),
            UserBrandCataloguePriceMapping.audience_type_id == audience.id,
        )
    ) if actor_id and catalogue.brand and catalogue.brand.strip() else None
    mapping = brand_mapping or (db.scalar(
        select(UserCataloguePriceMapping)
        .options(selectinload(UserCataloguePriceMapping.price_list))
        .where(
            UserCataloguePriceMapping.user_id == actor_id,
            UserCataloguePriceMapping.audience_type_id == audience.id,
        )
    ) if actor_id else None)
    price_list = mapping.price_list if mapping else audience.price_list
    link = CatalogueShareLink(
        catalogue_id=catalogue.id, audience_type_id=audience.id,
        customer_code=customer_code, customer_name=customer_name,
        price_list_id=price_list.id if price_list else None,
        show_prices=audience.show_prices and not bool(price_list and price_list.is_no_price),
        version_mode=values.version_mode,
        fixed_version_id=(_fixed_version(db, catalogue, values.version_mode, values.fixed_version_number).id if values.version_mode == "fixed_published" else None),
        token_hash=digest, encrypted_token=encrypted, status="active",
        expires_at=values.expires_at,
        password_hash=hash_password(values.password) if values.password else None,
        allow_pdf_download=values.allow_pdf_download, allow_print=values.allow_print,
        created_by_id=actor_id,
    )
    db.add(link)
    _audit(db, "catalogue_share_link_created", catalogue, actor_id, audience, {
        "version_mode": values.version_mode,
        "customer_code": customer_code or None,
    })
    return link


def ensure_share_links_for_catalogue(db: Session, catalogue: Catalogue, actor_id: uuid.UUID | None) -> None:
    existing = {item.audience_type_id: item for item in db.scalars(select(CatalogueShareLink).where(
        CatalogueShareLink.catalogue_id == catalogue.id,
        CatalogueShareLink.created_by_id == actor_id,
        CatalogueShareLink.customer_code == "",
    ))}
    audiences = db.scalars(select(CatalogueAudienceType).where(CatalogueAudienceType.is_active.is_(True)).order_by(CatalogueAudienceType.display_order)).all()
    for audience in audiences:
        link = existing.get(audience.id)
        if link is None:
            _create_link(db, catalogue, audience, actor_id)
        elif link.status == "revoked":
            link.status = "active"
            link.revoked_by_id = None
            link.revoked_at = None


@router.get("/catalogue-online-links", response_model=dict[str, str])
def online_catalogue_links(
    request: Request,
    actor: User = Depends(get_current_user),
    db: Session = Depends(get_db),
) -> dict[str, str]:
    """Return one safe public destination for each published catalogue a user can view."""

    catalogues = db.scalars(
        select(Catalogue)
        .options(
            noload(Catalogue.product_links),
            selectinload(Catalogue.share_links).selectinload(CatalogueShareLink.audience_type),
        )
        .where(Catalogue.status == "published")
    ).unique().all()
    now = _now()
    result: dict[str, str] = {}
    for catalogue in catalogues:
        if not has_record_access(db, actor, "catalogues.view", catalogue):
            continue
        preferred_codes = ("normal", "no_price") if catalogue.show_prices else ("no_price", "normal")
        eligible = [
            link
            for link in catalogue.share_links
            if not link.customer_code
            and link.status == "active"
            and not link.password_hash
            and link.audience_type.is_active
            and link.audience_type.code.casefold() in preferred_codes
            and (link.expires_at is None or _utc(link.expires_at) > now)
        ]
        eligible.sort(
            key=lambda link: (
                preferred_codes.index(link.audience_type.code.casefold()),
                str(link.id),
            )
        )
        if eligible:
            result[str(catalogue.id)] = _public_url(eligible[0], request=request)
        else:
            origin = _trusted_proxy_public_origin(request) or settings.public_app_url.rstrip("/")
            result[str(catalogue.id)] = f"{origin}/catalogues/{catalogue.id}/preview"
    return result


@router.get("/catalogue-audience-types", response_model=list[AudienceTypeResponse])
def list_audience_types(_: User = Depends(require_permission("catalogue_audience_types.view")), db: Session = Depends(get_db)) -> list[AudienceTypeResponse]:
    rows = db.scalars(select(CatalogueAudienceType).options(selectinload(CatalogueAudienceType.price_list)).order_by(CatalogueAudienceType.display_order, CatalogueAudienceType.id)).all()
    return [_audience_response(item) for item in rows]


@router.post("/catalogue-audience-types", response_model=AudienceTypeResponse, status_code=201)
def create_audience_type(payload: AudienceTypeCreate, actor: User = Depends(require_permission("catalogue_audience_types.manage")), db: Session = Depends(get_db)) -> AudienceTypeResponse:
    if payload.price_list_id and not db.get(PriceList, payload.price_list_id):
        raise HTTPException(status_code=422, detail="Price list not found.")
    item = CatalogueAudienceType(**payload.model_dump(), created_by_id=actor.id, updated_by_id=actor.id)
    db.add(item)
    try:
        db.commit()
    except IntegrityError:
        db.rollback(); raise HTTPException(status_code=409, detail="Audience code already exists.") from None
    return _audience_response(db.get(CatalogueAudienceType, item.id))


@router.patch("/catalogue-audience-types/{audience_id}", response_model=AudienceTypeResponse)
def update_audience_type(audience_id: int, payload: AudienceTypeUpdate, actor: User = Depends(require_permission("catalogue_audience_types.manage")), db: Session = Depends(get_db)) -> AudienceTypeResponse:
    item = db.get(CatalogueAudienceType, audience_id)
    if not item: raise HTTPException(status_code=404, detail="Audience type not found.")
    for field, value in payload.model_dump(exclude_unset=True).items(): setattr(item, field, value)
    item.updated_by_id = actor.id; db.commit(); return _audience_response(item)


@router.get("/catalogue-share-links/cards", response_model=dict[str, list[ShareLinkResponse]])
def card_links(
    request: Request,
    catalogue_ids: str | None = Query(default=None),
    actor: User = Depends(require_permission("catalogue_share_links.view")),
    db: Session = Depends(get_db),
) -> dict[str, list[ShareLinkResponse]]:
    expose = has_permission(actor, "catalogue_share_links.copy")
    audiences = db.scalars(select(CatalogueAudienceType).options(selectinload(CatalogueAudienceType.price_list)).where(CatalogueAudienceType.is_active.is_(True)).order_by(CatalogueAudienceType.display_order)).all()
    selected_ids: list[uuid.UUID] = []
    if catalogue_ids:
        raw_ids = [value.strip() for value in catalogue_ids.split(",") if value.strip()]
        if len(raw_ids) > 100:
            raise HTTPException(status_code=422, detail="Choose no more than 100 catalogue cards at once.")
        try:
            selected_ids = [uuid.UUID(value) for value in raw_ids]
        except ValueError as error:
            raise HTTPException(status_code=422, detail="A catalogue card identifier is invalid.") from error
    catalogue_query = select(Catalogue).options(
        noload(Catalogue.product_links),
        selectinload(Catalogue.share_links)
        .selectinload(CatalogueShareLink.audience_type)
        .selectinload(CatalogueAudienceType.price_list),
        selectinload(Catalogue.share_links).selectinload(CatalogueShareLink.fixed_version),
    )
    if selected_ids:
        catalogue_query = catalogue_query.where(Catalogue.id.in_(selected_ids))
    catalogues = [
        item
        for item in db.scalars(catalogue_query).unique().all()
        if has_record_access(db, actor, "catalogues.view", item)
    ]
    result: dict[str, list[ShareLinkResponse]] = {}
    mappings = {
        item.audience_type_id: item
        for item in db.scalars(
            select(UserCataloguePriceMapping)
            .options(selectinload(UserCataloguePriceMapping.price_list))
            .where(UserCataloguePriceMapping.user_id == actor.id)
        )
    }
    brand_mappings = {
        (item.brand_key, item.audience_type_id): item
        for item in db.scalars(
            select(UserBrandCataloguePriceMapping)
            .options(selectinload(UserBrandCataloguePriceMapping.price_list))
            .where(UserBrandCataloguePriceMapping.user_id == actor.id)
        )
    }
    for catalogue in catalogues:
        by_audience = {
            link.audience_type_id: link
            for link in catalogue.share_links
            if link.created_by_id == actor.id and not link.customer_code
        }
        rows = []
        for audience in audiences:
            link = by_audience.get(audience.id)
            if link:
                rows.append(_link_response(link, expose_url=expose, request=request))
            else:
                brand_key = catalogue.brand.strip().casefold() if catalogue.brand else ""
                mapping = brand_mappings.get((brand_key, audience.id)) or mappings.get(audience.id)
                price_list = mapping.price_list if mapping else audience.price_list
                rows.append(ShareLinkResponse(
                    catalogue_id=catalogue.id, audience_type_id=audience.id,
                    audience_code=audience.code, audience_name=audience.display_name,
                    price_list_id=price_list.id if price_list else None,
                    price_list_name=price_list.name if price_list else None,
                    show_prices=audience.show_prices and not bool(price_list and price_list.is_no_price),
                    button_style_key=audience.button_style_key, status="not_generated",
                    version_mode="latest_published", fixed_version_number=None, expires_at=None,
                    has_password=False, allow_pdf_download=True, allow_print=True,
                    created_at=None, updated_at=None, last_accessed_at=None, view_count=0,
                ))
        result[str(catalogue.id)] = rows
    return result


@router.get("/catalogues/{catalogue_id}/share-links", response_model=list[ShareLinkResponse])
def list_share_links(catalogue_id: uuid.UUID, request: Request, actor: User = Depends(require_permission("catalogue_share_links.view")), db: Session = Depends(get_db)) -> list[ShareLinkResponse]:
    _get_catalogue(db, catalogue_id, actor, "catalogues.view")
    expose = has_permission(actor, "catalogue_share_links.copy")
    rows = db.scalars(select(CatalogueShareLink).options(
        selectinload(CatalogueShareLink.audience_type).selectinload(CatalogueAudienceType.price_list), selectinload(CatalogueShareLink.fixed_version)
    ).where(
        CatalogueShareLink.catalogue_id == catalogue_id,
        CatalogueShareLink.created_by_id == actor.id,
    ).order_by(CatalogueShareLink.audience_type_id)).unique().all()
    return [_link_response(item, expose_url=expose, request=request) for item in rows]


@router.post("/catalogues/{catalogue_id}/share-links", response_model=ShareLinkResponse, status_code=201)
def create_share_link(
    catalogue_id: uuid.UUID,
    payload: ShareLinkCreate,
    request: Request,
    actor: User = Depends(
        require_any_permission(
            "catalogue_share_links.create",
            "catalogue_share_links.copy",
        )
    ),
    db: Session = Depends(get_db),
) -> ShareLinkResponse:
    # A Sales user may materialize a missing personal link as part of the copy
    # action. This does not grant catalogue editing or link-management access.
    catalogue = _get_catalogue(db, catalogue_id, actor, "catalogues.view")
    audience = db.get(CatalogueAudienceType, payload.audience_type_id)
    if not audience or not audience.is_active: raise HTTPException(status_code=422, detail="Audience type is unavailable.")
    customer_code, _ = _customer_identity(payload)
    duplicate_query = select(CatalogueShareLink.id).where(
        CatalogueShareLink.catalogue_id == catalogue.id,
        CatalogueShareLink.created_by_id == actor.id,
        CatalogueShareLink.customer_code == customer_code,
    )
    # Named recipients have one unambiguous price profile per catalogue. Base
    # shortcuts keep one row for each audience and use an empty customer code.
    if not customer_code:
        duplicate_query = duplicate_query.where(
            CatalogueShareLink.audience_type_id == audience.id,
        )
    if db.scalar(duplicate_query):
        raise HTTPException(
            status_code=409,
            detail=(
                "This customer already has a catalogue link for the selected price profile."
                if payload.customer_name or payload.customer_code
                else "This audience already has a catalogue link."
            ),
        )
    link = _create_link(db, catalogue, audience, actor.id, payload)
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="This customer catalogue link already exists.") from None
    return _link_response(_get_link(db, catalogue.id, link.id, actor.id), expose_url=True, request=request)


@router.patch("/catalogues/{catalogue_id}/share-links/{link_id}", response_model=ShareLinkResponse)
def update_share_link(catalogue_id: uuid.UUID, link_id: uuid.UUID, payload: ShareLinkUpdate, request: Request, actor: User = Depends(require_permission("catalogue_share_links.edit")), db: Session = Depends(get_db)) -> ShareLinkResponse:
    catalogue = _get_catalogue(db, catalogue_id, actor, "catalogues.edit"); link = _get_link(db, catalogue_id, link_id, actor.id)
    changed = payload.model_dump(exclude_unset=True)
    if "customer_name" in changed:
        if not link.customer_code:
            raise HTTPException(status_code=422, detail="Audience shortcut links cannot be assigned to a customer.")
        link.customer_name = (payload.customer_name or "").strip() or link.customer_name
    if "version_mode" in changed:
        link.version_mode = changed["version_mode"]
        link.fixed_version_id = (_fixed_version(db, catalogue, link.version_mode, payload.fixed_version_number).id if link.version_mode == "fixed_published" else None)
    if "expires_at" in changed: link.expires_at = payload.expires_at
    if "password" in changed: link.password_hash = hash_password(payload.password) if payload.password else None
    if payload.allow_pdf_download is not None: link.allow_pdf_download = payload.allow_pdf_download
    if payload.allow_print is not None: link.allow_print = payload.allow_print
    _audit(db, "catalogue_share_link_updated", catalogue, actor.id, link.audience_type, {
        "fields": sorted(changed),
        "customer_code": link.customer_code or None,
    })
    db.commit(); return _link_response(_get_link(db, catalogue_id, link_id, actor.id), expose_url=has_permission(actor, "catalogue_share_links.copy"), request=request)


@router.post("/catalogues/{catalogue_id}/share-links/{link_id}/regenerate", response_model=ShareLinkResponse)
def regenerate_share_link(catalogue_id: uuid.UUID, link_id: uuid.UUID, request: Request, actor: User = Depends(require_permission("catalogue_share_links.regenerate")), db: Session = Depends(get_db)) -> ShareLinkResponse:
    catalogue = _get_catalogue(db, catalogue_id, actor, "catalogues.edit"); link = _get_link(db, catalogue_id, link_id, actor.id)
    _, link.token_hash, link.encrypted_token = _new_token(); link.status = "active"; link.revoked_at = None; link.revoked_by_id = None
    _audit(db, "catalogue_share_link_regenerated", catalogue, actor.id, link.audience_type); db.commit(); return _link_response(link, expose_url=True, request=request)


@router.post("/catalogues/{catalogue_id}/share-links/{link_id}/revoke", response_model=ShareLinkResponse)
def revoke_share_link(catalogue_id: uuid.UUID, link_id: uuid.UUID, request: Request, actor: User = Depends(require_permission("catalogue_share_links.revoke")), db: Session = Depends(get_db)) -> ShareLinkResponse:
    catalogue = _get_catalogue(db, catalogue_id, actor, "catalogues.edit"); link = _get_link(db, catalogue_id, link_id, actor.id)
    link.status = "revoked"; link.revoked_at = _now(); link.revoked_by_id = actor.id
    _audit(db, "catalogue_share_link_revoked", catalogue, actor.id, link.audience_type); db.commit(); return _link_response(link, expose_url=False, request=request)


@router.post("/catalogues/{catalogue_id}/share-links/{link_id}/activate", response_model=ShareLinkResponse)
def activate_share_link(catalogue_id: uuid.UUID, link_id: uuid.UUID, request: Request, actor: User = Depends(require_permission("catalogue_share_links.edit")), db: Session = Depends(get_db)) -> ShareLinkResponse:
    catalogue = _get_catalogue(db, catalogue_id, actor, "catalogues.edit"); link = _get_link(db, catalogue_id, link_id, actor.id)
    if link.expires_at and _utc(link.expires_at) <= _now(): raise HTTPException(status_code=409, detail="Remove or change the expiration date before activating this link.")
    link.status = "active"; link.revoked_at = None; link.revoked_by_id = None
    _audit(db, "catalogue_share_link_activated", catalogue, actor.id, link.audience_type); db.commit(); return _link_response(link, expose_url=has_permission(actor, "catalogue_share_links.copy"), request=request)


@router.delete("/catalogues/{catalogue_id}/share-links/{link_id}", status_code=204)
def delete_share_link(catalogue_id: uuid.UUID, link_id: uuid.UUID, actor: User = Depends(require_permission("catalogue_share_links.delete")), db: Session = Depends(get_db)) -> None:
    catalogue = _get_catalogue(db, catalogue_id, actor, "catalogues.edit"); link = _get_link(db, catalogue_id, link_id, actor.id)
    _audit(db, "catalogue_share_link_deleted", catalogue, actor.id, link.audience_type); db.delete(link); db.commit()


def _media_signature(token: str, video_id: uuid.UUID, expires: int) -> str:
    message = f"{_token_hash(token)}:{video_id}:{expires}".encode()
    return hmac.new(settings.secret_key.get_secret_value().encode(), message, hashlib.sha256).hexdigest()


def _valid_media_signature(token: str, video_id: uuid.UUID, expires: int | None, signature: str | None) -> bool:
    return bool(expires and signature and expires >= int(time.time()) and hmac.compare_digest(signature, _media_signature(token, video_id, expires)))


def _online_cover_signature(token: str, asset_id: uuid.UUID, expires: int) -> str:
    message = f"online-cover:{_token_hash(token)}:{asset_id}:{expires}".encode()
    return hmac.new(settings.secret_key.get_secret_value().encode(), message, hashlib.sha256).hexdigest()


def _cover_asset_signature(token: str, asset_id: uuid.UUID, expires: int) -> str:
    message = f"catalogue-cover:{_token_hash(token)}:{asset_id}:{expires}".encode()
    return hmac.new(settings.secret_key.get_secret_value().encode(), message, hashlib.sha256).hexdigest()


def _published_online_cover(db: Session, snapshot: dict, token: str) -> CatalogueOnlineCover | None:
    cover = snapshot.get("online_cover")
    if not isinstance(cover, dict):
        return None
    try:
        asset_id = uuid.UUID(cover["asset_id"])
    except (KeyError, ValueError, TypeError):
        return None
    asset = db.get(DesignAsset, asset_id)
    if not asset or asset.deleted_at or not cover_storage.resolve(asset.storage_key).is_file():
        return None
    expires = int(time.time()) + 1800
    signature = _online_cover_signature(token, asset_id, expires)
    return CatalogueOnlineCover(asset_id=asset_id, file_name=asset.original_filename, width=asset.width, height=asset.height,
        url=f"{settings.api_prefix}/v1/public/catalogues/{token}/online-cover/content?asset_id={asset_id}&expires={expires}&signature={signature}")


def _mapped_public_price_lists(
    db: Session,
    link: CatalogueShareLink,
    products: list,
) -> dict[uuid.UUID, PriceList | None]:
    """Resolve the link creator's current audience mapping for every brand.

    A catalogue link selects the customer audience. The ERP price list is then
    resolved per product brand, so multi-brand catalogues do not incorrectly
    reuse one global price level.
    """
    audience = link.audience_type
    generic = None
    brand_mappings: dict[str, UserBrandCataloguePriceMapping] = {}
    catalogue_brand_key = (
        link.catalogue.brand.strip().casefold()
        if link.catalogue and link.catalogue.brand and link.catalogue.brand.strip()
        else None
    )

    def brand_key_for(product) -> str | None:
        # A single-brand catalogue's configured brand is authoritative. Multi-
        # brand Studio catalogues leave it empty and resolve each product brand.
        return catalogue_brand_key or (
            product.brand.strip().casefold()
            if product.brand and product.brand.strip()
            else None
        )

    if link.created_by_id:
        generic = db.scalar(
            select(UserCataloguePriceMapping)
            .options(selectinload(UserCataloguePriceMapping.price_list))
            .where(
                UserCataloguePriceMapping.user_id == link.created_by_id,
                UserCataloguePriceMapping.audience_type_id == audience.id,
            )
        )
        brand_keys = {
            brand_key_for(product)
            for product in products
            if brand_key_for(product)
        }
        if brand_keys:
            brand_mappings = {
                row.brand_key: row
                for row in db.scalars(
                    select(UserBrandCataloguePriceMapping)
                    .options(selectinload(UserBrandCataloguePriceMapping.price_list))
                    .where(
                        UserBrandCataloguePriceMapping.user_id == link.created_by_id,
                        UserBrandCataloguePriceMapping.audience_type_id == audience.id,
                        UserBrandCataloguePriceMapping.brand_key.in_(brand_keys),
                    )
                )
            }
    fallback = generic.price_list if generic else (link.price_list or audience.price_list)
    return {
        product.id: (
            brand_mappings[brand_key_for(product)].price_list
            if brand_key_for(product) in brand_mappings
            else fallback
        )
        for product in products
        if product.id
    }


def _snapshot_with_public_prices(snapshot: dict, presentation: PublicCatalogueResponse) -> dict:
    """Apply the link's resolved audience/brand prices to a Studio PDF snapshot.

    The immutable Studio version keeps its original design and ERP bindings. A
    public export must not reuse those saved amounts because the same catalogue
    can be shared to audiences whose brand mappings point at different ERP
    price lists.
    """
    hydrated = json.loads(json.dumps(snapshot or {}))
    products = {
        str(product.id): product
        for product in presentation.products
        if product.id is not None
    }

    def price_payload(product) -> dict | None:
        if product is None or product.price is None:
            return None
        return {
            "amount": str(product.price),
            "currency": product.currency or presentation.currency or "THB",
        }

    def price_label(product) -> str:
        payload = price_payload(product)
        if not payload:
            return ""
        try:
            amount = f"{float(payload['amount']):,.2f}"
        except (TypeError, ValueError):
            amount = str(payload["amount"])
        return f"{payload['currency']} {amount}".strip()

    for product_id, product_data in (hydrated.get("productData") or {}).items():
        product = products.get(str(product_id))
        payload = price_payload(product)
        product_data["price"] = payload["amount"] if payload else None
        product_data["price_currency"] = payload["currency"] if payload else None
        product_data["prices"] = {"1": payload} if payload else {}
        product_data["prices_by_list"] = {}

    for page in hydrated.get("pages", []):
        document_key = "pageData" if "pageData" in page else "page_data_json"
        document = page.get(document_key) or {}
        for element in document.get("elements", []):
            product = products.get(str(element.get("productId") or ""))
            style = element.setdefault("style", {})
            label = price_label(product)
            if element.get("type") == "product_card" and element.get("productId") and style.get("useErpPrice") is not False:
                style["productPrice"] = label
                style["productSecondaryPrice"] = ""
                style["showProductPrice"] = bool(
                    label
                    and style.get("priceMode") != "no_price"
                    and style.get("showProductPrice") is not False
                )
                style["showSecondaryPrice"] = False
                continue
            binding = str(element.get("binding") or "")
            if element.get("productId") and (
                style.get("priceSource") == "erp"
                or binding.startswith("{{product.price")
            ):
                element["text"] = label
                # The native renderer resolves a binding before element text.
                # Remove it only from this ephemeral public snapshot so the
                # audience-safe amount above is always used.
                element["binding"] = ""
                if not label:
                    element["visible"] = False
        page[document_key] = document
    return hydrated


def _public_link(db: Session, token: str, password: str | None, action: str, request: Request, *, signed_media: bool = False) -> tuple[CatalogueShareLink, PublicCatalogueResponse]:
    if not token or len(token) < 20: raise HTTPException(status_code=404, detail="Catalogue link not found.")
    link = db.scalar(select(CatalogueShareLink).options(
        selectinload(CatalogueShareLink.catalogue),
        selectinload(CatalogueShareLink.audience_type).selectinload(CatalogueAudienceType.price_list),
        selectinload(CatalogueShareLink.fixed_version),
    ).where(CatalogueShareLink.token_hash == _token_hash(token)))
    if not link: raise HTTPException(status_code=404, detail="Catalogue link not found.")
    now = _now()
    if link.status == "revoked": raise HTTPException(status_code=410, detail="This catalogue link has been revoked.")
    if link.expires_at and _utc(link.expires_at) <= now:
        link.status = "expired"; db.commit(); raise HTTPException(status_code=410, detail="This catalogue link has expired.")
    if link.status != "active": raise HTTPException(status_code=410, detail="This catalogue link is unavailable.")
    if link.password_hash and not signed_media and (not password or not verify_password(password, link.password_hash)):
        raise HTTPException(status_code=401, detail="A valid catalogue link password is required.")
    if action in {"public_catalogue_pdf_downloaded", "public_catalogue_category_excel_downloaded"} and not link.allow_pdf_download:
        raise HTTPException(status_code=403, detail="Downloads are disabled for this catalogue link.")
    catalogue = link.catalogue
    if catalogue.status not in {"published", "draft"} or catalogue.version < 1:
        raise HTTPException(status_code=404, detail="Published catalogue not found.")
    version_number = link.fixed_version.version_number if link.version_mode == "fixed_published" and link.fixed_version else catalogue.version
    snapshot, resolved_version = _presentation_source(db, catalogue, version_number)
    presentation = _safe_presentation(db, catalogue, snapshot, resolved_version, None)
    presentation.online_cover = _published_online_cover(db, snapshot, token)
    audience = link.audience_type
    price_lists_by_product = _mapped_public_price_lists(db, link, presentation.products)
    numeric_price_list_ids = {
        price_list.id
        for price_list in price_lists_by_product.values()
        if audience.show_prices and price_list and price_list.is_active and not price_list.is_no_price
    }
    presentation.audience = audience.display_name
    presentation.show_prices = bool(numeric_price_list_ids)
    display_price_list = next(
        (price_list for price_list in price_lists_by_product.values() if price_list and price_list.id in numeric_price_list_ids),
        None,
    )
    presentation.price_list = (
        CataloguePreviewPriceList(id=display_price_list.id, name=display_price_list.name, show_price=True)
        if display_price_list and len(numeric_price_list_ids) == 1
        else None
    )
    currencies = {
        price_list.currency
        for price_list in price_lists_by_product.values()
        if price_list and price_list.id in numeric_price_list_ids
    }
    presentation.currency = next(iter(currencies)) if len(currencies) == 1 else None
    prices: dict[tuple[uuid.UUID, int], ProductPrice] = {}
    product_ids = [product.id for product in presentation.products if product.id]
    if product_ids and numeric_price_list_ids:
        rows = db.scalars(select(ProductPrice).where(
            ProductPrice.product_id.in_(product_ids),
            ProductPrice.price_list_id.in_(numeric_price_list_ids),
            ProductPrice.status == "active", ProductPrice.effective_from <= now,
            or_(ProductPrice.expires_at.is_(None), ProductPrice.expires_at > now),
        ).order_by(ProductPrice.effective_from.desc())).all()
        for row in rows:
            prices.setdefault((row.product_id, row.price_list_id), row)
    for product in presentation.products:
        mapped_price_list = price_lists_by_product.get(product.id)
        price = (
            prices.get((product.id, mapped_price_list.id))
            if mapped_price_list and mapped_price_list.id in numeric_price_list_ids
            else None
        )
        product.price = price.amount if price else None
        product.currency = price.currency if price else None
    _apply_erp_card_prices(db, presentation)
    from app.promotion_service import apply_promotions_to_catalogue
    presentation.promotions = apply_promotions_to_catalogue(
        db, catalogue_id=catalogue.id, products=presentation.products,
        audience_type_id=audience.id,
        price_list_id=display_price_list.id if display_price_list and len(numeric_price_list_ids) == 1 else None,
        show_prices=presentation.show_prices,
        price_list_ids_by_product={
            product_id: price_list.id if price_list else None
            for product_id, price_list in price_lists_by_product.items()
        },
        show_prices_by_product={
            product_id: bool(price_list and price_list.id in numeric_price_list_ids)
            for product_id, price_list in price_lists_by_product.items()
        },
    )
    for product in presentation.products:
        if product.video and product.video.provider == "internal":
            expires = int(time.time()) + 300
            signature = _media_signature(token, product.video.id, expires)
            signed_query = f"?expires={expires}&signature={signature}"
            product.video.playback_url = f"{settings.api_prefix}/v1/public/catalogues/{token}/videos/{product.video.id}/content{signed_query}"
            snapshot_product = next((item for item in snapshot.get("products", []) if str(item.get("id")) == str(product.id)), None)
            snapshot_video = snapshot_product.get("video") if snapshot_product else None
            if snapshot_video and snapshot_video.get("thumbnail_storage_key"):
                product.video.thumbnail_url = f"{settings.api_prefix}/v1/public/catalogues/{token}/videos/{product.video.id}/thumbnail{signed_query}"
            if snapshot_video and snapshot_video.get("caption_storage_key"):
                product.video.caption_url = f"{settings.api_prefix}/v1/public/catalogues/{token}/videos/{product.video.id}/caption{signed_query}"
    db.execute(update(CatalogueShareLink).where(CatalogueShareLink.id == link.id).values(last_accessed_at=now, view_count=CatalogueShareLink.view_count + 1))
    _audit(db, action, catalogue, None, audience, {"version": resolved_version, "client_ip": request.client.host if request.client else None, "user_agent": request.headers.get("user-agent", "")[:240]})
    db.commit()
    public_presentation = PublicCatalogueResponse(
        **presentation.model_dump(),
        audience_type=audience.display_name,
        audience_code=audience.code,
        customer_name=link.customer_name,
        allow_pdf_download=link.allow_pdf_download,
        allow_print=link.allow_print,
        password_protected=bool(link.password_hash),
    )

    # Cover storage keys are deliberately excluded from API serialization, but
    # the PDF renderer still needs them to read private cover assets from disk.
    # Rebuilding the public response above strips those internal-only fields,
    # which previously left public PDF downloads with a plain colour cover even
    # though the browser catalogue displayed the saved artwork.
    if presentation.cover and public_presentation.cover:
        source_assets = {
            asset.id: asset
            for asset in presentation.cover.assets
            if asset.id is not None
        }
        for public_asset in public_presentation.cover.assets:
            source_asset = source_assets.get(public_asset.id)
            if source_asset:
                public_asset.storage_key = source_asset.storage_key
                public_asset.preview_storage_key = source_asset.preview_storage_key
                expires = int(time.time()) + 1800
                signature = _cover_asset_signature(token, source_asset.id, expires)
                public_url = (
                    f"{settings.api_prefix}/v1/public/catalogues/{token}/cover/assets/"
                    f"{source_asset.id}/content?expires={expires}&signature={signature}"
                )
                public_asset.file_url = public_url
                public_asset.preview_url = (
                    f"{public_url}&preview=true"
                    if source_asset.preview_storage_key
                    else public_url
                )

    return link, public_presentation


@router.get("/public/catalogues/{token}", response_model=PublicCatalogueResponse, response_model_exclude_none=True, response_model_exclude=PUBLIC_RESPONSE_EXCLUDE)
def public_catalogue_by_token(token: str, request: Request, x_catalogue_password: str | None = Header(default=None), db: Session = Depends(get_db)) -> PublicCatalogueResponse:
    return _public_link(db, token, x_catalogue_password, "public_catalogue_link_opened", request)[1]


@router.get("/public/catalogues/{token}/studio")
def public_catalogue_studio(token: str, request: Request, x_catalogue_password: str | None = Header(default=None), db: Session = Depends(get_db)):
    link, presentation = _public_link(db, token, x_catalogue_password, "public_catalogue_link_opened", request)
    design = db.scalar(select(CatalogueDesign).where(
        CatalogueDesign.catalogue_id == link.catalogue_id,
        CatalogueDesign.status == "published",
        CatalogueDesign.deleted_at.is_(None),
    ))
    if not design:
        raise HTTPException(status_code=404, detail="Published Studio catalogue not found.")
    response = _design_response_with_live_stock(db, design)
    response.online_cover_json = None
    response.published_online_cover_json = None
    return _filter_public_promotion_pages(
        response,
        presentation,
    )


@router.get("/public/catalogues/{token}/online-cover/content")
def public_online_cover_content(token: str, request: Request, asset_id: uuid.UUID | None = None,
    expires: int | None = None, signature: str | None = None,
    x_catalogue_password: str | None = Header(default=None), db: Session = Depends(get_db)):
    # Signed URLs let password-protected <img> tags load without disclosing the
    # password. They remain bound to the token, asset and selected published version.
    signed = bool(asset_id and expires and signature and expires >= int(time.time()) and
        hmac.compare_digest(signature, _online_cover_signature(token, asset_id, expires)))
    if signature and not signed:
        raise HTTPException(status_code=403, detail="Online cover access has expired. Reload the catalogue.")
    _, presentation = _public_link(db, token, x_catalogue_password, "public_catalogue_media_opened", request, signed_media=signed)
    cover = presentation.online_cover
    if signed and (not cover or cover.asset_id != asset_id):
        raise HTTPException(status_code=403, detail="This cover is no longer published for this link.")
    if not cover:
        raise HTTPException(status_code=404, detail="Published online cover not found.")
    asset = db.get(DesignAsset, cover.asset_id)
    return FileResponse(cover_storage.resolve(asset.storage_key), media_type=asset.mime_type,
        content_disposition_type="inline", headers={"Cache-Control": "private, no-store"})


@router.get("/public/catalogues/{token}/cover/assets/{asset_id}/content")
def public_catalogue_cover_asset(
    token: str,
    asset_id: uuid.UUID,
    request: Request,
    preview: bool = False,
    expires: int | None = None,
    signature: str | None = None,
    x_catalogue_password: str | None = Header(default=None),
    db: Session = Depends(get_db),
):
    signed = bool(
        expires
        and signature
        and expires >= int(time.time())
        and hmac.compare_digest(signature, _cover_asset_signature(token, asset_id, expires))
    )
    if signature and not signed:
        raise HTTPException(status_code=403, detail="Cover access has expired. Reload the catalogue.")
    _, presentation = _public_link(
        db,
        token,
        x_catalogue_password,
        "public_catalogue_media_opened",
        request,
        signed_media=signed,
    )
    asset = next(
        (
            item
            for item in (presentation.cover.assets if presentation.cover else [])
            if item.id == asset_id
        ),
        None,
    )
    if not asset:
        raise HTTPException(status_code=404, detail="Cover asset not found in this catalogue.")
    key = asset.preview_storage_key if preview and asset.preview_storage_key else asset.storage_key
    path = cover_storage.resolve(key)
    if not path.is_file():
        raise HTTPException(status_code=404, detail="Cover asset file is unavailable.")
    media_type = "image/webp" if key == asset.preview_storage_key else asset.mime_type
    return FileResponse(
        path,
        media_type=media_type,
        content_disposition_type="inline",
        headers={"Cache-Control": "private, max-age=300"},
    )


@router.get("/public/catalogues/{token}/studio/assets/{asset_id}/content")
def public_catalogue_studio_asset(token: str, asset_id: uuid.UUID, request: Request, x_catalogue_password: str | None = Header(default=None), db: Session = Depends(get_db)):
    link, _ = _public_link(db, token, x_catalogue_password, "public_catalogue_media_opened", request)
    design = db.scalar(select(CatalogueDesign).where(
        CatalogueDesign.catalogue_id == link.catalogue_id,
        CatalogueDesign.status == "published",
        CatalogueDesign.deleted_at.is_(None),
    ))
    if not design or not any(str(asset_id) in str(page.page_data_json) for page in design.pages):
        raise HTTPException(status_code=404, detail="Studio asset not found in this catalogue.")
    asset = db.get(DesignAsset, asset_id)
    if not asset or asset.deleted_at:
        raise HTTPException(status_code=404, detail="Studio asset not found.")
    path = cover_storage.resolve(asset.storage_key)
    if not path.is_file():
        raise HTTPException(status_code=404, detail="Studio asset file is unavailable.")
    return FileResponse(path, media_type=asset.mime_type, filename=asset.original_filename, content_disposition_type="inline", headers={"Cache-Control": "private, max-age=300"})


def _public_video_asset(token: str, video_id: uuid.UUID, kind: str, request: Request, password: str | None, expires: int | None, signature: str | None, db: Session):
    signed_media = _valid_media_signature(token, video_id, expires, signature)
    if not signed_media and not password:
        raise HTTPException(status_code=401, detail="This video playback link has expired. Reload the catalogue.")
    _, presentation = _public_link(db, token, password, "public_catalogue_video_opened", request, signed_media=signed_media)
    product = next((item for item in presentation.products if item.video and item.video.id == video_id), None)
    if not product or product.product_status != "active":
        raise HTTPException(status_code=404, detail="This video is currently unavailable.")
    video = db.scalar(select(ProductVideo).where(
        ProductVideo.id == video_id, ProductVideo.product_id == product.id,
        ProductVideo.is_active.is_(True), ProductVideo.deleted_at.is_(None),
        ProductVideo.show_in_public_catalogue.is_(True),
    ))
    if not video:
        raise HTTPException(status_code=404, detail="This video is currently unavailable.")
    key = (video.playback_storage_key or video.storage_key) if kind == "content" else (video.thumbnail_storage_key if kind == "thumbnail" else video.caption_storage_key)
    if not key:
        raise HTTPException(status_code=404, detail="This video asset is unavailable.")
    path = video_storage.resolve(key)
    if not path.is_file():
        raise HTTPException(status_code=404, detail="This video asset is unavailable.")
    mime = video.mime_type if kind == "content" else ("text/vtt" if kind == "caption" else None)
    headers = {"Cache-Control": "private, max-age=300" if password else "public, max-age=300", "Accept-Ranges": "bytes"}
    return FileResponse(path, media_type=mime, content_disposition_type="inline", headers=headers)


@router.get("/public/catalogues/{token}/videos/{video_id}/content")
def public_catalogue_video(token: str, video_id: uuid.UUID, request: Request, expires: int | None = None, signature: str | None = None, x_catalogue_password: str | None = Header(default=None), db: Session = Depends(get_db)):
    return _public_video_asset(token, video_id, "content", request, x_catalogue_password, expires, signature, db)


@router.get("/public/catalogues/{token}/videos/{video_id}/thumbnail")
def public_catalogue_video_thumbnail(token: str, video_id: uuid.UUID, request: Request, expires: int | None = None, signature: str | None = None, x_catalogue_password: str | None = Header(default=None), db: Session = Depends(get_db)):
    return _public_video_asset(token, video_id, "thumbnail", request, x_catalogue_password, expires, signature, db)


@router.get("/public/catalogues/{token}/videos/{video_id}/caption")
def public_catalogue_video_caption(token: str, video_id: uuid.UUID, request: Request, expires: int | None = None, signature: str | None = None, x_catalogue_password: str | None = Header(default=None), db: Session = Depends(get_db)):
    return _public_video_asset(token, video_id, "caption", request, x_catalogue_password, expires, signature, db)


def _excel_text(value: object) -> str:
    """Keep public text literal when Excel opens the downloaded workbook."""
    text = str(value or "").strip()
    return f"'{text}" if text.startswith(("=", "+", "-", "@")) else text


def _add_public_product_image(worksheet, product, row_number: int) -> None:
    image_url = str(product.main_image_url or "")
    if not image_url.startswith("/uploads/"):
        worksheet.cell(row_number, 1).value = "Image unavailable"
        return
    try:
        path = storage.resolve(image_url.removeprefix("/uploads/"))
        if not path.is_file():
            worksheet.cell(row_number, 1).value = "Image unavailable"
            return
        from PIL import Image as PillowImage, ImageOps, UnidentifiedImageError
        from openpyxl.drawing.image import Image as WorksheetImage

        # Setting WorksheetImage.width/height only changes the display size;
        # openpyxl would otherwise embed the original multi-megabyte asset.
        # Generate a real bounded thumbnail so large categories remain quick
        # to download and do not produce enormous workbooks.
        with PillowImage.open(path) as source:
            normalized = ImageOps.exif_transpose(source)
            normalized.thumbnail((144, 144), PillowImage.Resampling.LANCZOS)
            thumbnail = PillowImage.new("RGB", normalized.size, "white")
            if normalized.mode in {"RGBA", "LA"} or (
                normalized.mode == "P" and "transparency" in normalized.info
            ):
                rgba = normalized.convert("RGBA")
                thumbnail.paste(rgba, mask=rgba.getchannel("A"))
            else:
                thumbnail.paste(normalized.convert("RGB"))
            thumbnail_bytes = BytesIO()
            thumbnail.save(
                thumbnail_bytes,
                format="JPEG",
                quality=65,
                optimize=True,
            )
            thumbnail_bytes.seek(0)

        image = WorksheetImage(thumbnail_bytes)
        image.width = 72
        image.height = 72
        worksheet.add_image(image, f"A{row_number}")
        worksheet.row_dimensions[row_number].height = 58
    except (OSError, ValueError, UnidentifiedImageError):
        # Missing or unsupported media must not prevent the data export.
        worksheet.cell(row_number, 1).value = "Image unavailable"
        return


def _category_workbook(presentation: PublicCatalogueResponse, category) -> BytesIO:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    products = [
        product for product in presentation.products
        if product.product_status == "active" and (
            product.category_name == category.name or category.name in product.categories
        )
    ]
    if not products:
        raise HTTPException(status_code=404, detail="This category has no downloadable products.")

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = re.sub(r"[\\/*?:\[\]]", "-", str(category.name)).strip("'")[:31] or "Products"
    headings = [
        "Image", "Product code", "Product name", "Brand", "Category",
        "Model", "Barcode", "Description", "Stock", "Price", "Currency",
    ]
    worksheet.append(headings)
    header_fill = PatternFill("solid", fgColor="147A46")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.row_dimensions[1].height = 24

    for row_number, product in enumerate(products, start=2):
        model = product.erp_details.get("model", "")
        worksheet.append([
            "",
            _excel_text(product.code),
            _excel_text(product.name),
            _excel_text(product.brand),
            _excel_text(category.name),
            _excel_text(model),
            _excel_text(product.barcode),
            _excel_text(product.description),
            product.stock_quantity if product.stock_quantity is not None else 0,
            float(product.price) if product.price is not None else None,
            _excel_text(product.currency),
        ])
        for cell in worksheet[row_number]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        worksheet.cell(row_number, 9).number_format = "0"
        worksheet.cell(row_number, 10).number_format = '#,##0.00'
        _add_public_product_image(worksheet, product, row_number)

    widths = [13, 18, 36, 18, 22, 18, 20, 48, 12, 14, 12]
    for column, width in enumerate(widths, start=1):
        worksheet.column_dimensions[worksheet.cell(1, column).column_letter].width = width
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


@router.get("/public/catalogues/{token}/categories/{category_slug}/excel")
def public_catalogue_category_excel(token: str, category_slug: str, request: Request,
    x_catalogue_password: str | None = Header(default=None), db: Session = Depends(get_db)):
    link, presentation = _public_link(
        db, token, x_catalogue_password,
        "public_catalogue_category_excel_downloaded", request,
    )
    category = next(
        (item for item in presentation.categories if item.slug.casefold() == category_slug.casefold()),
        None,
    )
    if not category:
        raise HTTPException(status_code=404, detail="Catalogue category not found.")
    output = _category_workbook(presentation, category)
    filename = f"{link.catalogue.slug}-{category.slug}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/public/catalogues/{token}/pdf")
def public_catalogue_pdf(token: str, request: Request, inline: bool = False, x_catalogue_password: str | None = Header(default=None), db: Session = Depends(get_db)):
    link, presentation = _public_link(db, token, x_catalogue_password, "public_catalogue_link_opened" if inline else "public_catalogue_pdf_downloaded", request)
    design = db.scalar(select(CatalogueDesign).where(
        CatalogueDesign.catalogue_id == link.catalogue_id,
        CatalogueDesign.status == "published",
        CatalogueDesign.deleted_at.is_(None),
    ))
    studio_version = db.scalar(select(CatalogueDesignVersion).where(
        CatalogueDesignVersion.design_id == design.id,
        CatalogueDesignVersion.version_number == design.current_version,
    )) if design else None
    if design and studio_version:
        active_promotion_ids = _active_promotion_ids(presentation)
        latest_stock_sync = db.scalar(select(func.max(Product.stock_last_synced_at)))
        stock_revision = int(_utc(latest_stock_sync).timestamp()) if latest_stock_sync else 0
        price_signature = hashlib.sha256(json.dumps(
            [
                [str(product.id), str(product.price or ""), str(product.currency or "")]
                for product in presentation.products
            ],
            separators=(",", ":"),
        ).encode()).hexdigest()[:12]
        promotion_signature = hashlib.sha256(
            "|".join(sorted(active_promotion_ids)).encode()
        ).hexdigest()[:10]
        output_path = cover_storage.resolve(
            f"catalogue-studio/public/{link.id}-{studio_version.id}-revision-{design.revision}-stock-{stock_revision}-price-{price_signature}-promotions-{promotion_signature}-layout-v6.pdf"
        )
        cache_hit = output_path.is_file()
        if not cache_hit:
            studio_snapshot = _current_public_studio_snapshot(
                db,
                design,
                presentation,
            )
            try:
                render_browser_pdf(
                    db,
                    studio_snapshot,
                    output_path,
                    {"include_cover": True, "quality": 92},
                    design_id=design.id,
                    version_id=studio_version.id,
                    requested_by_id=studio_version.created_by_id or design.updated_by_id or design.created_by_id,
                    public_token=token,
                    public_password=x_catalogue_password,
                )
            except RuntimeError:
                # A product image may have changed in live ERP data after the
                # immutable version was published. The native renderer resolves
                # the saved files directly and still preserves Studio geometry.
                render_studio_pdf(db, studio_snapshot, output_path, {"include_cover": True, "quality": 92})
    else:
        output_path, cache_hit = _cached_catalogue_pdf(presentation, paper_size="A4", orientation="landscape", include_cover=True, include_table_of_contents=False)
    filename = f"{link.catalogue.slug}.pdf"
    return FileResponse(output_path, media_type="application/pdf", filename=filename, content_disposition_type="inline" if inline else "attachment", headers={"X-Catalogue-PDF-Cache": "HIT" if cache_hit else "MISS", "X-Catalogue-Renderer": "studio" if design and studio_version else "standard"})
