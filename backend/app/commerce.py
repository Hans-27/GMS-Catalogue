import colorsys
import io
import hashlib
import json
import os
import re
import threading
import time
import uuid
from concurrent.futures import ThreadPoolExecutor
from copy import deepcopy
from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace
from typing import Literal
from xml.sax.saxutils import escape

from fastapi import APIRouter, Depends, HTTPException, Query, Request, status
from fastapi.responses import FileResponse, StreamingResponse
from sqlalchemy import func, or_, select, update
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session, noload, selectinload

from app.access import (
    allowed_price_list_ids,
    effective_permission_access,
    has_permission,
    has_record_access,
    is_superadmin,
    require_permission,
)
from app.branding import LOGO_ALT_TEXT, LOGO_VERSION, default_logo_path
from app.catalogue_pdf import CATALOGUE_PDF_CONFIG
from app.commerce_models import (
    Catalogue,
    CatalogueAudienceType,
    CatalogueProduct,
    CatalogueShareLink,
    CatalogueVersion,
    PriceChangeRequest,
    PriceList,
    ProductPrice,
    UserBrandCataloguePriceMapping,
    UserCataloguePriceMapping,
)
from app.commerce_schemas import (
    CatalogueCreate,
    CatalogueGenerationResponse,
    CatalogueBrandOption,
    CataloguePreviewPriceList,
    CataloguePreviewResponse,
    CatalogueProductsUpdate,
    CatalogueResponse,
    CatalogueUpdate,
    CatalogueVersionResponse,
    ErpCustomerPriceLevelResponse,
    ErpCustomerPriceFilters,
    ErpPriceBrandOption,
    ErpPriceCategoryOption,
    ErpProductPriceMatrixItem,
    ErpProductPriceMatrixPage,
    ErpProductPriceMatrixValue,
    ErpProductCustomerPricePage,
    ErpProductCustomerPriceResponse,
    PriceListCreate,
    PriceListResponse,
    PriceProposalCreate,
    PriceRequestResponse,
    PriceReview,
    ProductPriceResponse,
    UserCataloguePriceMappingResponse,
    UserCataloguePriceMappingsUpdate,
    UserBrandCataloguePriceMappingSummary,
)
from app.config import settings
from app.catalogue_auto_generation import generate_erp_brand_catalogues
from app.database import get_db
from app.design_studio_models import CatalogueDesign
from app.erp_models import ErpCustomerPriceLevel, ErpProductCustomerPrice
from app.models import AuditLog, Brand, CatalogueEntry, Category, Product, ProductImage, ProductVideo, User, product_categories
from app.product_content import (
    catalogue_long_description,
    catalogue_short_description,
    erp_description,
    erp_details,
    erp_summary,
)
from app.purchase_order_api import apply_purchase_order_quantities
from app.product_card_models import GlobalProductCard
from app.platform_models import (
    CatalogueCategorySetting,
    CatalogueCoverAsset,
    CatalogueCoverFile,
    CatalogueCoverSetting,
)
from app.storage import cover_storage, storage


_ERP_CARD_VARIANTS = ("rounded", "editorial", "framed", "contrast", "minimal")


def _theme_hex(hue: float, saturation: float, lightness: float) -> str:
    red, green, blue = colorsys.hls_to_rgb(hue, lightness, saturation)
    return f"#{round(red * 255):02X}{round(green * 255):02X}{round(blue * 255):02X}"


def _erp_product_card_theme(brand: str) -> dict[str, str]:
    """Create a stable, accessible visual identity for one ERP brand."""

    normalized_brand = "".join(character for character in brand.casefold() if character.isalnum())
    if "glink" in normalized_brand:
        return {
            "key": "brand-glink",
            "variant": "contrast",
            "accent_color": "#8DE2DC",
            "strong_color": "#267A74",
            "surface_color": "#267A74",
            "border_color": "#164F4B",
            "text_color": "#164F4B",
        }
    if normalized_brand == "ega" or normalized_brand.startswith("egacatalogue"):
        return {
            "key": "brand-ega",
            "variant": "framed",
            "accent_color": "#65F58A",
            "strong_color": "#087A3D",
            "surface_color": "#EAFFEF",
            "border_color": "#159447",
            "text_color": "#073B22",
        }
    if normalized_brand == "nubwo" or normalized_brand.startswith("nubwocatalogue"):
        return {
            "key": "brand-nubwo",
            "variant": "framed",
            "accent_color": "#FFF05A",
            "strong_color": "#7A6400",
            "surface_color": "#FFFCE5",
            "border_color": "#C8A900",
            "text_color": "#332A00",
        }

    digest = hashlib.sha256(brand.strip().casefold().encode("utf-8")).digest()
    hue = int.from_bytes(digest[:2], "big") % 360 / 360
    variant = _ERP_CARD_VARIANTS[digest[2] % len(_ERP_CARD_VARIANTS)]
    return {
        "key": f"brand-{digest.hex()[:8]}",
        "variant": variant,
        "accent_color": _theme_hex(hue, .68, .74),
        "strong_color": _theme_hex(hue, .54, .29),
        "surface_color": _theme_hex(hue, .55, .965),
        "border_color": _theme_hex(hue, .38, .79),
        "text_color": _theme_hex(hue, .42, .20),
    }


router = APIRouter(tags=["Price and Catalogue Management"])

_PDF_CACHE_SCHEMA_VERSION = 11
_PDF_CACHE_LOCKS: dict[str, threading.Lock] = {}
_PDF_CACHE_LOCKS_GUARD = threading.Lock()
_PDF_CACHE_PRUNE_LOCK = threading.Lock()
_PDF_CACHE_LAST_PRUNE = 0.0


def _now() -> datetime:
    return datetime.now(UTC)


def _as_utc(value: datetime) -> datetime:
    return value.replace(tzinfo=UTC) if value.tzinfo is None else value.astimezone(UTC)


def _audit(
    db: Session,
    request: Request,
    actor: User,
    *,
    action: str,
    module: str,
    identifier: str | None = None,
    details: dict | None = None,
) -> None:
    forwarded_for = request.headers.get("x-forwarded-for")
    ip_address = (
        forwarded_for.split(",", maxsplit=1)[0].strip()
        if forwarded_for
        else (request.client.host if request.client else None)
    )
    db.add(
        AuditLog(
            user_id=actor.id,
            action=action,
            module=module,
            status="success",
            identifier=identifier,
            ip_address=ip_address,
            user_agent=request.headers.get("user-agent"),
            details=details,
        )
    )


def _slugify(value: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", value.casefold()).strip("-")
    return slug or f"catalogue-{uuid.uuid4().hex[:8]}"


def _get_product(db: Session, product_id: uuid.UUID) -> Product:
    product = db.scalar(
        select(Product)
        .options(
            selectinload(Product.catalogue_entry),
            selectinload(Product.categories),
            selectinload(Product.images),
            selectinload(Product.videos),
        )
        .where(Product.id == product_id)
    )
    if not product:
        raise HTTPException(status_code=404, detail="Product not found.")
    return product


def _get_price_list(db: Session, price_list_id: int) -> PriceList:
    price_list = db.get(PriceList, price_list_id)
    if not price_list or not price_list.is_active:
        raise HTTPException(status_code=404, detail="Price list not found.")
    return price_list


def _require_price_list_access(
    db: Session,
    actor: User,
    price_list_id: int | None,
) -> None:
    if price_list_id is None:
        return
    allowed = allowed_price_list_ids(db, actor)
    if allowed is not None and price_list_id not in allowed:
        raise HTTPException(status_code=404, detail="Price list not found.")


def _catalogue_query(*, include_product_details: bool = False):
    options = [
        selectinload(Catalogue.price_list),
        selectinload(Catalogue.product_links),
    ]
    if include_product_details:
        options.append(
            selectinload(Catalogue.product_links)
            .selectinload(CatalogueProduct.product)
            .options(
                selectinload(Product.catalogue_entry),
                selectinload(Product.categories),
                selectinload(Product.images),
                selectinload(Product.videos),
            )
        )
    return select(Catalogue).options(*options)


def _get_catalogue(db: Session, catalogue_id: uuid.UUID) -> Catalogue:
    catalogue = db.scalar(
        _catalogue_query(include_product_details=True).where(
            Catalogue.id == catalogue_id
        )
    )
    if not catalogue:
        raise HTTPException(status_code=404, detail="Catalogue not found.")
    return catalogue


def _get_scoped_catalogue(
    db: Session,
    catalogue_id: uuid.UUID,
    actor: User,
    permission_code: str,
) -> Catalogue:
    catalogue = _get_catalogue(db, catalogue_id)
    if not has_record_access(db, actor, permission_code, catalogue):
        raise HTTPException(status_code=404, detail="Catalogue not found.")
    return catalogue


def _current_price(
    db: Session,
    product_id: uuid.UUID,
    price_list_id: int,
    *,
    as_of: datetime | None = None,
) -> ProductPrice | None:
    effective_at = as_of or _now()
    return db.scalar(
        select(ProductPrice)
        .where(
            ProductPrice.product_id == product_id,
            ProductPrice.price_list_id == price_list_id,
            ProductPrice.status == "active",
            ProductPrice.effective_from <= effective_at,
            or_(
                ProductPrice.expires_at.is_(None),
                ProductPrice.expires_at > effective_at,
            ),
        )
        .order_by(ProductPrice.effective_from.desc(), ProductPrice.created_at.desc())
    )


def _insert_price(
    db: Session,
    *,
    product: Product,
    price_list: PriceList,
    amount: Decimal,
    effective_from: datetime,
    reason: str,
    actor: User,
) -> ProductPrice:
    previous = _current_price(
        db,
        product.id,
        price_list.id,
        as_of=effective_from,
    )
    if previous and _as_utc(previous.effective_from) >= _as_utc(effective_from):
        raise HTTPException(
            status_code=409,
            detail="A price already begins at or after that effective time.",
        )

    next_price = db.scalar(
        select(ProductPrice)
        .where(
            ProductPrice.product_id == product.id,
            ProductPrice.price_list_id == price_list.id,
            ProductPrice.status == "active",
            ProductPrice.effective_from > effective_from,
        )
        .order_by(ProductPrice.effective_from)
    )
    if previous:
        previous.expires_at = effective_from

    price = ProductPrice(
        product_id=product.id,
        price_list_id=price_list.id,
        currency=price_list.currency,
        amount=amount,
        effective_from=effective_from,
        expires_at=next_price.effective_from if next_price else None,
        status="active",
        reason=reason,
        created_by_id=actor.id,
        approved_by_id=actor.id,
        approved_at=_now(),
    )
    db.add(price)
    db.flush()
    return price


def _price_response(
    db: Session,
    price: ProductPrice,
) -> ProductPriceResponse:
    product = _get_product(db, price.product_id)
    price_list = db.get(PriceList, price.price_list_id)
    return ProductPriceResponse(
        id=price.id,
        product_id=price.product_id,
        product_sku=product.sku,
        product_name=(
            product.catalogue_entry.display_name
            if product.catalogue_entry and product.catalogue_entry.display_name
            else product.erp_name
        ),
        price_list_id=price.price_list_id,
        price_list_name=price_list.name if price_list else "Unavailable",
        currency=price.currency,
        amount=price.amount,
        effective_from=price.effective_from,
        expires_at=price.expires_at,
        status=price.status,
        reason=price.reason,
        created_by_id=price.created_by_id,
        approved_by_id=price.approved_by_id,
        created_at=price.created_at,
        approved_at=price.approved_at,
    )


def _request_response(
    db: Session,
    change_request: PriceChangeRequest,
) -> PriceRequestResponse:
    product = _get_product(db, change_request.product_id)
    price_list = db.get(PriceList, change_request.price_list_id)
    current = _current_price(db, product.id, change_request.price_list_id)
    return PriceRequestResponse(
        id=change_request.id,
        product_id=product.id,
        product_sku=product.sku,
        product_name=(
            product.catalogue_entry.display_name
            if product.catalogue_entry and product.catalogue_entry.display_name
            else product.erp_name
        ),
        price_list_id=change_request.price_list_id,
        price_list_name=price_list.name if price_list else "Unavailable",
        current_amount=current.amount if current else None,
        proposed_amount=change_request.proposed_amount,
        currency=change_request.currency,
        effective_from=change_request.effective_from,
        reason=change_request.reason,
        status=change_request.status,
        requested_by_id=change_request.requested_by_id,
        reviewed_by_id=change_request.reviewed_by_id,
        review_reason=change_request.review_reason,
        resulting_price_id=change_request.resulting_price_id,
        created_at=change_request.created_at,
        reviewed_at=change_request.reviewed_at,
    )


@router.get("/price-lists", response_model=list[PriceListResponse])
def list_price_lists(
    include_inactive: bool = False,
    actor: User = Depends(require_permission("prices.view")),
    db: Session = Depends(get_db),
) -> list[PriceListResponse]:
    query = select(PriceList)
    allowed = allowed_price_list_ids(db, actor)
    if allowed is not None:
        query = query.where(PriceList.id.in_(allowed))
    if not include_inactive:
        query = query.where(PriceList.is_active.is_(True))
    return [
        PriceListResponse.model_validate(item, from_attributes=True)
        for item in db.scalars(query.order_by(PriceList.name))
    ]


@router.get(
    "/price-levels",
    response_model=list[ErpCustomerPriceLevelResponse],
)
def list_erp_customer_price_levels(
    actor: User = Depends(require_permission("prices.view")),
    db: Session = Depends(get_db),
) -> list[ErpCustomerPriceLevelResponse]:
    allowed = allowed_price_list_ids(db, actor)
    query = (
        select(ErpCustomerPriceLevel, PriceList)
        .join(PriceList, PriceList.id == ErpCustomerPriceLevel.price_list_id)
        .order_by(ErpCustomerPriceLevel.sort_order, ErpCustomerPriceLevel.source_code)
    )
    if allowed is not None:
        query = query.where(ErpCustomerPriceLevel.price_list_id.in_(allowed))
    return [
        ErpCustomerPriceLevelResponse(
            id=level.id,
            erp_price_type_id=level.erp_price_type_id,
            source_code=level.source_code,
            source_name=level.source_name,
            price_list_id=price_list.id,
            price_list_code=price_list.code,
            price_list_name=price_list.name,
            product_count=level.product_count,
            sort_order=level.sort_order,
            is_active=level.is_active,
            last_synced_at=level.last_synced_at,
        )
        for level, price_list in db.execute(query)
    ]


def _user_catalogue_price_mappings(
    db: Session,
    actor: User,
    brand: str | None = None,
) -> list[UserCataloguePriceMappingResponse]:
    audiences = list(
        db.scalars(
            select(CatalogueAudienceType)
            .options(selectinload(CatalogueAudienceType.price_list))
            .where(CatalogueAudienceType.is_active.is_(True))
            .order_by(CatalogueAudienceType.display_order, CatalogueAudienceType.id)
        )
    )
    custom = {
        item.audience_type_id: item
        for item in db.scalars(
            select(UserCataloguePriceMapping)
            .options(
                selectinload(UserCataloguePriceMapping.price_list),
                selectinload(UserCataloguePriceMapping.audience_type),
            )
            .where(UserCataloguePriceMapping.user_id == actor.id)
        )
    }
    brand_name = brand.strip() if brand else None
    brand_key = brand_name.casefold() if brand_name else None
    brand_custom = {
        item.audience_type_id: item
        for item in db.scalars(
            select(UserBrandCataloguePriceMapping)
            .options(
                selectinload(UserBrandCataloguePriceMapping.price_list),
                selectinload(UserBrandCataloguePriceMapping.audience_type),
            )
            .where(
                UserBrandCataloguePriceMapping.user_id == actor.id,
                UserBrandCataloguePriceMapping.brand_key == brand_key,
            )
        )
    } if brand_key else {}
    allowed = allowed_price_list_ids(db, actor)
    level_rows = list(
        db.scalars(
            select(ErpCustomerPriceLevel)
            .where(ErpCustomerPriceLevel.is_active.is_(True))
            .order_by(ErpCustomerPriceLevel.sort_order, ErpCustomerPriceLevel.id)
        )
    )
    levels = {level.price_list_id: level for level in level_rows}
    first_allowed_price_list = next(
        (
            db.get(PriceList, level.price_list_id)
            for level in level_rows
            if allowed is None or level.price_list_id in allowed
        ),
        None,
    )
    fallback_no_price = db.scalar(
        select(PriceList).where(
            PriceList.is_active.is_(True),
            PriceList.is_no_price.is_(True),
        )
    )
    result: list[UserCataloguePriceMappingResponse] = []
    for audience in audiences:
        mapping = brand_custom.get(audience.id) or custom.get(audience.id)
        price_list = mapping.price_list if mapping else audience.price_list
        if (
            price_list is not None
            and not price_list.is_no_price
            and allowed is not None
            and price_list.id not in allowed
        ):
            price_list = first_allowed_price_list
        if price_list is None:
            price_list = fallback_no_price
        if price_list is None:
            continue
        level = levels.get(price_list.id)
        result.append(
            UserCataloguePriceMappingResponse(
                audience_type_id=audience.id,
                audience_code=audience.code,
                audience_name=audience.display_name,
                price_list_id=price_list.id,
                price_list_code=price_list.code,
                price_list_name=price_list.name,
                erp_source_code=level.source_code if level else None,
                erp_source_name=level.source_name if level else None,
                show_prices=audience.show_prices and not price_list.is_no_price,
                is_custom=mapping is not None,
                brand=brand_name,
                is_brand_override=audience.id in brand_custom,
            )
        )
    return result


@router.get(
    "/my-catalogue-price-mappings",
    response_model=list[UserCataloguePriceMappingResponse],
)
def get_my_catalogue_price_mappings(
    brand: str | None = Query(default=None, max_length=120),
    actor: User = Depends(require_permission("prices.view")),
    db: Session = Depends(get_db),
) -> list[UserCataloguePriceMappingResponse]:
    return _user_catalogue_price_mappings(db, actor, brand)


@router.put(
    "/my-catalogue-price-mappings",
    response_model=list[UserCataloguePriceMappingResponse],
)
def update_my_catalogue_price_mappings(
    payload: UserCataloguePriceMappingsUpdate,
    request: Request,
    actor: User = Depends(require_permission("prices.view")),
    db: Session = Depends(get_db),
) -> list[UserCataloguePriceMappingResponse]:
    brand_name = payload.brand.strip() if payload.brand else None
    brand_key = brand_name.casefold() if brand_name else None
    if payload.brand is not None and not brand_name:
        raise HTTPException(status_code=422, detail="Choose a valid brand.")
    audience_ids = [item.audience_type_id for item in payload.mappings]
    if len(audience_ids) != len(set(audience_ids)):
        raise HTTPException(status_code=422, detail="Each customer level can be configured only once.")
    audiences = {
        item.id: item
        for item in db.scalars(
            select(CatalogueAudienceType).where(
                CatalogueAudienceType.id.in_(audience_ids),
                CatalogueAudienceType.is_active.is_(True),
            )
        )
    }
    if len(audiences) != len(audience_ids):
        raise HTTPException(status_code=422, detail="A selected customer level is unavailable.")
    price_list_ids = {item.price_list_id for item in payload.mappings}
    price_lists = {
        item.id: item
        for item in db.scalars(
            select(PriceList).where(
                PriceList.id.in_(price_list_ids),
                PriceList.is_active.is_(True),
            )
        )
    }
    if len(price_lists) != len(price_list_ids):
        raise HTTPException(status_code=422, detail="A selected ERP price list is unavailable.")
    erp_price_list_ids = set(
        db.scalars(
            select(ErpCustomerPriceLevel.price_list_id).where(
                ErpCustomerPriceLevel.is_active.is_(True)
            )
        )
    )
    if any(
        price_list_id not in erp_price_list_ids and not price_lists[price_list_id].is_no_price
        for price_list_id in price_list_ids
    ):
        raise HTTPException(status_code=422, detail="Choose a synchronized ERP price level or No Price.")
    allowed = allowed_price_list_ids(db, actor)
    if allowed is not None and any(
        price_list_id not in allowed and not price_lists[price_list_id].is_no_price
        for price_list_id in price_list_ids
    ):
        raise HTTPException(status_code=403, detail="You cannot use one or more selected ERP price levels.")

    existing = {
        item.audience_type_id: item
        for item in db.scalars(
            select(UserBrandCataloguePriceMapping if brand_key else UserCataloguePriceMapping).where(
                (UserBrandCataloguePriceMapping.user_id if brand_key else UserCataloguePriceMapping.user_id) == actor.id,
                (UserBrandCataloguePriceMapping.audience_type_id if brand_key else UserCataloguePriceMapping.audience_type_id).in_(audience_ids),
                *((UserBrandCataloguePriceMapping.brand_key == brand_key,) if brand_key else ()),
            )
        )
    }
    for requested in payload.mappings:
        mapping = existing.get(requested.audience_type_id)
        if mapping:
            mapping.price_list_id = requested.price_list_id
        else:
            if brand_key:
                db.add(
                    UserBrandCataloguePriceMapping(
                        user_id=actor.id,
                        brand_key=brand_key,
                        brand_name=brand_name,
                        audience_type_id=requested.audience_type_id,
                        price_list_id=requested.price_list_id,
                    )
                )
            else:
                db.add(UserCataloguePriceMapping(
                    user_id=actor.id,
                    audience_type_id=requested.audience_type_id,
                    price_list_id=requested.price_list_id,
                ))
        price_list = price_lists[requested.price_list_id]
        db.execute(
            update(CatalogueShareLink)
            .where(
                CatalogueShareLink.created_by_id == actor.id,
                CatalogueShareLink.audience_type_id == requested.audience_type_id,
                *((CatalogueShareLink.catalogue_id.in_(
                    select(Catalogue.id).where(func.lower(Catalogue.brand) == brand_key)
                ),) if brand_key else ()),
                *((CatalogueShareLink.catalogue_id.not_in(
                    select(Catalogue.id)
                    .join(
                        UserBrandCataloguePriceMapping,
                        (UserBrandCataloguePriceMapping.user_id == actor.id)
                        & (UserBrandCataloguePriceMapping.brand_key == func.lower(Catalogue.brand))
                        & (UserBrandCataloguePriceMapping.audience_type_id == requested.audience_type_id),
                    )
                ),) if not brand_key else ()),
            )
            .values(
                price_list_id=price_list.id,
                show_prices=audiences[requested.audience_type_id].show_prices
                and not price_list.is_no_price,
            )
        )
    _audit(
        db,
        request,
        actor,
        action="catalogue_price_mapping_updated",
        module="prices",
        identifier=actor.username,
        details={
            "mappings": [
                {
                    "audience_type_id": item.audience_type_id,
                    "price_list_id": item.price_list_id,
                }
                for item in payload.mappings
            ],
            "brand": brand_name,
        },
    )
    db.commit()
    return _user_catalogue_price_mappings(db, actor, brand_name)


@router.get(
    "/erp-customer-prices",
    response_model=ErpProductCustomerPricePage,
)
def list_erp_product_customer_prices(
    q: str = Query(default="", max_length=120),
    price_level_id: int | None = None,
    brand: str | None = Query(default=None, max_length=120),
    category_id: int | None = None,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=100, ge=10, le=200),
    actor: User = Depends(require_permission("prices.view")),
    db: Session = Depends(get_db),
) -> ErpProductCustomerPricePage:
    query = (
        select(ErpProductCustomerPrice, Product, ErpCustomerPriceLevel, PriceList)
        .join(Product, Product.id == ErpProductCustomerPrice.product_id)
        .join(
            ErpCustomerPriceLevel,
            ErpCustomerPriceLevel.id == ErpProductCustomerPrice.price_level_id,
        )
        .join(PriceList, PriceList.id == ErpCustomerPriceLevel.price_list_id)
    )
    allowed = allowed_price_list_ids(db, actor)
    if allowed is not None:
        query = query.where(ErpCustomerPriceLevel.price_list_id.in_(allowed))
    if price_level_id:
        query = query.where(ErpProductCustomerPrice.price_level_id == price_level_id)
    if brand and brand.strip():
        query = query.where(func.lower(Product.brand) == brand.strip().casefold())
    if category_id is not None:
        query = query.where(Product.categories.any(Category.id == category_id))
    if q.strip():
        pattern = f"%{q.strip()}%"
        query = query.where(
            or_(
                Product.sku.ilike(pattern),
                Product.erp_name.ilike(pattern),
                Product.brand.ilike(pattern),
            )
        )
    total = db.scalar(select(func.count()).select_from(query.subquery())) or 0
    rows = db.execute(
        query.order_by(Product.sku, ErpCustomerPriceLevel.sort_order)
        .offset((page - 1) * page_size)
        .limit(page_size)
    ).all()
    return ErpProductCustomerPricePage(
        items=[
            ErpProductCustomerPriceResponse(
                id=price.id,
                product_id=product.id,
                product_sku=product.sku,
                product_name=product.erp_name,
                product_brand=product.brand,
                category_names=sorted(category.name for category in product.categories),
                price_level_id=level.id,
                source_code=level.source_code,
                source_name=level.source_name,
                price_list_code=price_list.code,
                price_list_name=price_list.name,
                amount=price.amount,
                currency=price.currency,
                source_updated_at=price.source_updated_at,
                last_synced_at=price.last_synced_at,
            )
            for price, product, level, price_list in rows
        ],
        total=total,
        page=page,
        page_size=page_size,
        pages=max(1, (total + page_size - 1) // page_size),
    )


@router.get(
    "/erp-customer-price-filters",
    response_model=ErpCustomerPriceFilters,
)
def get_erp_customer_price_filters(
    actor: User = Depends(require_permission("prices.view")),
    db: Session = Depends(get_db),
) -> ErpCustomerPriceFilters:
    allowed = allowed_price_list_ids(db, actor)
    priced_products = (
        select(ErpProductCustomerPrice.product_id)
        .join(
            ErpCustomerPriceLevel,
            ErpCustomerPriceLevel.id == ErpProductCustomerPrice.price_level_id,
        )
        .where(ErpCustomerPriceLevel.is_active.is_(True))
    )
    if allowed is not None:
        priced_products = priced_products.where(
            ErpCustomerPriceLevel.price_list_id.in_(allowed)
        )
    priced_products = priced_products.distinct().subquery()

    brand_rows = db.execute(
        select(Product.brand, func.count(Product.id))
        .join(priced_products, priced_products.c.product_id == Product.id)
        .where(Product.brand.is_not(None), Product.brand != "")
        .group_by(Product.brand)
        .order_by(func.lower(Product.brand))
    ).all()
    category_rows = db.execute(
        select(Category.id, Category.name, func.count(func.distinct(Product.id)))
        .join(product_categories, product_categories.c.category_id == Category.id)
        .join(Product, Product.id == product_categories.c.product_id)
        .join(priced_products, priced_products.c.product_id == Product.id)
        .where(Category.is_active.is_(True))
        .group_by(Category.id, Category.name)
        .order_by(func.lower(Category.name))
    ).all()
    return ErpCustomerPriceFilters(
        brands=[
            ErpPriceBrandOption(name=name, product_count=int(count))
            for name, count in brand_rows
            if name and name.strip()
        ],
        categories=[
            ErpPriceCategoryOption(id=category_id, name=name, product_count=int(count))
            for category_id, name, count in category_rows
        ],
    )


@router.get(
    "/my-catalogue-price-mapping-summary",
    response_model=list[UserBrandCataloguePriceMappingSummary],
)
def get_my_catalogue_price_mapping_summary(
    actor: User = Depends(require_permission("prices.view")),
    db: Session = Depends(get_db),
) -> list[UserBrandCataloguePriceMappingSummary]:
    """Show which ERP brands use explicit mappings and which inherit defaults."""
    brand_rows = db.execute(
        select(Product.brand, func.count(Product.id))
        .where(
            Product.status == "active",
            Product.brand.is_not(None),
            Product.brand != "",
        )
        .group_by(Product.brand)
        .order_by(func.lower(Product.brand))
    ).all()
    audience_count = db.scalar(
        select(func.count(CatalogueAudienceType.id)).where(
            CatalogueAudienceType.is_active.is_(True)
        )
    ) or 0
    override_counts = {
        brand_key: int(count)
        for brand_key, count in db.execute(
            select(
                UserBrandCataloguePriceMapping.brand_key,
                func.count(func.distinct(UserBrandCataloguePriceMapping.audience_type_id)),
            )
            .join(
                CatalogueAudienceType,
                CatalogueAudienceType.id == UserBrandCataloguePriceMapping.audience_type_id,
            )
            .where(
                UserBrandCataloguePriceMapping.user_id == actor.id,
                CatalogueAudienceType.is_active.is_(True),
            )
            .group_by(UserBrandCataloguePriceMapping.brand_key)
        )
    }
    result: list[UserBrandCataloguePriceMappingSummary] = []
    for brand_name, product_count in brand_rows:
        override_count = override_counts.get(brand_name.strip().casefold(), 0)
        status: Literal["custom", "partial", "default"] = (
            "custom"
            if audience_count > 0 and override_count >= audience_count
            else "partial"
            if override_count > 0
            else "default"
        )
        result.append(UserBrandCataloguePriceMappingSummary(
            brand=brand_name,
            product_count=int(product_count),
            override_count=override_count,
            audience_count=int(audience_count),
            status=status,
        ))
    return result


@router.get(
    "/erp-product-price-matrix",
    response_model=ErpProductPriceMatrixPage,
)
def get_erp_product_price_matrix(
    q: str = Query(default="", max_length=120),
    brand: str | None = Query(default=None, max_length=120),
    category_id: int | None = None,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=50, ge=10, le=100),
    actor: User = Depends(require_permission("prices.view")),
    db: Session = Depends(get_db),
) -> ErpProductPriceMatrixPage:
    allowed = allowed_price_list_ids(db, actor)
    product_query = (
        # PostgreSQL requires ORDER BY columns to be selected when DISTINCT is
        # used.  Keep SKU in this projection so pagination can remain stable
        # while scalars() continues to return the product id (the first column).
        select(Product.id, Product.sku)
        .join(
            ErpProductCustomerPrice,
            ErpProductCustomerPrice.product_id == Product.id,
        )
        .join(
            ErpCustomerPriceLevel,
            ErpCustomerPriceLevel.id == ErpProductCustomerPrice.price_level_id,
        )
        .where(ErpCustomerPriceLevel.is_active.is_(True))
    )
    if allowed is not None:
        product_query = product_query.where(
            ErpCustomerPriceLevel.price_list_id.in_(allowed)
        )
    if brand and brand.strip():
        product_query = product_query.where(
            func.lower(Product.brand) == brand.strip().casefold()
        )
    if category_id is not None:
        product_query = product_query.where(
            Product.categories.any(Category.id == category_id)
        )
    if q.strip():
        pattern = f"%{q.strip()}%"
        product_query = product_query.where(
            or_(
                Product.sku.ilike(pattern),
                Product.erp_name.ilike(pattern),
                Product.brand.ilike(pattern),
            )
        )

    product_query = product_query.distinct()
    total = db.scalar(select(func.count()).select_from(product_query.subquery())) or 0
    product_ids = list(
        db.scalars(
            product_query.order_by(Product.sku)
            .offset((page - 1) * page_size)
            .limit(page_size)
        )
    )
    products = {
        product.id: product
        for product in db.scalars(
            select(Product)
            .options(selectinload(Product.categories))
            .where(Product.id.in_(product_ids))
        ).unique()
    } if product_ids else {}

    price_query = (
        select(ErpProductCustomerPrice, ErpCustomerPriceLevel, PriceList)
        .join(
            ErpCustomerPriceLevel,
            ErpCustomerPriceLevel.id == ErpProductCustomerPrice.price_level_id,
        )
        .join(PriceList, PriceList.id == ErpCustomerPriceLevel.price_list_id)
        .where(
            ErpProductCustomerPrice.product_id.in_(product_ids),
            ErpCustomerPriceLevel.is_active.is_(True),
        )
        .order_by(
            ErpProductCustomerPrice.product_id,
            ErpCustomerPriceLevel.sort_order,
        )
    ) if product_ids else None
    if price_query is not None and allowed is not None:
        price_query = price_query.where(
            ErpCustomerPriceLevel.price_list_id.in_(allowed)
        )
    prices_by_product: dict[uuid.UUID, list[ErpProductPriceMatrixValue]] = {
        product_id: [] for product_id in product_ids
    }
    if price_query is not None:
        for price, level, price_list in db.execute(price_query):
            prices_by_product[price.product_id].append(
                ErpProductPriceMatrixValue(
                    price_level_id=level.id,
                    source_code=level.source_code,
                    source_name=level.source_name,
                    price_list_code=price_list.code,
                    price_list_name=price_list.name,
                    amount=price.amount,
                    currency=price.currency,
                    last_synced_at=price.last_synced_at,
                )
            )

    return ErpProductPriceMatrixPage(
        items=[
            ErpProductPriceMatrixItem(
                product_id=product.id,
                product_sku=product.sku,
                product_name=product.erp_name,
                product_brand=product.brand,
                category_names=sorted(category.name for category in product.categories),
                prices=prices_by_product.get(product.id, []),
            )
            for product_id in product_ids
            if (product := products.get(product_id)) is not None
        ],
        total=total,
        page=page,
        page_size=page_size,
        pages=max(1, (total + page_size - 1) // page_size),
    )


@router.post(
    "/price-lists",
    response_model=PriceListResponse,
    status_code=status.HTTP_201_CREATED,
)
def create_price_list(
    payload: PriceListCreate,
    request: Request,
    actor: User = Depends(require_permission("settings.manage")),
    db: Session = Depends(get_db),
) -> PriceListResponse:
    price_list = PriceList(**payload.model_dump())
    db.add(price_list)
    _audit(
        db,
        request,
        actor,
        action="price_list_created",
        module="pricing",
        identifier=payload.code,
        details={"name": payload.name, "is_no_price": payload.is_no_price},
    )
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="Price-list code or name exists.")
    return PriceListResponse.model_validate(price_list, from_attributes=True)


@router.get("/product-prices", response_model=list[ProductPriceResponse])
def list_product_prices(
    product_id: uuid.UUID | None = None,
    price_list_id: int | None = None,
    actor: User = Depends(require_permission("prices.view")),
    db: Session = Depends(get_db),
) -> list[ProductPriceResponse]:
    query = select(ProductPrice)
    allowed = allowed_price_list_ids(db, actor)
    if allowed is not None:
        query = query.where(ProductPrice.price_list_id.in_(allowed))
    if product_id:
        product = _get_product(db, product_id)
        if not has_record_access(db, actor, "prices.view", product):
            raise HTTPException(status_code=404, detail="Product not found.")
        query = query.where(ProductPrice.product_id == product_id)
    if price_list_id:
        _require_price_list_access(db, actor, price_list_id)
        query = query.where(ProductPrice.price_list_id == price_list_id)
    prices = db.scalars(
        query.order_by(ProductPrice.created_at.desc()).limit(500)
    ).all()
    return [_price_response(db, price) for price in prices]


@router.post(
    "/product-prices",
    response_model=ProductPriceResponse,
    status_code=status.HTTP_201_CREATED,
)
def create_product_price(
    payload: PriceProposalCreate,
    request: Request,
    actor: User = Depends(require_permission("prices.edit")),
    db: Session = Depends(get_db),
) -> ProductPriceResponse:
    product = _get_product(db, payload.product_id)
    if not has_record_access(db, actor, "prices.edit", product):
        raise HTTPException(status_code=404, detail="Product not found.")
    price_list = _get_price_list(db, payload.price_list_id)
    _require_price_list_access(db, actor, price_list.id)
    if price_list.is_no_price:
        raise HTTPException(status_code=422, detail="No Price lists cannot hold prices.")
    price = _insert_price(
        db,
        product=product,
        price_list=price_list,
        amount=payload.proposed_amount,
        effective_from=payload.effective_from,
        reason=payload.reason,
        actor=actor,
    )
    _audit(
        db,
        request,
        actor,
        action="product_price_created",
        module="pricing",
        identifier=product.sku,
        details={
            "price_list": price_list.code,
            "amount": str(price.amount),
            "reason": price.reason,
        },
    )
    db.commit()
    return _price_response(db, price)


@router.get("/price-change-requests", response_model=list[PriceRequestResponse])
def list_price_requests(
    request_status: str | None = Query(default=None, alias="status"),
    actor: User = Depends(require_permission("prices.view")),
    db: Session = Depends(get_db),
) -> list[PriceRequestResponse]:
    query = select(PriceChangeRequest)
    allowed = allowed_price_list_ids(db, actor)
    if allowed is not None:
        query = query.where(PriceChangeRequest.price_list_id.in_(allowed))
    if request_status:
        query = query.where(PriceChangeRequest.status == request_status)
    requests = db.scalars(
        query.order_by(PriceChangeRequest.created_at.desc()).limit(500)
    ).all()
    return [_request_response(db, item) for item in requests]


@router.post(
    "/price-change-requests",
    response_model=PriceRequestResponse,
    status_code=status.HTTP_201_CREATED,
)
def propose_price_change(
    payload: PriceProposalCreate,
    request: Request,
    actor: User = Depends(require_permission("prices.propose")),
    db: Session = Depends(get_db),
) -> PriceRequestResponse:
    product = _get_product(db, payload.product_id)
    if not has_record_access(db, actor, "prices.propose", product):
        raise HTTPException(status_code=404, detail="Product not found.")
    price_list = _get_price_list(db, payload.price_list_id)
    _require_price_list_access(db, actor, price_list.id)
    if price_list.is_no_price:
        raise HTTPException(status_code=422, detail="No Price lists cannot hold prices.")
    change_request = PriceChangeRequest(
        product_id=product.id,
        price_list_id=price_list.id,
        proposed_amount=payload.proposed_amount,
        currency=price_list.currency,
        effective_from=payload.effective_from,
        reason=payload.reason,
        requested_by_id=actor.id,
    )
    db.add(change_request)
    _audit(
        db,
        request,
        actor,
        action="price_change_proposed",
        module="pricing",
        identifier=product.sku,
        details={
            "price_list": price_list.code,
            "proposed_amount": str(payload.proposed_amount),
            "reason": payload.reason,
        },
    )
    db.commit()
    return _request_response(db, change_request)


@router.post(
    "/price-change-requests/{request_id}/review",
    response_model=PriceRequestResponse,
)
def review_price_change(
    request_id: uuid.UUID,
    payload: PriceReview,
    request: Request,
    actor: User = Depends(require_permission("prices.approve")),
    db: Session = Depends(get_db),
) -> PriceRequestResponse:
    change_request = db.get(PriceChangeRequest, request_id)
    if not change_request:
        raise HTTPException(status_code=404, detail="Price request not found.")
    _require_price_list_access(db, actor, change_request.price_list_id)
    request_product = _get_product(db, change_request.product_id)
    if not has_record_access(db, actor, "prices.approve", request_product):
        raise HTTPException(status_code=404, detail="Price request not found.")
    if change_request.status != "pending":
        raise HTTPException(status_code=409, detail="This request was already reviewed.")
    change_request.reviewed_by_id = actor.id
    change_request.reviewed_at = _now()
    change_request.review_reason = payload.reason
    if payload.approve:
        product = _get_product(db, change_request.product_id)
        price_list = _get_price_list(db, change_request.price_list_id)
        price = _insert_price(
            db,
            product=product,
            price_list=price_list,
            amount=change_request.proposed_amount,
            effective_from=change_request.effective_from,
            reason=change_request.reason,
            actor=actor,
        )
        change_request.status = "approved"
        change_request.resulting_price_id = price.id
        action = "price_change_approved"
    else:
        change_request.status = "rejected"
        action = "price_change_rejected"
    _audit(
        db,
        request,
        actor,
        action=action,
        module="pricing",
        identifier=str(change_request.id),
        details={"review_reason": payload.reason},
    )
    db.commit()
    return _request_response(db, change_request)


def _catalogue_product_response(db: Session, link: CatalogueProduct) -> dict:
    product = link.product
    entry = product.catalogue_entry or CatalogueEntry()
    primary_image = next(
        (image for image in product.images if image.is_primary),
        product.images[0] if product.images else None,
    )
    return {
        "product_id": product.id,
        "section_title": link.section_title,
        "override_description": link.override_description,
        "hide_price": link.hide_price,
        "include_video": link.include_video,
        "selected_video_id": link.selected_video_id,
        "video_title_override": link.video_title_override,
        "video_description_override": link.video_description_override,
        "video_display_mode": link.video_display_mode,
        "video_thumbnail_mode": link.video_thumbnail_mode,
        "sort_order": link.sort_order,
        "sku": product.sku,
        "name": entry.display_name or product.erp_name,
        "brand": product.brand,
        "primary_image_url": primary_image.public_url if primary_image else None,
        "product_status": product.status,
    }


def _catalogue_response(
    db: Session,
    catalogue: Catalogue,
    *,
    include_products: bool = True,
    actor: User | None = None,
    product_count: int | None = None,
    studio_design: CatalogueDesign | None = None,
    brand_logo_url: str | None = None,
    resolve_brand_logo: bool = True,
) -> CatalogueResponse:
    product_links = (
        sorted(catalogue.product_links, key=lambda item: item.sort_order)
        if include_products
        else []
    )
    if product_count is None:
        product_count = (
            len(product_links)
            if include_products
            else int(
                db.scalar(
                    select(func.count(CatalogueProduct.id)).where(
                        CatalogueProduct.catalogue_id == catalogue.id
                    )
                )
                or 0
            )
        )
    allowed_lists = allowed_price_list_ids(db, actor) if actor else None
    price_fields_allowed = bool(
        actor is None
        or (
            has_permission(actor, "prices.view")
            and (allowed_lists is None or catalogue.price_list_id in allowed_lists)
        )
    )
    if resolve_brand_logo:
        logo_asset = db.scalar(
            select(CatalogueCoverAsset)
            .where(
                CatalogueCoverAsset.catalogue_id == catalogue.id,
                CatalogueCoverAsset.asset_type == "brand_logo",
                CatalogueCoverAsset.deleted_at.is_(None),
            )
            .order_by(CatalogueCoverAsset.created_at.desc())
        )
        brand_logo_url = (
            f"{settings.api_prefix}/v1/catalogues/{catalogue.id}/cover/assets/"
            f"{logo_asset.id}/content?preview=true"
            if logo_asset
            else None
        )
    return CatalogueResponse(
        id=catalogue.id,
        title=catalogue.title,
        slug=catalogue.slug,
        description=catalogue.description,
        brand=catalogue.brand,
        audience=catalogue.audience,
        price_list_id=catalogue.price_list_id if price_fields_allowed else None,
        price_list_name=(catalogue.price_list.name if catalogue.price_list and price_fields_allowed else None),
        show_prices=catalogue.show_prices if price_fields_allowed else False,
        currency=catalogue.currency,
        language=catalogue.language,
        status=catalogue.status,
        version=catalogue.version,
        revision=catalogue.revision,
        valid_from=catalogue.valid_from,
        valid_until=catalogue.valid_until,
        is_public=catalogue.is_public,
        owner_id=catalogue.owner_id,
        product_count=product_count,
        products=(
            [_catalogue_product_response(db, link) for link in product_links]
            if include_products
            else []
        ),
        created_at=catalogue.created_at,
        updated_at=catalogue.updated_at,
        published_at=catalogue.published_at,
        studio_design_id=studio_design.id if studio_design else None,
        catalogue_type=studio_design.catalogue_type if studio_design else "standard",
        brand_logo_url=brand_logo_url,
        studio_editor_href=(
            f"/catalogue-studio/{studio_design.id}/editor"
            if studio_design
            else None
        ),
        studio_preview_href=(
            f"/catalogue-studio/{studio_design.id}/preview"
            if studio_design
            else None
        ),
    )


def _validate_catalogue_pricing(
    price_list: PriceList | None,
    show_prices: bool,
) -> bool:
    if price_list and price_list.is_no_price:
        return False
    if show_prices and not price_list:
        raise HTTPException(
            status_code=422,
            detail="Choose a price list or turn off price display.",
        )
    return show_prices


@router.get("/catalogues", response_model=list[CatalogueResponse])
def list_catalogues(
    q: str | None = None,
    catalogue_status: str | None = Query(default=None, alias="status"),
    actor: User = Depends(require_permission("catalogues.view")),
    db: Session = Depends(get_db),
) -> list[CatalogueResponse]:
    # Studio is the authoring source of truth. Reconcile every design so
    # Catalogue Management always reflects its current title, products and
    # published/draft lifecycle while preserving the exact Studio preview.
    from app.design_studio import synchronize_studio_catalogues

    if synchronize_studio_catalogues(db, actor):
        db.commit()

    product_counts = (
        select(
            CatalogueProduct.catalogue_id.label("catalogue_id"),
            func.count(CatalogueProduct.id).label("product_count"),
        )
        .group_by(CatalogueProduct.catalogue_id)
        .subquery()
    )
    query = (
        select(
            Catalogue,
            func.coalesce(product_counts.c.product_count, 0).label(
                "product_count"
            ),
        )
        .outerjoin(product_counts, product_counts.c.catalogue_id == Catalogue.id)
        .options(
            selectinload(Catalogue.price_list),
            noload(Catalogue.product_links),
            noload(Catalogue.share_links),
        )
    )
    if q:
        pattern = f"%{q.strip()}%"
        query = query.where(
            or_(Catalogue.title.ilike(pattern), Catalogue.description.ilike(pattern))
        )
    if catalogue_status:
        query = query.where(Catalogue.status == catalogue_status)
    catalogue_rows = db.execute(
        query.order_by(Catalogue.updated_at.desc())
    ).unique().all()
    catalogue_rows = [
        (item, int(product_count))
        for item, product_count in catalogue_rows
        if has_record_access(db, actor, "catalogues.view", item)
        and not (
            actor.data_scope
            and actor.data_scope.published_only
            and item.status != "published"
        )
    ]
    catalogue_ids = [catalogue.id for catalogue, _ in catalogue_rows]
    studio_designs_by_catalogue_id: dict[uuid.UUID, CatalogueDesign] = {}
    brand_logo_urls_by_catalogue_id: dict[uuid.UUID, str] = {}
    if catalogue_ids:
        studio_designs = list(
            db.scalars(
                select(CatalogueDesign).where(
                    CatalogueDesign.deleted_at.is_(None),
                    CatalogueDesign.catalogue_id.in_(catalogue_ids),
                )
            ).all()
        )
        studio_designs_by_catalogue_id = {
            design.catalogue_id: design
            for design in studio_designs
            if design.catalogue_id is not None
        }
        logo_assets = list(
            db.scalars(
                select(CatalogueCoverAsset)
                .where(
                    CatalogueCoverAsset.catalogue_id.in_(catalogue_ids),
                    CatalogueCoverAsset.asset_type == "brand_logo",
                    CatalogueCoverAsset.deleted_at.is_(None),
                )
                .order_by(CatalogueCoverAsset.created_at.desc())
            ).all()
        )
        for asset in logo_assets:
            brand_logo_urls_by_catalogue_id.setdefault(
                asset.catalogue_id,
                f"{settings.api_prefix}/v1/catalogues/{asset.catalogue_id}/cover/assets/"
                f"{asset.id}/content?preview=true",
            )
    return [
        _catalogue_response(
            db,
            catalogue,
            include_products=False,
            actor=actor,
            product_count=product_count,
            studio_design=studio_designs_by_catalogue_id.get(catalogue.id),
            brand_logo_url=brand_logo_urls_by_catalogue_id.get(catalogue.id),
            resolve_brand_logo=False,
        )
        for catalogue, product_count in catalogue_rows
    ]


@router.get("/catalogues/brand-options", response_model=list[CatalogueBrandOption])
def list_catalogue_brand_options(
    _: User = Depends(require_permission("catalogues.view")),
    db: Session = Depends(get_db),
) -> list[CatalogueBrandOption]:
    """Return active ERP and user-created brands for catalogue setup."""
    brands = list(
        db.scalars(
            select(Brand).where(Brand.is_active.is_(True)).order_by(Brand.name)
        ).all()
    )
    brands_by_name = {brand.name.casefold(): brand for brand in brands}
    product_counts: dict[str, int] = {}
    product_names: dict[str, str] = {}
    for name, count in db.execute(
        select(Product.brand, func.count(Product.id))
        .where(
            Product.source_system == "gms_erp",
            Product.source_record_exists.is_(True),
            Product.brand.is_not(None),
            Product.brand != "",
        )
        .group_by(Product.brand)
    ).all():
        normalized = str(name).strip()
        if not normalized:
            continue
        key = normalized.casefold()
        product_names.setdefault(key, normalized)
        product_counts[key] = product_counts.get(key, 0) + int(count)

    # User-created brands remain selectable while their first ERP products are
    # still being prepared. ERP-backed brands retain their source spelling and
    # include their current synchronized product count.
    available = [
        (
            brand,
            product_names.get(key, brand.name),
            product_counts.get(key, 0),
        )
        for key, brand in brands_by_name.items()
    ]
    available.sort(key=lambda item: item[1].casefold())
    return [
        CatalogueBrandOption(
            id=brand.id,
            name=source_name,
            code=brand.code,
            product_count=product_count,
        )
        for brand, source_name, product_count in available
    ]


@router.post(
    "/catalogues",
    response_model=CatalogueResponse,
    status_code=status.HTTP_201_CREATED,
)
def create_catalogue(
    payload: CatalogueCreate,
    request: Request,
    actor: User = Depends(require_permission("catalogues.create")),
    db: Session = Depends(get_db),
) -> CatalogueResponse:
    price_list = (
        _get_price_list(db, payload.price_list_id) if payload.price_list_id else None
    )
    if price_list:
        _require_price_list_access(db, actor, price_list.id)
    show_prices = _validate_catalogue_pricing(price_list, payload.show_prices)
    catalogue = Catalogue(
        **payload.model_dump(exclude={"slug", "show_prices"}),
        slug=payload.slug or _slugify(payload.title),
        show_prices=show_prices,
        owner_id=actor.id,
        created_by_id=actor.id,
        updated_by_id=actor.id,
    )
    db.add(catalogue)
    _audit(
        db,
        request,
        actor,
        action="catalogue_created",
        module="catalogues",
        identifier=catalogue.slug,
        details={"title": catalogue.title, "show_prices": show_prices},
    )
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="Catalogue slug already exists.")
    return _catalogue_response(db, _get_catalogue(db, catalogue.id), actor=actor)


@router.post(
    "/catalogues/generate-erp-brands",
    response_model=CatalogueGenerationResponse,
)
def generate_catalogues_for_erp_brands(
    request: Request,
    actor: User = Depends(require_permission("catalogues.create")),
    db: Session = Depends(get_db),
) -> CatalogueGenerationResponse:
    if not has_permission(actor, "catalogues.edit"):
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Generating brand catalogues also requires catalogue edit access.",
        )
    try:
        result = generate_erp_brand_catalogues(db, actor.id)
        _audit(
            db,
            request,
            actor,
            action="erp_brand_catalogues_generated",
            module="catalogues",
            identifier="all-erp-brands",
            details=result,
        )
        db.commit()
    except Exception:
        db.rollback()
        raise
    return CatalogueGenerationResponse(**result)


@router.get("/catalogues/{catalogue_id}", response_model=CatalogueResponse)
def get_catalogue(
    catalogue_id: uuid.UUID,
    actor: User = Depends(require_permission("catalogues.view")),
    db: Session = Depends(get_db),
) -> CatalogueResponse:
    return _catalogue_response(db, _get_scoped_catalogue(db, catalogue_id, actor, "catalogues.view"), actor=actor)


@router.put("/catalogues/{catalogue_id}", response_model=CatalogueResponse)
def update_catalogue(
    catalogue_id: uuid.UUID,
    payload: CatalogueUpdate,
    request: Request,
    actor: User = Depends(require_permission("catalogues.edit")),
    db: Session = Depends(get_db),
) -> CatalogueResponse:
    catalogue = _get_scoped_catalogue(db, catalogue_id, actor, "catalogues.edit")
    if catalogue.revision != payload.expected_revision:
        raise HTTPException(
            status_code=409,
            detail="This catalogue changed in another session. Refresh and try again.",
        )
    price_list = (
        _get_price_list(db, payload.price_list_id) if payload.price_list_id else None
    )
    if price_list:
        _require_price_list_access(db, actor, price_list.id)
    show_prices = _validate_catalogue_pricing(price_list, payload.show_prices)
    update_values = payload.model_dump(
        exclude={"slug", "show_prices", "status", "expected_revision"}
    )
    for field, value in update_values.items():
        setattr(catalogue, field, value)
    catalogue.slug = payload.slug or catalogue.slug
    catalogue.show_prices = show_prices
    catalogue.status = (
        "draft" if catalogue.status == "published" else payload.status
    )
    catalogue.updated_by_id = actor.id
    catalogue.revision += 1
    _audit(
        db,
        request,
        actor,
        action="catalogue_updated",
        module="catalogues",
        identifier=catalogue.slug,
        details={"revision": catalogue.revision},
    )
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="Catalogue slug already exists.")
    return _catalogue_response(db, _get_catalogue(db, catalogue.id), actor=actor)


@router.put("/catalogues/{catalogue_id}/products", response_model=CatalogueResponse)
def set_catalogue_products(
    catalogue_id: uuid.UUID,
    payload: CatalogueProductsUpdate,
    request: Request,
    actor: User = Depends(require_permission("catalogues.edit")),
    db: Session = Depends(get_db),
) -> CatalogueResponse:
    catalogue = _get_scoped_catalogue(db, catalogue_id, actor, "catalogues.edit")
    products = {
        product.id: product
        for product in db.scalars(
            select(Product).where(
                Product.id.in_([item.product_id for item in payload.products])
            )
        )
    }
    if len(products) != len(payload.products):
        raise HTTPException(status_code=404, detail="One or more products were not found.")
    selected_video_ids = {item.selected_video_id for item in payload.products if item.selected_video_id}
    if selected_video_ids:
        selected_rows = {
            video.id: video
            for video in db.scalars(select(ProductVideo).where(ProductVideo.id.in_(selected_video_ids), ProductVideo.deleted_at.is_(None)))
        }
        for item in payload.products:
            if item.selected_video_id and (
                item.selected_video_id not in selected_rows
                or selected_rows[item.selected_video_id].product_id != item.product_id
                or not selected_rows[item.selected_video_id].is_active
            ):
                raise HTTPException(status_code=422, detail="A selected video is unavailable for its product.")
    if any(not has_record_access(db, actor, "products.view", product) for product in products.values()):
        raise HTTPException(status_code=404, detail="One or more products were not found.")
    selected_inactive = sorted(
        product.sku for product in products.values() if product.status != "active"
    )
    if selected_inactive:
        raise HTTPException(
            status_code=422,
            detail="Inactive products cannot be saved in a catalogue: "
            + ", ".join(selected_inactive[:5]),
        )
    for existing_link in list(catalogue.product_links):
        db.delete(existing_link)
    db.flush()
    catalogue.product_links = [
        CatalogueProduct(
            product_id=item.product_id,
            section_title=item.section_title,
            override_description=item.override_description,
            hide_price=item.hide_price,
            include_video=item.include_video,
            selected_video_id=item.selected_video_id,
            video_title_override=item.video_title_override,
            video_description_override=item.video_description_override,
            video_display_mode=item.video_display_mode,
            video_thumbnail_mode=item.video_thumbnail_mode,
            sort_order=index,
        )
        for index, item in enumerate(payload.products, start=1)
    ]
    if catalogue.status == "published":
        catalogue.status = "draft"
    catalogue.updated_by_id = actor.id
    catalogue.revision += 1
    _audit(
        db,
        request,
        actor,
        action="catalogue_products_updated",
        module="catalogues",
        identifier=catalogue.slug,
        details={"product_count": len(payload.products), "videos_included": sum(item.include_video for item in payload.products)},
    )
    if any(item.include_video or item.video_title_override or item.video_description_override for item in payload.products):
        _audit(db, request, actor, action="catalogue_video_visibility_changed", module="product_videos", identifier=catalogue.slug, details={"catalogue_id": str(catalogue.id), "included_product_ids": [str(item.product_id) for item in payload.products if item.include_video]})
    db.commit()
    return _catalogue_response(db, _get_catalogue(db, catalogue.id), actor=actor)


def _build_snapshot(
    db: Session,
    catalogue: Catalogue,
    *,
    strict_prices: bool,
) -> dict:
    show_prices = bool(
        catalogue.show_prices
        and catalogue.price_list
        and not catalogue.price_list.is_no_price
    )
    inactive_brand_names = {
        str(brand.name).strip().casefold()
        for brand in db.scalars(
            select(Brand).where(Brand.is_active.is_(False))
        ).all()
        if getattr(brand, "name", None)
    }
    items: list[dict] = []
    for link in sorted(catalogue.product_links, key=lambda item: item.sort_order):
        product = link.product
        if product.status != "active":
            continue
        if (product.brand or "").strip().casefold() in inactive_brand_names:
            continue
        entry = product.catalogue_entry or CatalogueEntry()
        primary_image: ProductImage | None = next(
            (image for image in product.images if image.is_primary),
            product.images[0] if product.images else None,
        )
        selected_video = None
        if link.include_video:
            selected_video = next(
                (
                    video for video in product.videos
                    if video.deleted_at is None and video.is_active and video.show_in_catalogue
                    and (video.id == link.selected_video_id if link.selected_video_id else video.is_featured)
                ),
                None,
            )
        item = {
            "id": str(product.id),
            "sku": product.sku,
            "name": entry.display_name or product.erp_name,
            "name_en": entry.display_name or product.erp_name,
            "name_th": product.erp_name_th,
            "brand": product.brand,
            "description": catalogue_short_description(product, entry, link.override_description),
            "long_description": catalogue_long_description(product, entry),
            "description_en": link.override_description or entry.short_description or erp_summary(product, "en"),
            "description_th": link.override_description or erp_summary(product, "th") or entry.short_description,
            "long_description_en": entry.long_description or erp_description(product, "en"),
            "long_description_th": erp_description(product, "th") or entry.long_description,
            "erp_details": erp_details(product),
            "categories": [category.name for category in product.categories],
            "category_refs": [
                {"id": category.id, "name": category.name, "slug": category.slug, "description": category.description}
                for category in product.categories
            ],
            "image_url": primary_image.public_url if primary_image else None,
            "image_urls": [image.public_url for image in sorted(product.images, key=lambda row: (not row.is_primary, row.sort_order, row.created_at))],
            "barcode": product.barcode,
            "stock_quantity": product.stock_quantity,
            "section_title": link.section_title,
            "sort_order": link.sort_order,
            "featured": entry.is_featured,
            "product_status": product.status,
        }
        if selected_video:
            item["video"] = {
                "id": str(selected_video.id), "source_type": selected_video.source_type,
                "provider": selected_video.provider, "external_video_id": selected_video.external_video_id,
                "external_url": selected_video.external_url,
                "title": link.video_title_override or selected_video.title_en or selected_video.title_th,
                "description": link.video_description_override or selected_video.description,
                "alt_text": selected_video.alt_text, "storage_key": selected_video.playback_storage_key or selected_video.storage_key,
                "thumbnail_storage_key": selected_video.thumbnail_storage_key,
                "thumbnail_mode": link.video_thumbnail_mode,
                "thumbnail_product_image_url": primary_image.public_url if link.video_thumbnail_mode == "product_image" and primary_image else None,
                "caption_storage_key": selected_video.caption_storage_key, "mime_type": selected_video.mime_type,
                "duration_seconds": float(selected_video.duration_seconds) if selected_video.duration_seconds is not None else None,
                "width": selected_video.width, "height": selected_video.height,
                "display_mode": link.video_display_mode, "show_controls": selected_video.show_controls,
                "allow_download": selected_video.allow_download, "autoplay": selected_video.autoplay,
                "muted": selected_video.muted, "loop": selected_video.loop,
                "show_in_public_catalogue": selected_video.show_in_public_catalogue,
            }
        if show_prices and not link.hide_price:
            price = _current_price(db, product.id, catalogue.price_list_id)
            if strict_prices and product.status == "active" and not price:
                raise HTTPException(
                    status_code=422,
                    detail=f"{product.sku} has no active {catalogue.price_list.name} price.",
                )
            if price:
                item["price"] = str(price.amount)
                item["currency"] = price.currency
            else:
                item["price_missing"] = True
        items.append(item)
    category_ids: list[int] = []
    for item in items:
        for category in item.get("category_refs", []):
            if category["id"] not in category_ids:
                category_ids.append(category["id"])
    category_settings = {
        item.category_id: item
        for item in db.scalars(
            select(CatalogueCategorySetting).where(
                CatalogueCategorySetting.catalogue_id == catalogue.id
            )
        ).all()
    }
    category_master = {
        item.id: item
        for item in db.scalars(select(Category).where(Category.id.in_(category_ids))).all()
    } if category_ids else {}
    categories = []
    for fallback_order, category_id in enumerate(category_ids, 1):
        master = category_master.get(category_id)
        setting = category_settings.get(category_id)
        if not master or not master.is_active or (setting and not setting.is_visible):
            continue
        categories.append({
            "id": master.id,
            "slug": master.slug,
            "name": setting.display_name or master.name if setting else master.name,
            "description": setting.description or master.description if setting else master.description,
            "banner_url": storage.public_url(setting.banner_storage_key) if setting and setting.banner_storage_key else None,
            "display_order": setting.display_order if setting else fallback_order,
            "show_product_count": setting.show_product_count if setting else True,
            "default_expanded": setting.default_expanded if setting else False,
        })
    categories.sort(key=lambda item: item["display_order"])
    visible_category_ids = {item["id"] for item in categories}
    visible_items: list[dict] = []
    for item in items:
        original_refs = list(item.get("category_refs", []))
        refs = [ref for ref in original_refs if ref["id"] in visible_category_ids]
        if original_refs and not refs:
            continue
        refs.sort(key=lambda ref: next((cat["display_order"] for cat in categories if cat["id"] == ref["id"]), 9999))
        item["category_refs"] = refs
        item["categories"] = [next((cat["name"] for cat in categories if cat["id"] == ref["id"]), ref["name"]) for ref in refs]
        visible_items.append(item)
    items = visible_items
    for category in categories:
        category["product_count"] = sum(any(ref["id"] == category["id"] for ref in item.get("category_refs", [])) for item in items)

    cover_record = db.scalar(select(CatalogueCoverSetting).where(CatalogueCoverSetting.catalogue_id == catalogue.id))
    cover = None
    if cover_record:
        cover_fields = (
            "cover_mode", "catalogue_name", "catalogue_year", "subtitle", "company_name", "collection_name",
            "background_color", "overlay_color", "overlay_opacity", "background_fit", "show_catalogue_name",
            "show_catalogue_year", "show_subtitle", "show_brand_logo", "show_company_logo", "show_start_button",
            "title_color", "title_font_size", "title_alignment", "title_position_x_percent", "title_position_y_percent",
            "title_width_percent", "title_z_index", "subtitle_color", "subtitle_font_size",
            "subtitle_position_x_percent", "subtitle_position_y_percent", "cover_alt_text",
        )
        cover = {field: getattr(cover_record, field) for field in cover_fields}
        cover["default_company_logo_version"] = LOGO_VERSION
        cover["assets"] = []
        active_assets = db.scalars(select(CatalogueCoverAsset).where(
            CatalogueCoverAsset.catalogue_id == catalogue.id,
            CatalogueCoverAsset.deleted_at.is_(None),
        )).all()
        for asset in active_assets:
            asset_url = f"{settings.api_prefix}/v1/catalogues/{catalogue.id}/cover/assets/{asset.id}/content"
            cover["assets"].append({
                "id": str(asset.id), "asset_type": asset.asset_type, "file_url": asset_url,
                "preview_url": f"{asset_url}?preview=true", "mime_type": asset.mime_type,
                "original_filename": asset.original_filename, "file_size": asset.file_size,
                "width": asset.width, "height": asset.height, "alt_text": asset.alt_text,
                "position_x_percent": asset.position_x_percent, "position_y_percent": asset.position_y_percent,
                "width_percent": asset.width_percent, "height_percent": asset.height_percent,
                "opacity": asset.opacity, "rotation": asset.rotation, "z_index": asset.z_index,
                "storage_key": asset.storage_key, "preview_storage_key": asset.preview_storage_key,
            })
    snapshot = {
        "catalogue_id": str(catalogue.id),
        "title": catalogue.title,
        "description": catalogue.description,
        "audience": catalogue.audience,
        "language": catalogue.language,
        "status": catalogue.status,
        "show_prices": show_prices,
        "products": items,
        "categories": categories,
        "cover": cover,
        "valid_from": catalogue.valid_from.isoformat() if catalogue.valid_from else None,
        "valid_until": catalogue.valid_until.isoformat() if catalogue.valid_until else None,
        "generated_at": _now().isoformat(),
    }
    if show_prices:
        snapshot["currency"] = catalogue.currency
        snapshot["price_list"] = catalogue.price_list.name
        snapshot["price_list_id"] = catalogue.price_list.id
    return snapshot


def _selected_version(
    db: Session,
    catalogue_id: uuid.UUID,
    version_number: int,
) -> CatalogueVersion:
    version = db.scalar(
        select(CatalogueVersion).where(
            CatalogueVersion.catalogue_id == catalogue_id,
            CatalogueVersion.version_number == version_number,
        )
    )
    if not version:
        raise HTTPException(
            status_code=404,
            detail="The selected catalogue version is unavailable.",
        )
    return version


def _requested_version(
    version: int | None,
    version_number: int | None,
) -> int | None:
    if version is not None and version_number is not None and version != version_number:
        raise HTTPException(
            status_code=422,
            detail="Choose one catalogue version.",
        )
    return version if version is not None else version_number


def _presentation_source(
    db: Session,
    catalogue: Catalogue,
    version_number: int | None,
) -> tuple[dict, int | None]:
    selected_number = version_number
    if selected_number is None and catalogue.status in {"published", "archived"}:
        selected_number = catalogue.version or None
    if selected_number is not None:
        selected = _selected_version(db, catalogue.id, selected_number)
        return deepcopy(selected.snapshot), selected.version_number
    if catalogue.status == "archived":
        raise HTTPException(
            status_code=409,
            detail="This archived catalogue has no published version to preview.",
        )
    return _build_snapshot(db, catalogue, strict_prices=False), None


def _contains_thai(value: str) -> bool:
    return any("\u0e00" <= character <= "\u0e7f" for character in value)


def _normalized_cover(snapshot: dict, catalogue: Catalogue) -> dict:
    source = deepcopy(snapshot.get("cover") or {})
    if "catalogue_name" in source:
        return source
    title = str(snapshot.get("title") or catalogue.title)
    match = re.match(r"^(.*?)\s+((?:19|20)\d{2})$", title.strip())
    name, year = (match.group(1), match.group(2)) if match else (title, "")
    normalized = {
        "cover_mode": "custom", "catalogue_name": name, "catalogue_year": year,
        "subtitle": str(source.get("subtitle") or snapshot.get("description") or ""),
        "company_name": "", "collection_name": "", "background_color": "#164f35",
        "overlay_color": "#081f14", "overlay_opacity": 0.28,
        "background_fit": source.get("display_mode") or "cover", "show_catalogue_name": True,
        "show_catalogue_year": True, "show_subtitle": True,
        "show_brand_logo": bool(source.get("show_logo_overlay", False)), "show_company_logo": False,
        "show_start_button": True, "title_color": "#ffffff", "title_font_size": 72,
        "title_alignment": "left", "title_position_x_percent": 10, "title_position_y_percent": 68,
        "title_width_percent": 80, "title_z_index": 20, "subtitle_color": "#ffffff",
        "subtitle_font_size": 20, "subtitle_position_x_percent": 10,
        "subtitle_position_y_percent": 86, "cover_alt_text": source.get("alt_text") or title,
        "default_company_logo_version": None,
        "assets": [],
    }
    if source.get("file_url"):
        normalized["cover_mode"] = "full_image"
        normalized["show_catalogue_name"] = bool(source.get("show_title_overlay"))
        normalized["assets"].append({
            "asset_type": "full_cover", "file_url": source["file_url"],
            "preview_url": source.get("preview_url") or source["file_url"],
            "mime_type": source.get("mime_type") or "image/jpeg",
            "original_filename": source.get("original_filename") or "legacy-cover",
            "file_size": source.get("file_size") or 0, "width": source.get("width") or 1,
            "height": source.get("height") or 1, "alt_text": source.get("alt_text") or title,
            "position_x_percent": source.get("position_x", 50), "position_y_percent": source.get("position_y", 50),
            "width_percent": 100, "height_percent": 100, "opacity": 1, "rotation": 0, "z_index": 0,
        })
    return normalized


_ERP_CARD_PRICE_FIELDS = {
    "SP6": "wholesale_price",
    "SRP": "online_price",
    "SP5": "retail_price",
}


def _apply_erp_card_prices(
    db: Session,
    presentation: CataloguePreviewResponse,
) -> CataloguePreviewResponse:
    """Attach live ERP card prices without changing audience pricing."""

    for product in presentation.products:
        product.wholesale_price = None
        product.online_price = None
        product.retail_price = None
    if not presentation.show_prices:
        return presentation

    product_ids = [product.id for product in presentation.products if product.id]
    if not product_ids:
        return presentation

    rows = db.execute(
        select(
            ErpProductCustomerPrice.product_id,
            ErpCustomerPriceLevel.source_code,
            ErpProductCustomerPrice.amount,
        )
        .join(
            ErpCustomerPriceLevel,
            ErpCustomerPriceLevel.id == ErpProductCustomerPrice.price_level_id,
        )
        .where(
            ErpProductCustomerPrice.product_id.in_(product_ids),
            ErpCustomerPriceLevel.source_code.in_(_ERP_CARD_PRICE_FIELDS),
            ErpCustomerPriceLevel.is_active.is_(True),
        )
    ).all()
    prices_by_product: dict[uuid.UUID, dict[str, Decimal]] = {}
    for product_id, source_code, amount in rows:
        field_name = _ERP_CARD_PRICE_FIELDS.get(str(source_code).upper())
        if field_name:
            prices_by_product.setdefault(product_id, {})[field_name] = amount

    for product in presentation.products:
        if not product.id:
            continue
        for field_name, amount in prices_by_product.get(product.id, {}).items():
            setattr(product, field_name, amount)
    return presentation


def _safe_presentation(
    db: Session,
    catalogue: Catalogue,
    snapshot: dict,
    version_number: int | None,
    actor: User | None,
    include_inactive: bool = False,
) -> CataloguePreviewResponse:
    prices_allowed = bool(
        snapshot.get("show_prices")
        and actor is not None
        and has_permission(actor, "prices.view")
    )
    allowed_price_lists = allowed_price_list_ids(db, actor) if actor is not None else set()
    if (
        prices_allowed
        and allowed_price_lists is not None
        and snapshot.get("price_list_id") not in allowed_price_lists
    ):
        prices_allowed = False
    snapshot_products = snapshot.get("products", [])
    inactive_brand_names = {
        str(brand.name).strip().casefold()
        for brand in db.scalars(
            select(Brand).where(Brand.is_active.is_(False))
        ).all()
        if getattr(brand, "name", None)
    }
    snapshot_categories = list(snapshot.get("categories", []))
    snapshot_category_ids: list[int] = []
    snapshot_category_names_by_id: dict[int, str] = {}
    for source_category in snapshot_categories:
        try:
            category_id = int(source_category.get("id"))
        except (TypeError, ValueError):
            continue
        if category_id not in snapshot_category_ids:
            snapshot_category_ids.append(category_id)
        snapshot_category_names_by_id[category_id] = str(
            source_category.get("name") or ""
        )
    live_category_activity = {
        category.id: category.is_active
        for category in db.scalars(
            select(Category).where(Category.id.in_(snapshot_category_ids))
        ).all()
    } if snapshot_category_ids else {}
    inactive_category_ids = {
        category_id
        for category_id, is_active in live_category_activity.items()
        if not is_active
    }
    inactive_category_names = {
        snapshot_category_names_by_id[category_id].casefold()
        for category_id in inactive_category_ids
        if category_id in snapshot_category_names_by_id
    }
    snapshot_ids = []
    for item in snapshot_products:
        try:
            if item.get("id"):
                snapshot_ids.append(uuid.UUID(str(item["id"])))
        except (ValueError, TypeError):
            continue
    # Published versions keep their order, selection, pricing and catalogue copy
    # immutable. ERP-owned facts are filled from the live product master when an
    # older snapshot predates those fields, so existing customer links benefit
    # from names/details without silently replacing authored catalogue content.
    live_products = {
        str(product.id): product
        for product in db.scalars(
            select(Product).options(noload("*")).where(Product.id.in_(snapshot_ids))
        )
    } if snapshot_ids else {}
    live_statuses = {
        product_id: product.status for product_id, product in live_products.items()
    }
    published_cards = {
        str(card.product_id): card
        for card in db.scalars(
            select(GlobalProductCard).where(
                GlobalProductCard.product_id.in_(snapshot_ids),
                GlobalProductCard.published_json.is_not(None),
            )
        )
    } if snapshot_ids else {}
    snapshot_video_ids: list[uuid.UUID] = []
    for item in snapshot_products:
        try:
            if item.get("video", {}).get("id"):
                snapshot_video_ids.append(uuid.UUID(str(item["video"]["id"])))
        except (ValueError, TypeError):
            continue
    live_video_visibility = {
        str(video_id): (active, deleted_at, public)
        for video_id, active, deleted_at, public in db.execute(
            select(ProductVideo.id, ProductVideo.is_active, ProductVideo.deleted_at, ProductVideo.show_in_public_catalogue)
            .where(ProductVideo.id.in_(snapshot_video_ids))
        )
    } if snapshot_video_ids else {}
    product_rows: list[dict] = []
    for source in sorted(
        snapshot_products,
        key=lambda item: int(item.get("sort_order", 0)),
    ):
        item = deepcopy(source)
        if str(item.get("brand") or "").strip().casefold() in inactive_brand_names:
            continue
        live_product = live_products.get(str(item.get("id")))
        published_card = published_cards.get(str(item.get("id")))
        live_status = live_statuses.get(str(item.get("id")), "inactive")
        if live_status != "active" and not include_inactive:
            continue
        if live_product:
            item.setdefault("name_en", live_product.erp_name)
            if live_product.erp_name_th:
                item.setdefault("name_th", live_product.erp_name_th)
            item.setdefault("description_en", erp_summary(live_product, "en"))
            item.setdefault("description_th", erp_summary(live_product, "th"))
            item.setdefault("long_description_en", erp_description(live_product, "en"))
            item.setdefault("long_description_th", erp_description(live_product, "th"))
            item.setdefault("erp_details", erp_details(live_product))
            # These fields are owned by the ERP mirror and are never editable
            # through the global card presentation editor.
            item["sku"] = live_product.sku
            item["barcode"] = live_product.barcode
            # Stock is operational live data, not immutable catalogue content.
            # Always replace the published snapshot value with the latest ERP
            # mirror value so publishing never freezes inventory.
            item["stock_quantity"] = live_product.stock_quantity
        card_presentation = (
            deepcopy(published_card.published_json)
            if published_card and published_card.published_json
            else None
        )
        if card_presentation:
            language = str(snapshot.get("language") or catalogue.language or "en")
            display_name = card_presentation.get(
                "display_name_th" if language == "th" else "display_name"
            ) or card_presentation.get("display_name")
            if display_name:
                item["name"] = str(display_name)
                item["name_th" if language == "th" else "name_en"] = str(display_name)
            description = card_presentation.get(
                "description_th" if language == "th" else "description"
            ) or card_presentation.get("description")
            if description:
                item["description_th" if language == "th" else "description_en"] = str(description)
            image_urls = [
                str(value)
                for value in card_presentation.get("image_urls", [])
                if isinstance(value, str) and value.startswith("/uploads/")
            ]
            if image_urls:
                item["image_url"] = image_urls[0]
                item["image_urls"] = image_urls
        language = str(snapshot.get("language") or catalogue.language or "en")
        raw_name = str(
            (item.get("name_th") if language == "th" else item.get("name_en"))
            or item.get("name")
            or "Unnamed product"
        )
        source_categories = [str(value) for value in item.get("categories", [])]
        categories = [
            value
            for value in source_categories
            if value.casefold() not in inactive_category_names
        ]
        if source_categories and not categories:
            continue
        row = {
            "id": item.get("id"),
            "code": str(item.get("sku") or ""),
            "name": raw_name,
            "name_th": item.get("name_th") or (raw_name if _contains_thai(raw_name) else None),
            "name_en": item.get("name_en") or (None if _contains_thai(raw_name) else raw_name),
            "brand": item.get("brand"),
            "category_name": categories[0] if categories else None,
            "categories": categories,
            "description": str(
                item.get("description_th" if language == "th" else "description_en")
                or item.get("description")
                or ""
            ),
            "long_description": str(
                item.get("long_description_th" if language == "th" else "long_description_en")
                or item.get("long_description")
                or ""
            ),
            "description_en": str(item.get("description_en") or item.get("description") or ""),
            "description_th": str(item.get("description_th") or item.get("description") or ""),
            "long_description_en": str(item.get("long_description_en") or item.get("long_description") or ""),
            "long_description_th": str(item.get("long_description_th") or item.get("long_description") or ""),
            "erp_details": dict(item.get("erp_details") or {}),
            "main_image_url": item.get("image_url"),
            "image_urls": [str(value) for value in item.get("image_urls", []) if value],
            "barcode": item.get("barcode"),
            "stock_quantity": item.get("stock_quantity"),
            "card_presentation": card_presentation,
            "section_title": str(item.get("section_title") or ""),
            "display_order": int(item.get("sort_order", len(product_rows) + 1)),
            "featured": bool(item.get("featured", False)),
            "product_status": live_status,
        }
        source_video = item.get("video")
        if source_video:
            visibility = live_video_visibility.get(str(source_video.get("id")))
            video_allowed = bool(visibility and visibility[0] and visibility[1] is None)
            if actor is None:
                video_allowed = video_allowed and bool(source_video.get("show_in_public_catalogue")) and bool(visibility and visibility[2])
            if video_allowed:
                provider = str(source_video.get("provider") or "internal")
                playback_url = None
                if provider == "youtube": playback_url = f"https://www.youtube-nocookie.com/embed/{source_video.get('external_video_id')}"
                elif provider == "vimeo": playback_url = f"https://player.vimeo.com/video/{source_video.get('external_video_id')}"
                elif provider == "direct_url": playback_url = source_video.get("external_url")
                elif actor is not None: playback_url = f"{settings.api_prefix}/v1/products/{item.get('id')}/videos/{source_video.get('id')}/content"
                row["video"] = {
                    "id": source_video.get("id"), "source_type": source_video.get("source_type"), "provider": provider,
                    "title": source_video.get("title") or "", "description": source_video.get("description") or "",
                    "alt_text": source_video.get("alt_text") or "", "playback_url": playback_url,
                    "thumbnail_url": (source_video.get("thumbnail_product_image_url") if source_video.get("thumbnail_mode") == "product_image" else (f"{settings.api_prefix}/v1/products/{item.get('id')}/videos/{source_video.get('id')}/thumbnail/content" if source_video.get("thumbnail_storage_key") and actor is not None else None)),
                    "caption_url": (f"{settings.api_prefix}/v1/products/{item.get('id')}/videos/{source_video.get('id')}/caption/content" if source_video.get("caption_storage_key") and actor is not None else None),
                    "mime_type": source_video.get("mime_type"), "duration_seconds": source_video.get("duration_seconds"),
                    "width": source_video.get("width"), "height": source_video.get("height"),
                    "display_mode": source_video.get("display_mode") or "product_detail",
                    "show_controls": bool(source_video.get("show_controls", True)), "allow_download": bool(source_video.get("allow_download", False)),
                    "autoplay": bool(source_video.get("autoplay", False)), "muted": bool(source_video.get("muted", False)), "loop": bool(source_video.get("loop", False)),
                }
        if prices_allowed and item.get("price") is not None:
            row["price"] = item["price"]
            row["currency"] = item.get("currency") or snapshot.get("currency")
        product_rows.append(row)

    price_list = None
    if prices_allowed and snapshot.get("price_list"):
        price_list = {
            "id": snapshot.get("price_list_id"),
            "name": str(snapshot["price_list"]),
            "show_price": True,
        }
    generated_at = snapshot.get("generated_at") or _now()
    presentation_status = "published" if version_number is not None else catalogue.status
    visible_category_counts: dict[str, int] = {}
    for product in product_rows:
        for category_name in product.get("categories", []):
            visible_category_counts[category_name] = visible_category_counts.get(category_name, 0) + 1
    visible_categories = []
    for source_category in snapshot_categories:
        try:
            source_category_id = int(source_category.get("id"))
        except (TypeError, ValueError):
            source_category_id = None
        if source_category_id in inactive_category_ids:
            continue
        category = deepcopy(source_category)
        category["product_count"] = visible_category_counts.get(str(category.get("name") or ""), 0)
        visible_categories.append(category)
    is_erp_brand_catalogue = str(catalogue.slug or "").startswith("erp-brand-")
    studio_design = db.scalar(
        select(CatalogueDesign)
        .where(
            CatalogueDesign.catalogue_id == catalogue.id,
            CatalogueDesign.deleted_at.is_(None),
        )
        .order_by(CatalogueDesign.updated_at.desc())
        .limit(1)
    )
    presentation = CataloguePreviewResponse(
        catalogue_id=catalogue.id,
        studio_design_id=studio_design.id if studio_design else None,
        studio_preview_href=(
            f"/catalogue-studio/{studio_design.id}/preview"
            if studio_design
            else None
        ),
        version=version_number,
        title=str(snapshot.get("title") or catalogue.title),
        description=str(snapshot.get("description") or ""),
        audience=str(snapshot.get("audience") or catalogue.audience),
        language=str(snapshot.get("language") or catalogue.language),
        status=presentation_status,
        is_draft=version_number is None,
        product_card_style=(
            "erp_detail"
            if is_erp_brand_catalogue
            else "standard"
        ),
        product_card_theme=(
            _erp_product_card_theme(catalogue.brand or catalogue.title)
            if is_erp_brand_catalogue
            else None
        ),
        price_list=price_list,
        show_prices=prices_allowed,
        currency=(str(snapshot.get("currency")) if prices_allowed else None),
        product_count=len(product_rows),
        products=product_rows,
        categories=visible_categories,
        cover=_normalized_cover(snapshot, catalogue),
        valid_from=snapshot.get("valid_from"),
        valid_until=snapshot.get("valid_until"),
        generated_at=generated_at,
    )
    apply_purchase_order_quantities(presentation.products)
    from app.promotion_service import apply_promotions_to_catalogue
    presentation.promotions = apply_promotions_to_catalogue(
        db, catalogue_id=catalogue.id, products=presentation.products,
        audience_type_id=None, price_list_id=snapshot.get("price_list_id"),
        show_prices=prices_allowed,
    )
    return presentation


def _apply_user_catalogue_price_mapping(
    db: Session,
    presentation: CataloguePreviewResponse,
    actor: User,
    audience_type_id: int | None,
    brand: str | None = None,
) -> CataloguePreviewResponse:
    """Apply the signed-in user's selected ERP price to preview/export output."""

    if audience_type_id is None:
        return _apply_erp_card_prices(db, presentation)
    audience = db.scalar(
        select(CatalogueAudienceType)
        .options(selectinload(CatalogueAudienceType.price_list))
        .where(
            CatalogueAudienceType.id == audience_type_id,
            CatalogueAudienceType.is_active.is_(True),
        )
    )
    if audience is None:
        raise HTTPException(status_code=422, detail="The selected customer level is unavailable.")
    brand_mapping = db.scalar(
        select(UserBrandCataloguePriceMapping)
        .options(selectinload(UserBrandCataloguePriceMapping.price_list))
        .where(
            UserBrandCataloguePriceMapping.user_id == actor.id,
            UserBrandCataloguePriceMapping.brand_key == brand.strip().casefold(),
            UserBrandCataloguePriceMapping.audience_type_id == audience.id,
        )
    ) if brand and brand.strip() else None
    mapping = brand_mapping or db.scalar(
        select(UserCataloguePriceMapping)
        .options(selectinload(UserCataloguePriceMapping.price_list))
        .where(
            UserCataloguePriceMapping.user_id == actor.id,
            UserCataloguePriceMapping.audience_type_id == audience.id,
        )
    )
    price_list = mapping.price_list if mapping else audience.price_list
    allowed = allowed_price_list_ids(db, actor)
    if price_list and allowed is not None and price_list.id not in allowed:
        raise HTTPException(status_code=403, detail="You cannot use the mapped ERP price level.")

    show_prices = bool(
        has_permission(actor, "prices.view")
        and audience.show_prices
        and price_list
        and not price_list.is_no_price
    )
    presentation.audience = audience.display_name
    presentation.show_prices = show_prices
    presentation.price_list = (
        CataloguePreviewPriceList(
            id=price_list.id,
            name=price_list.name,
            show_price=True,
        )
        if show_prices and price_list
        else None
    )
    presentation.currency = price_list.currency if show_prices and price_list else None

    prices: dict[uuid.UUID, ProductPrice] = {}
    if show_prices and price_list:
        product_ids = [product.id for product in presentation.products if product.id]
        now = _now()
        rows = db.scalars(
            select(ProductPrice)
            .where(
                ProductPrice.product_id.in_(product_ids),
                ProductPrice.price_list_id == price_list.id,
                ProductPrice.status == "active",
                ProductPrice.effective_from <= now,
                or_(ProductPrice.expires_at.is_(None), ProductPrice.expires_at > now),
            )
            .order_by(ProductPrice.effective_from.desc())
        ).all()
        for row in rows:
            prices.setdefault(row.product_id, row)

    for product in presentation.products:
        price = prices.get(product.id) if show_prices else None
        product.price = price.amount if price else None
        product.currency = price.currency if price else None
        product.original_price = None
        product.promotion_price = None
        product.discount_percent = None
        product.promotion_code = None
        product.promotion_name = None
        product.promotion_badge = None
        product.promotion_end_at = None

    _apply_erp_card_prices(db, presentation)
    from app.promotion_service import apply_promotions_to_catalogue
    presentation.promotions = apply_promotions_to_catalogue(
        db,
        catalogue_id=presentation.catalogue_id,
        products=presentation.products,
        audience_type_id=audience.id,
        price_list_id=price_list.id if price_list else None,
        show_prices=show_prices,
    )
    return presentation


@router.get(
    "/public/catalogues/by-slug/{slug}",
    response_model=CataloguePreviewResponse,
    response_model_exclude_none=True,
)
def public_catalogue(slug: str, db: Session = Depends(get_db)) -> CataloguePreviewResponse:
    catalogue = db.scalar(select(Catalogue).where(
        Catalogue.slug == slug,
        Catalogue.is_public.is_(True),
        Catalogue.status == "published",
    ))
    if not catalogue:
        raise HTTPException(status_code=404, detail="Catalogue not found.")
    snapshot, resolved_version = _presentation_source(db, catalogue, None)
    return _safe_presentation(db, catalogue, snapshot, resolved_version, None)


@router.get("/public/products/{sku}")
def public_product(sku: str, db: Session = Depends(get_db)) -> dict:
    product = db.scalar(select(Product).options(
        selectinload(Product.catalogue_entry),
        selectinload(Product.images),
        selectinload(Product.categories),
    ).join(CatalogueEntry).where(
        Product.sku == sku,
        Product.status == "active",
        CatalogueEntry.visibility == "public",
        CatalogueEntry.workflow_status == "published",
    ))
    if not product:
        raise HTTPException(status_code=404, detail="Product not found.")
    entry = product.catalogue_entry
    image = next((item for item in product.images if item.is_primary), product.images[0] if product.images else None)
    return {
        "id": str(product.id), "code": product.sku,
        "name": entry.display_name if entry and entry.display_name else product.erp_name,
        "brand": product.brand,
        "short_description": catalogue_short_description(product, entry or CatalogueEntry()),
        "erp_details": erp_details(product),
        "image_url": image.public_url if image else None,
        "categories": [category.name for category in product.categories],
    }


@router.get(
    "/catalogues/{catalogue_id}/preview",
    response_model=CataloguePreviewResponse,
    response_model_exclude_none=True,
)
def preview_catalogue(
    catalogue_id: uuid.UUID,
    request: Request,
    version: int | None = Query(default=None, ge=1),
    version_number: int | None = Query(default=None, ge=1),
    audience_type_id: int | None = Query(default=None, ge=1),
    include_inactive: bool = False,
    actor: User = Depends(require_permission("catalogues.preview")),
    db: Session = Depends(get_db),
) -> CataloguePreviewResponse:
    catalogue = _get_scoped_catalogue(db, catalogue_id, actor, "catalogues.preview")
    selected_number = _requested_version(version, version_number)
    snapshot, resolved_version = _presentation_source(
        db,
        catalogue,
        selected_number,
    )
    if include_inactive and not is_superadmin(actor):
        raise HTTPException(
            status_code=403,
            detail="SuperAdmin access is required to view inactive products.",
        )
    presentation = _safe_presentation(db, catalogue, snapshot, resolved_version, actor, include_inactive=include_inactive)
    presentation = _apply_user_catalogue_price_mapping(
        db, presentation, actor, audience_type_id, catalogue.brand
    )
    audit_details = {
        "catalogue_id": str(catalogue.id),
        "version": resolved_version,
        "price_list": (
            presentation.price_list.name if presentation.price_list else None
        ),
        "prices_included": presentation.show_prices,
    }
    _audit(
        db,
        request,
        actor,
        action="catalogue_preview_opened",
        module="catalogues",
        identifier=catalogue.slug,
        details=audit_details,
    )
    if selected_number is not None:
        _audit(
            db,
            request,
            actor,
            action="catalogue_version_selected",
            module="catalogues",
            identifier=catalogue.slug,
            details=audit_details,
        )
    db.commit()
    return presentation


@router.post(
    "/catalogues/{catalogue_id}/publish",
    response_model=CatalogueVersionResponse,
)
def publish_catalogue(
    catalogue_id: uuid.UUID,
    request: Request,
    actor: User = Depends(require_permission("catalogues.publish")),
    db: Session = Depends(get_db),
) -> CatalogueVersionResponse:
    catalogue = _get_scoped_catalogue(db, catalogue_id, actor, "catalogues.publish")
    if not catalogue.product_links:
        raise HTTPException(status_code=422, detail="Add products before publishing.")
    snapshot = _build_snapshot(db, catalogue, strict_prices=True)
    next_version = (
        db.scalar(
            select(func.max(CatalogueVersion.version_number)).where(
                CatalogueVersion.catalogue_id == catalogue.id
            )
        )
        or 0
    ) + 1
    snapshot["version"] = next_version
    version = CatalogueVersion(
        catalogue_id=catalogue.id,
        version_number=next_version,
        snapshot=snapshot,
        published_by_id=actor.id,
        published_at=_now(),
    )
    catalogue.version = next_version
    catalogue.status = "published"
    catalogue.published_by_id = actor.id
    catalogue.published_at = version.published_at
    catalogue.updated_by_id = actor.id
    catalogue.revision += 1
    db.add(version)
    db.flush()
    from app.catalogue_share_links import ensure_share_links_for_catalogue
    ensure_share_links_for_catalogue(db, catalogue, actor.id)
    _audit(
        db,
        request,
        actor,
        action="catalogue_published",
        module="catalogues",
        identifier=catalogue.slug,
        details={"version": next_version, "show_prices": snapshot["show_prices"], "videos_included": sum(bool(item.get("video")) for item in snapshot["products"])},
    )
    if any(item.get("video") for item in snapshot["products"]):
        _audit(db, request, actor, action="product_videos_included_in_catalogue_version", module="product_videos", identifier=catalogue.slug, details={"catalogue_id": str(catalogue.id), "version": next_version, "video_ids": [item["video"]["id"] for item in snapshot["products"] if item.get("video")]})
    db.commit()
    return CatalogueVersionResponse.model_validate(version, from_attributes=True)


@router.get(
    "/catalogues/{catalogue_id}/versions",
    response_model=list[CatalogueVersionResponse],
)
def list_catalogue_versions(
    catalogue_id: uuid.UUID,
    actor: User = Depends(require_permission("catalogues.view_all_versions")),
    db: Session = Depends(get_db),
) -> list[CatalogueVersionResponse]:
    _get_scoped_catalogue(db, catalogue_id, actor, "catalogues.view_all_versions")
    versions = db.scalars(
        select(CatalogueVersion)
        .where(CatalogueVersion.catalogue_id == catalogue_id)
        .order_by(CatalogueVersion.version_number.desc())
    ).all()
    return [
        CatalogueVersionResponse.model_validate(item, from_attributes=True)
        for item in versions
    ]


@router.get(
    "/catalogues/{catalogue_id}/versions/{version_number}",
    response_model=CatalogueVersionResponse,
)
def get_catalogue_version(
    catalogue_id: uuid.UUID,
    version_number: int,
    actor: User = Depends(require_permission("catalogues.view_all_versions")),
    db: Session = Depends(get_db),
) -> CatalogueVersionResponse:
    _get_scoped_catalogue(db, catalogue_id, actor, "catalogues.view_all_versions")
    version = db.scalar(
        select(CatalogueVersion).where(
            CatalogueVersion.catalogue_id == catalogue_id,
            CatalogueVersion.version_number == version_number,
        )
    )
    if not version:
        raise HTTPException(status_code=404, detail="Catalogue version not found.")
    return CatalogueVersionResponse.model_validate(version, from_attributes=True)


def _version_for_export(
    db: Session,
    catalogue_id: uuid.UUID,
    version_number: int | None,
) -> CatalogueVersion:
    query = select(CatalogueVersion).where(
        CatalogueVersion.catalogue_id == catalogue_id
    )
    if version_number is not None:
        query = query.where(CatalogueVersion.version_number == version_number)
    else:
        query = query.order_by(CatalogueVersion.version_number.desc())
    version = db.scalar(query)
    if not version:
        raise HTTPException(
            status_code=404,
            detail="Publish the catalogue before exporting it.",
        )
    return version


@router.get("/catalogues/{catalogue_id}/export/excel")
def export_catalogue_excel(
    catalogue_id: uuid.UUID,
    request: Request,
    version_number: int | None = None,
    actor: User = Depends(require_permission("catalogues.export_excel")),
    db: Session = Depends(get_db),
):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    catalogue = _get_scoped_catalogue(db, catalogue_id, actor, "catalogues.export_excel")
    version = _version_for_export(db, catalogue_id, version_number)
    snapshot = version.snapshot
    allowed_lists = allowed_price_list_ids(db, actor)
    include_prices = bool(
        snapshot.get("show_prices")
        and has_permission(actor, "prices.view")
        and (allowed_lists is None or snapshot.get("price_list_id") in allowed_lists)
    )
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Catalogue"
    headers = ["Order", "SKU", "Product", "Brand", "Category", "Description"]
    if include_prices:
        headers.extend(["Price", "Currency"])
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="176536")
    product_ids = []
    for item in snapshot.get("products", []):
        try:
            if item.get("id"):
                product_ids.append(uuid.UUID(str(item["id"])))
        except (ValueError, TypeError):
            continue
    active_ids = set(db.scalars(select(Product.id).where(Product.id.in_(product_ids), Product.status == "active"))) if product_ids else set()
    for item in snapshot.get("products", []):
        try:
            if uuid.UUID(str(item.get("id"))) not in active_ids:
                continue
        except (ValueError, TypeError):
            continue
        row = [
            item.get("sort_order"),
            item.get("sku"),
            item.get("name"),
            item.get("brand"),
            ", ".join(item.get("categories", [])),
            item.get("description"),
        ]
        if include_prices:
            row.extend([item.get("price"), item.get("currency")])
        sheet.append(row)
    sheet.freeze_panes = "A2"
    for column in sheet.columns:
        width = min(max(len(str(cell.value or "")) for cell in column) + 2, 55)
        sheet.column_dimensions[column[0].column_letter].width = width
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    _audit(
        db,
        request,
        actor,
        action="catalogue_exported_excel",
        module="catalogues",
        identifier=catalogue.slug,
        details={"version": version.version_number, "prices_included": include_prices},
    )
    db.commit()
    filename = f"{catalogue.slug}-v{version.version_number}.xlsx"
    return StreamingResponse(
        output,
        media_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _pdf_font_names() -> tuple[str, str]:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    candidates = [
        (
            Path(__file__).parent / "assets" / "fonts" / "NotoSansThai-Regular.ttf",
            Path(__file__).parent / "assets" / "fonts" / "NotoSansThai-Bold.ttf",
        ),
        (Path("C:/Windows/Fonts/tahoma.ttf"), Path("C:/Windows/Fonts/tahomabd.ttf")),
        (
            Path("/usr/share/fonts/truetype/noto/NotoSansThai-Regular.ttf"),
            Path("/usr/share/fonts/truetype/noto/NotoSansThai-Bold.ttf"),
        ),
    ]
    for regular_path, bold_path in candidates:
        if not regular_path.exists():
            continue
        if "GMSCatalogue" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont("GMSCatalogue", str(regular_path)))
        resolved_bold = bold_path if bold_path.exists() else regular_path
        if "GMSCatalogueBold" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont("GMSCatalogueBold", str(resolved_bold)))
        return "GMSCatalogue", "GMSCatalogueBold"
    return "Helvetica", "Helvetica-Bold"


def _pdf_product_image(image_url: str | None, width: float, height: float):
    from reportlab.lib import colors
    from reportlab.lib.utils import ImageReader
    from reportlab.platypus import Image, Table, TableStyle

    if image_url and image_url.startswith("/uploads/"):
        upload_root = Path(settings.upload_dir).resolve()
        relative = image_url.removeprefix("/uploads/")
        image_path = (upload_root / relative).resolve()
        if upload_root in image_path.parents and image_path.exists():
            try:
                printable_path = _pdf_product_image_cache_path(image_path, relative)
                image_width, image_height = ImageReader(str(printable_path)).getSize()
                scale = min(width / image_width, height / image_height)
                return Image(
                    str(printable_path),
                    width=image_width * scale,
                    height=image_height * scale,
                )
            except Exception:
                pass
    placeholder = Table([["Image"]], colWidths=[width], rowHeights=[height])
    placeholder.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#EEF5F0")),
                ("TEXTCOLOR", (0, 0), (-1, -1), colors.HexColor("#176536")),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("FONTSIZE", (0, 0), (-1, -1), 16),
                ("BOX", (0, 0), (-1, -1), 0.4, colors.HexColor("#D8E5DC")),
            ]
        )
    )
    return placeholder


def _cache_lock(key: str) -> threading.Lock:
    with _PDF_CACHE_LOCKS_GUARD:
        return _PDF_CACHE_LOCKS.setdefault(key, threading.Lock())


def _pdf_cache_root() -> Path:
    root = Path(settings.catalogue_pdf_cache_dir).resolve()
    root.mkdir(parents=True, exist_ok=True)
    return root


def _pdf_product_image_cache_path(source: Path, relative: str) -> Path:
    """Create a small, print-ready JPEG once instead of embedding ERP originals."""

    from PIL import Image, ImageOps

    stat = source.stat()
    digest = hashlib.sha256(
        f"{relative}|{stat.st_size}|{stat.st_mtime_ns}|pdf-image-v3".encode()
    ).hexdigest()
    target = _pdf_cache_root() / "product-images" / f"{digest}.jpg"
    if target.is_file():
        target.touch()
        return target
    with _cache_lock(f"image:{digest}"):
        if target.is_file():
            target.touch()
            return target
        target.parent.mkdir(parents=True, exist_ok=True)
        with Image.open(source) as opened:
            # Resize while the decoder still owns the source image. Pillow can
            # use format-specific reduced decoding here instead of allocating a
            # full-resolution RGBA copy for every ERP image.
            image = ImageOps.exif_transpose(opened)
            image.thumbnail((300, 300), Image.Resampling.LANCZOS)
            if image.mode in {"RGBA", "LA"} or "transparency" in image.info:
                rgba = image.convert("RGBA")
                flattened = Image.new("RGB", rgba.size, "white")
                flattened.paste(rgba, mask=rgba.getchannel("A"))
            else:
                flattened = image.convert("RGB")
            temporary = target.with_name(f"{target.stem}-{uuid.uuid4().hex}.tmp")
            flattened.save(temporary, format="JPEG", quality=68, subsampling=2)
        os.replace(temporary, target)
    return target


def _prepare_pdf_product_images(presentation: CataloguePreviewResponse) -> None:
    """Warm unique print thumbnails concurrently before ReportLab layout."""

    upload_root = Path(settings.upload_dir).resolve()
    sources: dict[str, Path] = {}
    for product in presentation.products:
        image_url = product.main_image_url
        if not image_url or not image_url.startswith("/uploads/"):
            continue
        relative = image_url.removeprefix("/uploads/")
        candidate = (upload_root / relative).resolve()
        if upload_root in candidate.parents and candidate.is_file():
            sources.setdefault(relative, candidate)
    if not sources:
        return

    def prepare(item: tuple[str, Path]) -> None:
        relative, source = item
        try:
            _pdf_product_image_cache_path(source, relative)
        except Exception:
            # The card renderer will use its normal placeholder if one source
            # image is malformed; other products should still be downloadable.
            return

    workers = min(4, len(sources))
    with ThreadPoolExecutor(max_workers=workers, thread_name_prefix="pdf-image") as executor:
        list(executor.map(prepare, sources.items()))


def _local_media_path(media_url: str | None) -> Path | None:
    if not media_url or not media_url.startswith("/uploads/"):
        return None
    root = Path(settings.upload_dir).resolve()
    candidate = (root / media_url.removeprefix("/uploads/")).resolve()
    if root not in candidate.parents or not candidate.is_file():
        return None
    return candidate


def _cover_asset_path(asset) -> Path | None:
    packaged_path = getattr(asset, "local_path", None)
    if packaged_path:
        candidate = Path(packaged_path).resolve()
        return candidate if candidate.is_file() else None
    if getattr(asset, "storage_key", None):
        try:
            candidate = cover_storage.resolve(asset.storage_key)
            return candidate if candidate.is_file() else None
        except HTTPException:
            return None
    return _local_media_path(getattr(asset, "file_url", None))


def _open_cover_asset(asset):
    from PIL import Image

    path = _cover_asset_path(asset)
    if not path:
        return None
    try:
        if str(getattr(asset, "mime_type", "")).casefold() == "image/svg+xml":
            import cairosvg

            png = cairosvg.svg2png(url=str(path), output_width=max(int(asset.width), 1), output_height=max(int(asset.height), 1))
            return Image.open(io.BytesIO(png)).convert("RGBA")
        return Image.open(path).convert("RGBA")
    except Exception:
        return None


def _cover_page_pdf(
    cover,
    page_size: tuple[float, float],
    *,
    hidden_labels: set[str] | None = None,
) -> io.BytesIO:
    """Render the saved responsive cover configuration as one PDF page."""
    from PIL import Image, ImageColor, ImageDraw, ImageEnhance, ImageFont
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas as pdf_canvas

    canvas_width = 1400 if page_size[0] <= page_size[1] else 1980
    canvas_height = round(canvas_width * page_size[1] / page_size[0])
    try:
        background_rgb = ImageColor.getrgb(cover.background_color)
    except ValueError:
        background_rgb = (22, 79, 53)
    image = Image.new("RGB", (canvas_width, canvas_height), background_rgb)
    assets = list(cover.assets)
    brand_asset = next((item for item in assets if item.asset_type == "brand_logo"), None)
    company_asset = next((item for item in assets if item.asset_type == "secondary_logo"), None)
    packaged_logo = default_logo_path()
    use_packaged_logo = (
        packaged_logo is not None
        and cover.default_company_logo_version == LOGO_VERSION
        and ((cover.show_brand_logo and brand_asset is None) or (cover.show_company_logo and brand_asset is not None and company_asset is None))
    )
    if use_packaged_logo:
        is_primary = brand_asset is None
        assets.append(SimpleNamespace(
            asset_type="default_company_logo", local_path=packaged_logo, mime_type="image/jpeg",
            width=447, height=447, alt_text=LOGO_ALT_TEXT,
            position_x_percent=15 if is_primary else 85, position_y_percent=13,
            width_percent=24 if is_primary else 18, height_percent=16 if is_primary else 12,
            opacity=1, rotation=0, z_index=10 if is_primary else 11,
        ))
    assets.sort(key=lambda item: item.z_index)
    background_type = "full_cover" if cover.cover_mode == "full_image" else "background"
    background_asset = next((item for item in assets if item.asset_type == background_type), None)
    if background_asset:
        source = _open_cover_asset(background_asset)
        if source:
            fit = cover.background_fit
            if fit == "contain":
                source.thumbnail((canvas_width, canvas_height), Image.Resampling.LANCZOS)
                image.paste(source.convert("RGB"), ((canvas_width - source.width) // 2, (canvas_height - source.height) // 2))
            else:
                scale = max(canvas_width / source.width, canvas_height / source.height)
                resized = source.resize((round(source.width * scale), round(source.height * scale)), Image.Resampling.LANCZOS)
                focus_x = background_asset.position_x_percent / 100
                focus_y = background_asset.position_y_percent / 100
                left = round((resized.width - canvas_width) * focus_x)
                top = round((resized.height - canvas_height) * focus_y)
                image.paste(resized.crop((left, top, left + canvas_width, top + canvas_height)).convert("RGB"), (0, 0))

    if cover.overlay_opacity > 0:
        try:
            overlay_rgb = ImageColor.getrgb(cover.overlay_color)
        except ValueError:
            overlay_rgb = (8, 31, 20)
        overlay = Image.new("RGBA", image.size, (*overlay_rgb, round(255 * cover.overlay_opacity)))
        image = Image.alpha_composite(image.convert("RGBA"), overlay)
    else:
        image = image.convert("RGBA")

    def paste_cover_asset(asset) -> None:
        if asset.asset_type in {"background", "full_cover"}:
            return
        if asset.asset_type == "brand_logo" and not cover.show_brand_logo:
            return
        if asset.asset_type == "secondary_logo" and not cover.show_company_logo:
            return
        source = _open_cover_asset(asset)
        if not source:
            return
        max_width = max(1, round(canvas_width * asset.width_percent / 100))
        max_height = max(1, round(canvas_height * asset.height_percent / 100))
        scale = min(max_width / source.width, max_height / source.height)
        resized = source.resize((max(1, round(source.width * scale)), max(1, round(source.height * scale))), Image.Resampling.LANCZOS)
        if asset.opacity < 1:
            alpha = ImageEnhance.Brightness(resized.getchannel("A")).enhance(asset.opacity)
            resized.putalpha(alpha)
        if asset.rotation:
            resized = resized.rotate(-asset.rotation, expand=True, resample=Image.Resampling.BICUBIC)
        x = round(canvas_width * asset.position_x_percent / 100 - resized.width / 2)
        y = round(canvas_height * asset.position_y_percent / 100 - resized.height / 2)
        image.alpha_composite(resized, (x, y))

    for asset in assets:
        paste_cover_asset(asset)

    draw = ImageDraw.Draw(image)
    font_candidates = [Path("C:/Windows/Fonts/tahomabd.ttf"), Path(__file__).parent / "assets" / "fonts" / "NotoSansThai-Bold.ttf"]
    regular_candidates = [Path("C:/Windows/Fonts/tahoma.ttf"), Path(__file__).parent / "assets" / "fonts" / "NotoSansThai-Regular.ttf"]

    def font(paths: list[Path], size: int):
        for path in paths:
            if path.exists():
                return ImageFont.truetype(str(path), max(8, size))
        return ImageFont.load_default()

    def draw_positioned_text(value: str, *, x_percent: float, y_percent: float, width_percent: float, size: int, color: str, alignment: str, bold: bool) -> None:
        if not value:
            return
        resolved_font = font(font_candidates if bold else regular_candidates, round(size * canvas_width / 1100))
        max_width = canvas_width * width_percent / 100
        words = value.split()
        lines: list[str] = []
        current = ""
        for word in words:
            attempt = f"{current} {word}".strip()
            if current and draw.textbbox((0, 0), attempt, font=resolved_font)[2] > max_width:
                lines.append(current)
                current = word
            else:
                current = attempt
        if current:
            lines.append(current)
        rendered = "\n".join(lines)
        spacing = max(4, round(size * canvas_width / 1100 * 0.12))
        box = draw.multiline_textbbox((0, 0), rendered, font=resolved_font, spacing=spacing, align=alignment)
        text_width = box[2] - box[0]
        x = canvas_width * x_percent / 100
        if alignment == "center":
            x -= text_width / 2
        elif alignment == "right":
            x -= text_width
        y = canvas_height * y_percent / 100
        draw.multiline_text((round(x), round(y)), rendered, font=resolved_font, fill=color, spacing=spacing, align=alignment)

    title_parts = []
    if cover.show_catalogue_name:
        title_parts.append(cover.catalogue_name)
    if cover.show_catalogue_year and cover.catalogue_year:
        title_parts.append(cover.catalogue_year)
    def draw_title() -> None:
        draw_positioned_text("\n".join(title_parts), x_percent=cover.title_position_x_percent,
                             y_percent=cover.title_position_y_percent, width_percent=cover.title_width_percent,
                             size=cover.title_font_size, color=cover.title_color,
                             alignment=cover.title_alignment, bold=True)

    draw_title()
    for asset in assets:
        if asset.z_index > cover.title_z_index:
            paste_cover_asset(asset)
    hidden_labels = {value.strip().casefold() for value in (hidden_labels or set()) if value.strip()}
    subtitle = str(cover.subtitle or "").strip()
    if cover.show_subtitle and subtitle.casefold() not in hidden_labels:
        draw_positioned_text(subtitle, x_percent=cover.subtitle_position_x_percent,
                             y_percent=cover.subtitle_position_y_percent, width_percent=80,
                             size=cover.subtitle_font_size, color=cover.subtitle_color,
                             alignment=cover.title_alignment, bold=False)
        for asset in assets:
            if asset.z_index > 25:
                paste_cover_asset(asset)
    if cover.title_z_index > 25:
        draw_title()
        for asset in assets:
            if asset.z_index > cover.title_z_index:
                paste_cover_asset(asset)

    rendered_image = io.BytesIO()
    image.convert("RGB").save(
        rendered_image,
        format="JPEG",
        quality=88,
        subsampling=1,
    )
    rendered_image.seek(0)
    result = io.BytesIO()
    canvas = pdf_canvas.Canvas(result, pagesize=page_size)
    canvas.drawImage(ImageReader(rendered_image), 0, 0, width=page_size[0], height=page_size[1], preserveAspectRatio=False)
    canvas.showPage()
    canvas.save()
    result.seek(0)
    return result


def _build_catalogue_pdf(
    presentation: CataloguePreviewResponse,
    *,
    paper_size: str,
    orientation: str,
    include_cover: bool,
    include_table_of_contents: bool,
) -> io.BytesIO:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        Image,
        KeepInFrame,
        KeepTogether,
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    if paper_size != CATALOGUE_PDF_CONFIG.paper_size:
        raise HTTPException(status_code=422, detail="Only A4 paper is currently supported.")
    if orientation != CATALOGUE_PDF_CONFIG.orientation:
        raise HTTPException(status_code=422, detail="Catalogue PDFs are always A4 landscape.")
    page_size = CATALOGUE_PDF_CONFIG.reportlab_page_size
    _prepare_pdf_product_images(presentation)
    regular_font, bold_font = _pdf_font_names()
    output = io.BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=page_size,
        leftMargin=CATALOGUE_PDF_CONFIG.content_margin_mm * mm,
        rightMargin=CATALOGUE_PDF_CONFIG.content_margin_mm * mm,
        topMargin=14 * mm,
        bottomMargin=14 * mm,
        title=presentation.title,
        author="Catalogue Management Platform",
    )
    base_styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "CatalogueTitle",
        parent=base_styles["Title"],
        fontName=bold_font,
        fontSize=24,
        leading=31,
        textColor=colors.HexColor("#10261A"),
        alignment=TA_LEFT,
        spaceAfter=8,
        wordWrap="CJK",
    )
    body_style = ParagraphStyle(
        "CatalogueBody",
        parent=base_styles["BodyText"],
        fontName=regular_font,
        fontSize=8.5,
        leading=12,
        textColor=colors.HexColor("#53645A"),
        wordWrap="CJK",
    )
    small_style = ParagraphStyle(
        "CatalogueSmall",
        parent=body_style,
        fontSize=7,
        leading=9,
        textColor=colors.HexColor("#66776D"),
    )
    product_title_style = ParagraphStyle(
        "ProductTitle",
        parent=body_style,
        fontName=bold_font,
        fontSize=10,
        leading=13,
        textColor=colors.HexColor("#10261A"),
        spaceAfter=3,
    )
    category_title_style = ParagraphStyle(
        "CategoryTitle",
        parent=title_style,
        fontSize=20,
        leading=25,
        keepWithNext=True,
    )
    price_style = ParagraphStyle(
        "ProductPrice",
        parent=body_style,
        fontName=bold_font,
        fontSize=11,
        textColor=colors.HexColor("#176536"),
        spaceBefore=4,
    )

    def page_footer(canvas, doc) -> None:
        canvas.saveState()
        canvas.setFillColor(colors.HexColor("#176536"))
        canvas.rect(0, page_size[1] - 7 * mm, page_size[0], 7 * mm, fill=1, stroke=0)
        canvas.setFillColor(colors.HexColor("#66776D"))
        canvas.setFont(regular_font, 7)
        canvas.drawString(13 * mm, 8 * mm, presentation.title)
        canvas.drawRightString(
            page_size[0] - 13 * mm,
            8 * mm,
            f"Page {doc.page}",
        )
        canvas.restoreState()

    story = []
    version_label = (
        f"Version {presentation.version} · Published"
        if presentation.version is not None
        else "Draft Preview · Not published"
    )
    requested_cover = include_cover
    hidden_cover_labels = {
        str(presentation.audience or ""),
        str(presentation.price_list.name if presentation.price_list else ""),
    }
    rendered_cover = (
        _cover_page_pdf(
            presentation.cover,
            page_size,
            hidden_labels=hidden_cover_labels,
        )
        if requested_cover and presentation.cover
        else None
    )
    uploaded_cover = None
    cover_is_pdf = False
    cover_image_path = None
    include_cover = False

    if include_cover and cover_image_path:
        from reportlab.lib.utils import ImageReader

        image_width, image_height = ImageReader(str(cover_image_path)).getSize()
        available_width = page_size[0] - document.leftMargin - document.rightMargin
        available_height = page_size[1] - document.topMargin - document.bottomMargin - 12 * mm
        scale = min(available_width / image_width, available_height / image_height)
        story.append(Image(str(cover_image_path), width=image_width * scale, height=image_height * scale, hAlign="CENTER"))
        if uploaded_cover.show_title_overlay:
            story.extend([Spacer(1, 4 * mm), Paragraph(escape(presentation.title), title_style)])
        if uploaded_cover.subtitle:
            story.append(Paragraph(escape(uploaded_cover.subtitle), body_style))
        story.append(PageBreak())
    elif include_cover and not cover_is_pdf:
        logo = Table([["", "CATALOGUE"]], colWidths=[13 * mm, 58 * mm])
        logo.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (0, 0), colors.HexColor("#176536")),
                    ("TEXTCOLOR", (0, 0), (0, 0), colors.white),
                    ("TEXTCOLOR", (1, 0), (1, 0), colors.HexColor("#176536")),
                    ("FONTNAME", (0, 0), (-1, -1), bold_font),
                    ("FONTSIZE", (0, 0), (0, 0), 14),
                    ("FONTSIZE", (1, 0), (1, 0), 10),
                    ("ALIGN", (0, 0), (0, 0), "CENTER"),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("LEFTPADDING", (1, 0), (1, 0), 7),
                ]
            )
        )
        story.extend(
            [
                logo,
                Spacer(1, 20 * mm),
                Paragraph(escape(presentation.title), title_style),
                Paragraph(escape(presentation.description or "Product catalogue"), body_style),
                Spacer(1, 8 * mm),
                Paragraph(escape(version_label), body_style),
                Paragraph(
                    escape(
                        f"Generated {presentation.generated_at.strftime('%d %b %Y %H:%M')} · "
                        f"{presentation.product_count} products"
                    ),
                    small_style,
                ),
                PageBreak(),
            ]
        )
    elif not cover_is_pdf:
        story.extend(
            [
                Paragraph(escape(presentation.title), title_style),
                Paragraph(escape(version_label), body_style),
                Spacer(1, 6 * mm),
            ]
        )

    usable_width = page_size[0] - document.leftMargin - document.rightMargin
    gap = 4 * mm
    erp_detail_cards = presentation.product_card_style == "erp_detail"
    card_theme = (
        presentation.product_card_theme.model_dump()
        if presentation.product_card_theme is not None
        else {}
    )
    card_variant = str(card_theme.get("variant") or "rounded")
    card_accent = colors.HexColor(str(card_theme.get("accent_color") or "#FFF45A"))
    card_strong = colors.HexColor(str(card_theme.get("strong_color") or "#176536"))
    card_surface = colors.HexColor(str(card_theme.get("surface_color") or "#F8FFD8"))
    card_border = colors.HexColor(str(card_theme.get("border_color") or "#D9E3BA"))
    card_text = colors.HexColor(str(card_theme.get("text_color") or "#263425"))
    columns = 2 if erp_detail_cards else CATALOGUE_PDF_CONFIG.product_columns
    product_rows_per_page = (
        3 if erp_detail_cards else CATALOGUE_PDF_CONFIG.product_rows_per_page
    )
    card_width = (usable_width - gap * (columns - 1)) / columns

    def product_card(product):
        category = product.section_title or product.category_name or product.brand or "Product"
        product_name = product.name_th or product.name_en or product.name
        secondary_name = (
            product.name_en
            if product.name_th and product.name_en and product.name_en != product.name_th
            else None
        )
        description = product.description or product.long_description or "Product details coming soon."
        if len(description) > 220:
            description = description[:217].rstrip() + "..."
        if erp_detail_cards:
            from reportlab.graphics.barcode import createBarcodeDrawing

            header_style = ParagraphStyle(
                "ErpCardHeader",
                parent=product_title_style,
                fontSize=10,
                leading=12,
                spaceAfter=0,
                textColor=card_text,
            )
            fact_header_style = ParagraphStyle(
                "ErpCardFactHeader",
                parent=small_style,
                fontName=bold_font,
                fontSize=6.5,
                leading=8,
                textColor=card_text,
                alignment=TA_CENTER,
            )
            fact_style = ParagraphStyle(
                "ErpCardFact",
                parent=small_style,
                fontSize=6.2,
                leading=8,
                textColor=card_text,
                alignment=TA_CENTER,
            )
            media_width = 43 * mm
            details_width = card_width - media_width - 10 * mm
            image = _pdf_product_image(
                product.main_image_url,
                media_width - 5 * mm,
                28 * mm,
            )
            media = Table(
                [[image], [Paragraph(escape(product.code or "—"), fact_style)]],
                colWidths=[media_width],
            )
            media.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, -1), colors.white),
                        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 1.5 * mm),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 1.5 * mm),
                        ("TOPPADDING", (0, 0), (-1, -1), 1.5 * mm),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 1.5 * mm),
                    ]
                )
            )
            facts = Table(
                [
                    [
                        Paragraph("Code", fact_header_style),
                        Paragraph("Barcode", fact_header_style),
                        Paragraph("Stock", fact_header_style),
                    ],
                    [
                        Paragraph(escape(product.code or "—"), fact_style),
                        Paragraph(escape(product.barcode or "—"), fact_style),
                        Paragraph(
                            escape(
                                f"{product.stock_quantity:,}"
                                if product.stock_quantity is not None
                                else "—"
                            ),
                            fact_style,
                        ),
                    ],
                ],
                colWidths=[details_width * .27, details_width * .48, details_width * .25],
            )
            facts.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), card_accent),
                        ("BACKGROUND", (0, 1), (-1, 1), colors.white),
                        ("GRID", (0, 0), (-1, -1), .35, card_border),
                        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 1 * mm),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 1 * mm),
                        ("TOPPADDING", (0, 0), (-1, -1), 1.2 * mm),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 1.2 * mm),
                    ]
                )
            )
            details_content = [facts]
            if product.barcode:
                barcode = createBarcodeDrawing(
                    "Code128",
                    value=str(product.barcode),
                    barHeight=8 * mm,
                    barWidth=.22 * mm,
                    humanReadable=True,
                )
                details_content.extend(
                    [
                        Spacer(1, 2 * mm),
                        KeepInFrame(
                            details_width,
                            12 * mm,
                            [barcode],
                            mode="shrink",
                            hAlign="RIGHT",
                        ),
                    ]
                )
            body = Table(
                [[media, details_content]],
                colWidths=[media_width, details_width],
            )
            body.setStyle(
                TableStyle(
                    [
                        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 1.5 * mm),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 1.5 * mm),
                        ("TOPPADDING", (0, 0), (-1, -1), 1.5 * mm),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 1.5 * mm),
                    ]
                )
            )
            footer_copy = description[:115].rstrip()
            if len(description) > 115:
                footer_copy += "..."
            footer_copy_style = (
                ParagraphStyle(
                    "ErpCardContrastFooter",
                    parent=small_style,
                    textColor=colors.white,
                )
                if card_variant == "contrast"
                else small_style
            )
            footer_left = [
                Paragraph(escape(footer_copy), footer_copy_style),
                Paragraph(escape(product.brand or category), fact_style),
            ]
            footer_right = ""
            if presentation.show_prices and product.price is not None:
                footer_right = Paragraph(
                    escape(
                        f"{product.currency or presentation.currency or ''} {product.price}".strip()
                    ),
                    ParagraphStyle(
                        "ErpCardPrice",
                        parent=price_style,
                        alignment=TA_CENTER,
                        textColor=card_text,
                        spaceBefore=0,
                    ),
                )
            footer = Table(
                [[footer_left, footer_right]],
                colWidths=[card_width - 43 * mm, 38 * mm],
            )
            footer.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (1, 0), (1, 0), card_accent),
                        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 2 * mm),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 2 * mm),
                        ("TOPPADDING", (0, 0), (-1, -1), 1.5 * mm),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 1.5 * mm),
                    ]
                )
            )
            card = Table(
                [
                    [Paragraph(escape(product_name), header_style)],
                    [body],
                    [footer],
                ],
                colWidths=[card_width - 3 * mm],
            )
            card_commands = [
                ("BACKGROUND", (0, 0), (0, 0), card_accent),
                (
                    "BACKGROUND",
                    (0, 1),
                    (0, -1),
                    card_strong if card_variant == "contrast" else card_surface,
                ),
                (
                    "BOX",
                    (0, 0),
                    (-1, -1),
                    1.5 if card_variant == "framed" else .55,
                    card_strong if card_variant == "framed" else card_border,
                ),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 2.5 * mm),
                ("RIGHTPADDING", (0, 0), (-1, -1), 2.5 * mm),
                ("TOPPADDING", (0, 0), (0, 0), 2.2 * mm),
                ("BOTTOMPADDING", (0, 0), (0, 0), 2.2 * mm),
                ("TOPPADDING", (0, 1), (0, -1), 1 * mm),
                ("BOTTOMPADDING", (0, 1), (0, -1), 1 * mm),
            ]
            if card_variant == "editorial":
                card_commands.extend(
                    [
                        ("LINEBEFORE", (0, 0), (0, -1), 4, card_strong),
                        ("BOX", (0, 0), (-1, -1), .25, card_border),
                    ]
                )
            elif card_variant == "framed":
                card_commands.append(("BACKGROUND", (0, 0), (0, 0), colors.white))
            elif card_variant == "minimal":
                card_commands.extend(
                    [
                        ("BACKGROUND", (0, 0), (0, -1), colors.white),
                        ("BOX", (0, 0), (-1, -1), .25, card_border),
                        ("LINEBELOW", (0, 0), (0, 0), 2.5, card_accent),
                    ]
                )
            card.setStyle(TableStyle(card_commands))
            return KeepInFrame(
                card_width - 2 * mm,
                152,
                [card],
                mode="shrink",
                hAlign="LEFT",
                vAlign="TOP",
            )
        content = [
            [_pdf_product_image(product.main_image_url, card_width - 10 * mm, 30 * mm)],
            [Paragraph(escape(category.upper()), small_style)],
            [Paragraph(escape(product_name), product_title_style)],
        ]
        if secondary_name:
            content.append([Paragraph(escape(secondary_name), small_style)])
        content.extend(
            [
                [Paragraph(escape(description), body_style)],
                [Paragraph(escape(product.code), small_style)],
            ]
        )
        facts = []
        if product.barcode:
            facts.append(f"Barcode {product.barcode}")
        if product.stock_quantity is not None:
            facts.append(f"Stock {product.stock_quantity:,}")
        if facts:
            content.append([Paragraph(escape("  •  ".join(facts)), small_style)])
        if product.video:
            content.append([
                Paragraph("▶ Product video available online", ParagraphStyle(
                    "ProductVideoAvailable", parent=small_style, fontName=bold_font,
                    textColor=colors.HexColor("#176536"), spaceBefore=3,
                ))
            ])
        if presentation.show_prices and product.price is not None:
            content.append(
                [
                    Paragraph(
                        escape(f"{product.currency or presentation.currency or ''} {product.price}".strip()),
                        price_style,
                    )
                ]
            )
        card = Table(content, colWidths=[card_width - 8 * mm])
        card.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, -1), colors.white),
                    ("BOX", (0, 0), (-1, -1), 0.45, colors.HexColor("#D8E5DC")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("ALIGN", (0, 0), (0, 0), "CENTER"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 4 * mm),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 4 * mm),
                    ("TOPPADDING", (0, 0), (-1, -1), 2.5 * mm),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5 * mm),
                ]
            )
        )
        return KeepInFrame(
            card_width - 2 * mm,
            230,
            [card],
            mode="shrink",
            hAlign="LEFT",
            vAlign="TOP",
        )

    if include_table_of_contents and presentation.categories:
        story.append(Paragraph("Contents", title_style))
        estimated_start = 2 if requested_cover else 1
        for index, category in enumerate(presentation.categories):
            story.append(Paragraph(escape(f"{category.name}  ·  Page {estimated_start + index + 1}"), product_title_style))
        story.append(PageBreak())

    def card_grid(page_cards):
        card_rows = []
        for index in range(0, len(page_cards), columns):
            row = page_cards[index : index + columns]
            while len(row) < columns:
                row.append("")
            card_rows.append(row)
        grid = Table(
            card_rows,
            colWidths=[card_width] * columns,
            hAlign="LEFT",
            splitByRow=False,
        )
        grid.setStyle(
            TableStyle(
                [
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 0),
                    ("RIGHTPADDING", (0, 0), (-1, -1), gap),
                    ("TOPPADDING", (0, 0), (-1, -1), 0),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), gap),
                    ("RIGHTPADDING", (columns - 1, 0), (columns - 1, -1), 0),
                ]
            )
        )
        return grid

    def append_card_grid(products, *, heading=None) -> None:
        cards = [product_card(product) for product in products]
        if not cards:
            return
        cards_per_page = columns * product_rows_per_page
        page_start = 0
        if heading:
            # Keep the category identity and at least the first product row on
            # the same physical page.  Without this block, ReportLab can leave
            # the heading behind when the product grid moves to the next page.
            first_row_end = min(columns, len(cards))
            story.append(
                KeepTogether([*heading, card_grid(cards[:first_row_end])])
            )
            page_start = first_row_end
            # Use the remaining row on the category's first page when space is
            # available; ReportLab will move this complete row if it is not.
            if page_start < len(cards):
                second_row_end = min(page_start + columns, len(cards))
                story.append(card_grid(cards[page_start:second_row_end]))
                page_start = second_row_end
            if page_start < len(cards):
                story.append(PageBreak())

        while page_start < len(cards):
            page_cards = cards[page_start : page_start + cards_per_page]
            # Keeping each table to one physical page prevents ReportLab from
            # repeatedly splitting hundreds of nested product-card rows. The
            # old unbounded table became nonlinear for 2,000+ product PDFs.
            story.append(card_grid(page_cards))
            page_start += len(page_cards)
            if page_start < len(cards):
                story.append(PageBreak())

    categorized_codes: set[str] = set()
    visible_categories = [category for category in presentation.categories if category.product_count > 0]
    for category_index, category in enumerate(visible_categories):
        category_products = [product for product in presentation.products if product.category_name == category.name or category.name in product.categories]
        if not category_products:
            continue
        if category_index > 0:
            story.append(PageBreak())
        heading = []
        banner_path = _local_media_path(category.banner_url)
        if banner_path:
            from reportlab.lib.utils import ImageReader

            banner_width, banner_height = ImageReader(str(banner_path)).getSize()
            banner_scale = min(usable_width / banner_width, (28 * mm) / banner_height)
            heading.append(
                Image(
                    str(banner_path),
                    width=banner_width * banner_scale,
                    height=banner_height * banner_scale,
                    hAlign="CENTER",
                )
            )
            heading.append(Spacer(1, 3 * mm))
        heading.append(Paragraph(escape(category.name), category_title_style))
        if category.description:
            heading.append(Paragraph(escape(category.description), body_style))
            heading.append(Spacer(1, 5 * mm))
        append_card_grid(category_products, heading=heading)
        categorized_codes.update(product.code for product in category_products)
    remaining_products = [product for product in presentation.products if product.code not in categorized_codes]
    if remaining_products:
        if categorized_codes:
            story.append(PageBreak())
        append_card_grid(
            remaining_products,
            heading=[Paragraph("Products", title_style)],
        )
    if not presentation.products:
        empty_style = ParagraphStyle(
            "EmptyCatalogue",
            parent=body_style,
            alignment=TA_CENTER,
            fontSize=12,
            leading=18,
        )
        story.append(Paragraph("This catalogue does not contain any products.", empty_style))

    # The final page uses the same ReportLab document page size, so it can never
    # fall back to portrait even for empty, draft or no-price catalogues.
    story.extend(
        [
            PageBreak(),
            Spacer(1, 35 * mm),
            Paragraph(
                escape(presentation.title),
                ParagraphStyle("FinalTitle", parent=title_style, alignment=TA_CENTER),
            ),
            Paragraph(
                "End of catalogue",
                ParagraphStyle(
                    "FinalCopy", parent=body_style, alignment=TA_CENTER,
                    fontSize=11, leading=16,
                ),
            ),
        ]
    )

    try:
        document.build(story, onFirstPage=page_footer, onLaterPages=page_footer)
    except Exception as error:
        raise HTTPException(
            status_code=500,
            detail="The PDF could not be generated. Please try again.",
        ) from error
    output.seek(0)
    if rendered_cover:
        from pypdf import PdfReader, PdfWriter

        merged = io.BytesIO()
        writer = PdfWriter()
        writer.add_page(PdfReader(rendered_cover).pages[0])
        for page in PdfReader(output).pages:
            writer.add_page(page)
        writer.write(merged)
        merged.seek(0)
        return merged
    return output


def _prune_pdf_cache() -> None:
    global _PDF_CACHE_LAST_PRUNE
    now = time.time()
    if now - _PDF_CACHE_LAST_PRUNE < 3600 or not _PDF_CACHE_PRUNE_LOCK.acquire(blocking=False):
        return
    try:
        _PDF_CACHE_LAST_PRUNE = now
        cutoff = now - settings.catalogue_pdf_cache_retention_days * 86400
        for path in _pdf_cache_root().rglob("*"):
            if path.is_file() and path.stat().st_mtime < cutoff:
                path.unlink(missing_ok=True)
    finally:
        _PDF_CACHE_PRUNE_LOCK.release()


def _cached_catalogue_pdf(
    presentation: CataloguePreviewResponse,
    *,
    paper_size: str,
    orientation: str,
    include_cover: bool,
    include_table_of_contents: bool,
) -> tuple[Path, bool]:
    """Return a content-addressed PDF, generating it once when necessary."""

    cache_payload = {
        "schema": _PDF_CACHE_SCHEMA_VERSION,
        "paper_size": paper_size,
        "orientation": orientation,
        "include_cover": include_cover,
        "include_table_of_contents": include_table_of_contents,
        "presentation": presentation.model_dump(mode="json"),
    }
    digest = hashlib.sha256(
        json.dumps(
            cache_payload,
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
        ).encode("utf-8")
    ).hexdigest()
    target = _pdf_cache_root() / "documents" / digest[:2] / f"{digest}.pdf"
    if target.is_file():
        target.touch()
        return target, True
    with _cache_lock(f"pdf:{digest}"):
        if target.is_file():
            target.touch()
            return target, True
        output = _build_catalogue_pdf(
            presentation,
            paper_size=paper_size,
            orientation=orientation,
            include_cover=include_cover,
            include_table_of_contents=include_table_of_contents,
        )
        target.parent.mkdir(parents=True, exist_ok=True)
        temporary = target.with_name(f"{target.stem}-{uuid.uuid4().hex}.tmp")
        with temporary.open("xb") as stream:
            stream.write(output.getbuffer())
        os.replace(temporary, target)
    _prune_pdf_cache()
    return target, False


@router.get("/catalogues/{catalogue_id}/export/pdf")
def export_catalogue_pdf(
    catalogue_id: uuid.UUID,
    request: Request,
    version: int | None = Query(default=None, ge=1),
    version_number: int | None = Query(default=None, ge=1),
    audience_type_id: int | None = Query(default=None, ge=1),
    language: str | None = Query(default=None, pattern=r"^(en|th|en-th)$"),
    include_cover: bool = True,
    paper_size: str = Query(default="A4", pattern=r"^A4$"),
    orientation: str = Query(default=CATALOGUE_PDF_CONFIG.orientation, pattern=r"^landscape$"),
    include_table_of_contents: bool = False,
    actor: User = Depends(require_permission("catalogues.export_pdf")),
    db: Session = Depends(get_db),
):
    catalogue = _get_scoped_catalogue(db, catalogue_id, actor, "catalogues.export_pdf")
    selected_number = _requested_version(version, version_number)
    snapshot, resolved_version = _presentation_source(db, catalogue, selected_number)
    presentation = _safe_presentation(db, catalogue, snapshot, resolved_version, actor)
    presentation = _apply_user_catalogue_price_mapping(
        db, presentation, actor, audience_type_id, catalogue.brand
    )
    if language:
        presentation.language = language
    output_path, cache_hit = _cached_catalogue_pdf(
        presentation,
        paper_size=paper_size,
        orientation=orientation,
        include_cover=include_cover,
        include_table_of_contents=include_table_of_contents,
    )
    _audit(
        db,
        request,
        actor,
        action="catalogue_exported_pdf",
        module="catalogues",
        identifier=catalogue.slug,
        details={
            "catalogue_id": str(catalogue.id),
            "version": resolved_version,
            "price_list": (
                presentation.price_list.name if presentation.price_list else None
            ),
            "prices_included": presentation.show_prices,
            "language": presentation.language,
            "paper_size": paper_size,
            "orientation": orientation,
            "cache_hit": cache_hit,
        },
    )
    db.commit()
    version_label = f"v{resolved_version}" if resolved_version else "Draft"
    cover_title = presentation.title
    if presentation.cover:
        cover_title = " ".join(value for value in (presentation.cover.catalogue_name, presentation.cover.catalogue_year) if value).strip()
    safe_title = re.sub(r"[^A-Za-z0-9\u0E00-\u0E7F]+", "-", cover_title.upper()).strip("-") or "CATALOGUE"
    filename = f"{safe_title}-{version_label}.pdf"
    return FileResponse(
        output_path,
        media_type="application/pdf",
        filename=filename,
        content_disposition_type="attachment",
        headers={"X-Catalogue-PDF-Cache": "HIT" if cache_hit else "MISS"},
    )


@router.post(
    "/catalogues/{catalogue_id}/print",
    status_code=status.HTTP_204_NO_CONTENT,
)
def record_catalogue_print(
    catalogue_id: uuid.UUID,
    request: Request,
    version: int | None = Query(default=None, ge=1),
    actor: User = Depends(require_permission("catalogues.print")),
    db: Session = Depends(get_db),
) -> None:
    catalogue = _get_scoped_catalogue(db, catalogue_id, actor, "catalogues.print")
    snapshot, resolved_version = _presentation_source(db, catalogue, version)
    presentation = _safe_presentation(db, catalogue, snapshot, resolved_version, actor)
    _audit(
        db,
        request,
        actor,
        action="catalogue_printed",
        module="catalogues",
        identifier=catalogue.slug,
        details={
            "catalogue_id": str(catalogue.id),
            "version": resolved_version,
            "price_list": (
                presentation.price_list.name if presentation.price_list else None
            ),
            "prices_included": presentation.show_prices,
        },
    )
    db.commit()


@router.post(
    "/catalogues/{catalogue_id}/duplicate",
    response_model=CatalogueResponse,
    status_code=status.HTTP_201_CREATED,
)
def duplicate_catalogue(
    catalogue_id: uuid.UUID,
    request: Request,
    actor: User = Depends(require_permission("catalogues.duplicate")),
    db: Session = Depends(get_db),
) -> CatalogueResponse:
    source = _get_scoped_catalogue(db, catalogue_id, actor, "catalogues.duplicate")
    suffix = uuid.uuid4().hex[:6]
    duplicate = Catalogue(
        title=f"{source.title} Copy",
        slug=f"{source.slug}-copy-{suffix}",
        description=source.description,
        brand=source.brand,
        audience=source.audience,
        price_list_id=source.price_list_id,
        show_prices=source.show_prices,
        currency=source.currency,
        language=source.language,
        status="draft",
        valid_from=source.valid_from,
        valid_until=source.valid_until,
        is_public=False,
        owner_id=actor.id,
        created_by_id=actor.id,
        updated_by_id=actor.id,
        product_links=[
            CatalogueProduct(
                product_id=link.product_id,
                section_title=link.section_title,
                override_description=link.override_description,
                hide_price=link.hide_price,
                include_video=link.include_video,
                selected_video_id=link.selected_video_id,
                video_title_override=link.video_title_override,
                video_description_override=link.video_description_override,
                video_display_mode=link.video_display_mode,
                video_thumbnail_mode=link.video_thumbnail_mode,
                sort_order=link.sort_order,
            )
            for link in source.product_links
        ],
    )
    db.add(duplicate)
    _audit(
        db,
        request,
        actor,
        action="catalogue_duplicated",
        module="catalogues",
        identifier=duplicate.slug,
        details={"source_catalogue_id": str(source.id)},
    )
    db.commit()
    return _catalogue_response(db, _get_catalogue(db, duplicate.id), actor=actor)


@router.post("/catalogues/{catalogue_id}/archive", response_model=CatalogueResponse)
def archive_catalogue(
    catalogue_id: uuid.UUID,
    request: Request,
    actor: User = Depends(require_permission("catalogues.archive")),
    db: Session = Depends(get_db),
) -> CatalogueResponse:
    catalogue = _get_scoped_catalogue(db, catalogue_id, actor, "catalogues.archive")
    catalogue.status = "archived"
    catalogue.updated_by_id = actor.id
    catalogue.revision += 1
    _audit(
        db,
        request,
        actor,
        action="catalogue_archived",
        module="catalogues",
        identifier=catalogue.slug,
    )
    db.commit()
    return _catalogue_response(db, _get_catalogue(db, catalogue.id), actor=actor)


@router.delete(
    "/catalogues/{catalogue_id}",
    status_code=status.HTTP_204_NO_CONTENT,
)
def delete_catalogue(
    catalogue_id: uuid.UUID,
    request: Request,
    actor: User = Depends(require_permission("catalogues.delete")),
    db: Session = Depends(get_db),
) -> None:
    """Permanently delete one catalogue and its catalogue-owned records."""

    catalogue = _get_scoped_catalogue(
        db, catalogue_id, actor, "catalogues.delete"
    )
    cover_assets = list(
        db.scalars(
            select(CatalogueCoverAsset).where(
                CatalogueCoverAsset.catalogue_id == catalogue.id
            )
        )
    )
    legacy_cover = db.scalar(
        select(CatalogueCoverFile).where(
            CatalogueCoverFile.catalogue_id == catalogue.id
        )
    )
    category_settings = list(
        db.scalars(
            select(CatalogueCategorySetting).where(
                CatalogueCategorySetting.catalogue_id == catalogue.id
            )
        )
    )
    cover_keys = {
        key
        for asset in cover_assets
        for key in (asset.storage_key, asset.preview_storage_key)
        if key
    }
    if legacy_cover:
        cover_keys.update(
            key
            for key in (
                legacy_cover.storage_key,
                legacy_cover.preview_storage_key,
            )
            if key
        )
    banner_keys = {
        item.banner_storage_key
        for item in category_settings
        if item.banner_storage_key
    }
    studio_designs = list(
        db.scalars(
            select(CatalogueDesign).where(
                CatalogueDesign.catalogue_id == catalogue.id,
                CatalogueDesign.deleted_at.is_(None),
            )
        )
    )
    deleted_at = _now()
    for design in studio_designs:
        design.deleted_at = deleted_at
        design.catalogue_id = None
        design.revision += 1
        design.updated_by_id = actor.id
        _audit(
            db,
            request,
            actor,
            action="catalogue_design_deleted",
            module="catalogue_studio",
            identifier=str(design.id),
            details={"catalogue_id": str(catalogue.id)},
        )
    title = catalogue.title
    slug = catalogue.slug
    _audit(
        db,
        request,
        actor,
        action="catalogue_deleted",
        module="catalogues",
        identifier=slug,
        details={
            "catalogue_id": str(catalogue.id),
            "title": title,
            "product_count": len(catalogue.product_links),
            "version_count": len(catalogue.versions),
            "share_link_count": len(catalogue.share_links),
            "studio_design_ids": [str(design.id) for design in studio_designs],
        },
    )
    db.delete(catalogue)
    db.commit()

    # The database deletion is authoritative. Media cleanup is best-effort so
    # a transient filesystem issue cannot resurrect a deleted catalogue.
    for key in cover_keys:
        try:
            cover_storage.delete(key)
        except OSError:
            pass
    for key in banner_keys:
        try:
            storage.delete(key)
        except OSError:
            pass
